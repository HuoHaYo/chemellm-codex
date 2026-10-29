"""ChemELLM local kernel (v1).

A local-only HTTP service that wraps this repository's ``app`` toolchain
behind a stable contract also used by the Codex MCP adapter. It reuses the original
deterministic calculation engine, document generators and evidence tools instead
of rewriting them.

Design rules:
- loopback only (127.0.0.1), optional shared token;
- project state is local JSON under ``data/projects/<id>/project.json``;
- the archived SQLite database is used only as the calculation/artifact store;
- every mutating call also refreshes a read-only snapshot under ``<repo>/.chem``.
"""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import subprocess
import sys
import threading
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
HERE = Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))


if not (ROOT / "app" / "documents.py").is_file():
    raise RuntimeError("ChemELLM requires its own app/ package alongside kernel/")
DATA_ROOT = Path(os.getenv("CHEMELLM_DATA_ROOT") or (ROOT / "data"))
KERNEL_DATA = Path(os.getenv("CHEMELLM_KERNEL_DATA") or (DATA_ROOT / "kernel"))
PROJECTS_DIR = DATA_ROOT / "projects"
BUNDLES_DIR = DATA_ROOT / "bundles"
SNAPSHOT_DIR = Path(os.getenv("CHEMELLM_SNAPSHOT_DIR") or (ROOT / ".chem"))

# The archived app reads its data directory at import time.
os.environ.setdefault("CHEMELLM_DATA_DIR", str(KERNEL_DATA))
for folder in (KERNEL_DATA, PROJECTS_DIR, BUNDLES_DIR, SNAPSHOT_DIR):
    folder.mkdir(parents=True, exist_ok=True)
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from fastapi import FastAPI, HTTPException, Request  # noqa: E402
from fastapi.middleware.cors import CORSMiddleware  # noqa: E402
from fastapi.responses import FileResponse, JSONResponse  # noqa: E402
from pydantic import BaseModel, Field  # noqa: E402

from app.db import connection, execute, init_db, now, one, rows, uid  # noqa: E402
from app.documents import artifact_record, ingest as old_ingest  # noqa: E402
from app.documents import public_artifact  # noqa: E402
import drawing_tools as drawings  # noqa: E402
import pid_sheet  # noqa: E402
import pid_standard  # noqa: E402
import pid_data  # noqa: E402
import datasheet  # noqa: E402
import relief  # noqa: E402
import hazop  # noqa: E402
import compliance  # noqa: E402
import gear  # noqa: E402
import dwsim_bridge  # noqa: E402
import cad_engineering  # noqa: E402
from app.toolbox import SCHEMAS, run_tool  # noqa: E402
from app.catalog import TOOLS  # noqa: E402

import app.documents as _legacy_documents  # noqa: E402

KERNEL_VERSION = "0.1.0"
PHASE_GATE = os.getenv("CHEMELLM_PHASE_GATE", "advisory").lower()
KERNEL_TOKEN = os.getenv("CHEMELLM_KERNEL_TOKEN", "")
MAX_INTAKE = 25 * 1024 * 1024
MAX_TOOL_RESULT = 40000
_write_lock = threading.RLock()

# LibreOffice conversion with a fixed, normally-permissioned profile directory.
# The archived implementation uses Python's mkdtemp for the profile, which some
# sandboxed environments create with permissions the child process cannot use.
LO_PROFILE = KERNEL_DATA / "lo-profile"

def _local_office_pdf(path: Path) -> Path | None:
    candidates = []
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        candidates.append(found)
    candidates.extend([
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ])
    LO_PROFILE.mkdir(parents=True, exist_ok=True)
    for executable in candidates:
        if not executable or not Path(executable).is_file():
            continue
        try:
            subprocess.run(
                [executable, "-env:UserInstallation=" + LO_PROFILE.as_uri(),
                 "--headless", "--convert-to", "pdf", "--outdir", str(path.parent), str(path)],
                capture_output=True, timeout=180, check=True,
            )
        except (OSError, subprocess.SubprocessError):
            continue
        target = path.with_suffix(".pdf")
        if target.is_file():
            return target
    return None

_legacy_documents.office_pdf = _local_office_pdf

PHASES = ["intake", "requirements", "plan", "basis", "execute", "review", "bundle", "done"]
CALC_TOOLS = {
    "properties", "steam", "convert_units", "material_balance", "heat_exchanger", "flash",
    "economics", "reaction_thermo", "mixture_reaction", "distillation_shortcut",
    "pipe_pressure_drop", "pinch_analysis", "column_mesh", "aqueous_flash", "condenser_flash",
    "frontend_process", "wet_column", "wet_separation_train", "recycle_process", "fit_data",
    "cantera_equilibrium", "cantera_reactor",
}
SPECIAL_TOOLS = {"export_xlsx", "export_freecad", "pid", "pid_check", "pid_data",
                 "cad_part", "equipment_datasheet", "relief_sizing", "hazop_lopa",
                 "compliance_matrix", "dwsim"}
GENERATE_TOOLS = {"export_report", "pfd"}
EVIDENCE_TOOLS = {
    "literature_search", "pubchem", "open_literature_search", "project_search",
    "fulltext_import", "source_fetch", "document_extract", "document_read",
    "claim_save", "evidence_add", "evidence_check", "evidence_export",
}
# Minimum phase required for the mutating tools when phaseGate=enforce.
MIN_PHASE = {
    **{name: "basis" for name in CALC_TOOLS},
    **{name: "basis" for name in GENERATE_TOOLS},
    "export_xlsx": "basis",
    "export_freecad": "basis",
    "pid": "basis",
    "pid_data": "basis",
    "equipment_datasheet": "basis",
    "relief_sizing": "basis",
    "hazop_lopa": "basis",
    "compliance_matrix": "basis",
    "dwsim": "basis",
    "cad_part": "basis",
    "engineering_requirements": "intake",
    "engineering_plan": "requirements",
}

init_db()

app = FastAPI(title="ChemELLM local kernel", version=KERNEL_VERSION)

# 可选本地 HTTP 客户端的跨源配置；Codex MCP 在进程内调用相同路由。
# 若设置了 CHEMELLM_CORS_ORIGINS（逗号分隔），则只允许这些来源。
_cors_env = os.environ.get("CHEMELLM_CORS_ORIGINS", "").strip()
_cors_origins = [item.strip() for item in _cors_env.split(",") if item.strip()] or ["*"]
app.add_middleware(
    CORSMiddleware,
    allow_origins=_cors_origins,
    allow_methods=["*"],
    allow_headers=["*"],
)


# --------------------------------------------------------------------------- #
# project state
# --------------------------------------------------------------------------- #
def _utc() -> str:
    return datetime.now(timezone.utc).isoformat()


def project_dir(project_id: str) -> Path:
    return PROJECTS_DIR / project_id


def project_file(project_id: str) -> Path:
    return project_dir(project_id) / "project.json"


def load_project(project_id: str) -> dict:
    path = project_file(project_id)
    if not path.is_file():
        raise HTTPException(status_code=404, detail="项目不存在：" + project_id)
    return json.loads(path.read_text(encoding="utf-8"))


def derive_phase(project: dict) -> str:
    requirements = project.get("requirements") or []
    plan = project.get("plan") or {}
    objects = [o for o in (project.get("objects") or {}).values() if o.get("quantities")]
    calculations = project.get("calculations") or []
    files = project.get("files") or []
    generated = [f for f in files if f.get("origin") == "generated"]
    review = project.get("review") or {}
    if not requirements:
        return "intake" if not files else "requirements"
    if not plan.get("deliverables"):
        return "plan"
    if not objects:
        return "basis"
    if not calculations or not generated:
        return "execute"
    if not review.get("ready"):
        return "review"
    return "done" if project.get("bundle") else "bundle"


def snapshot_project(project: dict) -> None:
    """Write the read-only workspace snapshot consumed by panels and diffs."""
    folder = SNAPSHOT_DIR / project["id"]
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "project.json").write_text(
        json.dumps(project, ensure_ascii=False, indent=2), encoding="utf-8")
    (folder / "requirements.json").write_text(
        json.dumps(project.get("requirements") or [], ensure_ascii=False, indent=2), encoding="utf-8")
    (folder / "plan.json").write_text(
        json.dumps(project.get("plan") or {}, ensure_ascii=False, indent=2), encoding="utf-8")
    for key, value in (project.get("objects") or {}).items():
        target = folder / "objects" / (key.replace("/", "__") + ".json")
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    (folder / "calculations.json").write_text(
        json.dumps(project.get("calculations") or [], ensure_ascii=False, indent=2), encoding="utf-8")
    (folder / "files.json").write_text(
        json.dumps(project.get("files") or [], ensure_ascii=False, indent=2), encoding="utf-8")
    (folder / "reviews.json").write_text(
        json.dumps(project.get("review") or {}, ensure_ascii=False, indent=2), encoding="utf-8")


def save_project(project: dict) -> dict:
    project["updatedAt"] = _utc()
    project["phase"] = derive_phase(project)
    path = project_file(project["id"])
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(project, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)
    snapshot_project(project)
    return project


def ensure_legacy_project(project: dict) -> None:
    """The archived DB is the calculation/artifact store, so it needs the project row."""
    if one("SELECT id FROM projects WHERE id=?", (project["id"],)):
        return
    execute("INSERT INTO projects VALUES(?,?,?)", (project["id"], project["name"], now()))


def summary_of(project: dict) -> dict:
    return {
        "projectId": project["id"],
        "name": project["name"],
        "phase": project.get("phase"),
        "updatedAt": project.get("updatedAt"),
        "counts": {
            "requirements": len(project.get("requirements") or []),
            "deliverables": len((project.get("plan") or {}).get("deliverables") or []),
            "objects": len([o for o in (project.get("objects") or {}).values() if o.get("quantities")]),
            "calculations": len(project.get("calculations") or []),
            "files": len(project.get("files") or []),
        },
        "ready": bool((project.get("review") or {}).get("ready")),
    }


# --------------------------------------------------------------------------- #
# artifacts / validation helpers
# --------------------------------------------------------------------------- #
def _artifact_path(artifact_id: str) -> Path:
    row = one("SELECT * FROM artifacts WHERE id=?", (artifact_id,))
    if not row:
        raise HTTPException(status_code=404, detail="文件不存在：" + artifact_id)
    path = Path(row["path"]).resolve()
    artifacts_root = KERNEL_DATA.joinpath("artifacts").resolve()
    if not path.is_file() or not path.is_relative_to(artifacts_root):
        raise HTTPException(status_code=404, detail="文件路径无效")
    return path


def _normalize_artifacts(value: Any, found: list) -> list:
    """Collect public artifact records out of an arbitrary tool result."""
    if isinstance(value, dict):
        if "download_url" in value and "id" in value and "name" in value:
            metadata = value.get("metadata") or {}
            aid = value["id"]
            found.append({
                "artifactId": aid,
                "name": value.get("name"),
                "kind": value.get("kind"),
                "bytes": metadata.get("bytes"),
                "sha256": metadata.get("sha256"),
                "downloadPath": f"/v1/artifacts/{aid}/download",
                "previewPath": f"/v1/artifacts/{aid}/preview",
                "hasVisualPreview": bool(value.get("has_visual_preview")),
                "previewMode": metadata.get("preview_mode"),
            })
            return found
        for item in value.values():
            _normalize_artifacts(item, found)
    elif isinstance(value, list):
        for item in value:
            _normalize_artifacts(item, found)
    return found


def _validation_of(value: Any) -> dict | None:
    if isinstance(value, dict):
        if isinstance(value.get("validation"), dict):
            return value["validation"]
        for item in value.values():
            found = _validation_of(item)
            if found is not None:
                return found
    elif isinstance(value, list):
        for item in value:
            found = _validation_of(item)
            if found is not None:
                return found
    return None


def _calc_summary(tool: str, result: Any) -> dict:
    validation = _validation_of(result)
    model = result.get("model") if isinstance(result, dict) else None
    return {"model": model or tool, "validation": validation}


def _truncate_result(result: Any) -> dict:
    encoded = json.dumps(result, ensure_ascii=False, allow_nan=False)
    if len(encoded) <= MAX_TOOL_RESULT:
        return {"result": result, "truncated": False}
    keys = list(result.keys())[:20] if isinstance(result, dict) else []
    return {
        "result": {k: result[k] for k in keys if k not in {"stages", "streams", "profiles"}},
        "truncated": True,
        "resultBytes": len(encoded),
        "note": "结果过大，已截断；完整结果可用 chem_calc/read 或面板查看。",
    }


# --------------------------------------------------------------------------- #
# request models
# --------------------------------------------------------------------------- #
class ProjectCreate(BaseModel):
    name: str = Field(min_length=1, max_length=120)


class RequirementItem(BaseModel):
    code: str = Field(min_length=1, max_length=30)
    title: str = Field(min_length=1, max_length=300)
    mandatory: bool = True
    applicability: str = ""
    quote: str = ""
    source: str = ""


class RequirementsIn(BaseModel):
    items: list[RequirementItem] = Field(min_length=1, max_length=200)


class PlanObject(BaseModel):
    key: str = Field(min_length=1, max_length=80)
    kind: str = "basis"
    title: str = ""
    note: str = ""


class PlanDeliverable(BaseModel):
    key: str = Field(min_length=1, max_length=80)
    title: str = Field(min_length=1, max_length=200)
    requirementCodes: list[str] = Field(default_factory=list, max_length=100)
    formats: list[str] = Field(default_factory=list, max_length=20)


class PlanIn(BaseModel):
    objects: list[PlanObject] = Field(default_factory=list, max_length=200)
    deliverables: list[PlanDeliverable] = Field(default_factory=list, max_length=200)


class ObjectIn(BaseModel):
    kind: str = Field(min_length=1, max_length=40)
    key: str = Field(min_length=1, max_length=80)
    title: str = ""
    quantities: list[dict] = Field(default_factory=list, max_length=200)
    attributes: dict = Field(default_factory=dict)
    dependsOn: list[str] = Field(default_factory=list, max_length=100)


class ExecuteIn(BaseModel):
    tool: str = Field(min_length=1, max_length=60)
    args: dict = Field(default_factory=dict)


class IntakeIn(BaseModel):
    path: str = Field(min_length=1, max_length=2000)
    name: str | None = Field(default=None, max_length=200)


class ReviewIn(BaseModel):
    note: str = ""


# --------------------------------------------------------------------------- #
# auth / health / catalog
# --------------------------------------------------------------------------- #
@app.middleware("http")
async def kernel_token(request: Request, call_next):
    if KERNEL_TOKEN and request.url.path.startswith("/v1/"):
        supplied = request.headers.get("x-kernel-token", "")
        if supplied != KERNEL_TOKEN:
            return JSONResponse({"detail": "内核令牌无效"}, status_code=401)
    return await call_next(request)


@app.get("/v1/health")
def health():
    return {
        "ok": True,
        "version": KERNEL_VERSION,
        "phaseGate": PHASE_GATE,
        "dataDir": str(DATA_ROOT),
        "projectCount": len(list(PROJECTS_DIR.glob("*/project.json"))),
    }


@app.get("/v1/tools")
def tools():
    return {
        "tools": TOOLS,
        "schemas": {name: schema.model_json_schema() for name, schema in SCHEMAS.items()},
        "executable": sorted(list(SCHEMAS.keys() - {"ask_user", "acceptance"}) + list(SPECIAL_TOOLS)),
        "calcTools": sorted(CALC_TOOLS),
        "generateTools": sorted(GENERATE_TOOLS),
        "evidenceTools": sorted(EVIDENCE_TOOLS),
    }


# --------------------------------------------------------------------------- #
# projects
# --------------------------------------------------------------------------- #
@app.get("/v1/projects")
def list_projects():
    result = []
    for path in sorted(PROJECTS_DIR.glob("*/project.json")):
        try:
            result.append(summary_of(json.loads(path.read_text(encoding="utf-8"))))
        except (OSError, ValueError):
            continue
    result.sort(key=lambda item: item.get("updatedAt") or "", reverse=True)
    return {"projects": result}


@app.post("/v1/projects")
def create_project(p: ProjectCreate):
    ident = uid()
    project = {
        "id": ident,
        "name": p.name,
        "createdAt": _utc(),
        "updatedAt": _utc(),
        "requirements": [],
        "plan": {"objects": [], "deliverables": []},
        "objects": {},
        "calculations": [],
        "files": [],
        "review": {},
        "bundle": None,
    }
    with _write_lock:
        project["phase"] = derive_phase(project)
        save_project(project)
        ensure_legacy_project(project)
    return {"ok": True, "project": summary_of(project)}


@app.get("/v1/projects/{project_id}")
def get_project(project_id: str):
    project = load_project(project_id)
    return {"ok": True, "project": project, "summary": summary_of(project)}


@app.delete("/v1/projects/{project_id}")
def delete_project(project_id: str):
    """Delete a local project: state, snapshot, calculation rows and artifact files."""
    project = load_project(project_id)
    artifact_rows = rows("SELECT path FROM artifacts WHERE project_id=?", (project_id,))
    with _write_lock:
        shutil.rmtree(project_dir(project_id), ignore_errors=True)
        shutil.rmtree(SNAPSHOT_DIR / project_id, ignore_errors=True)
        for row in artifact_rows:
            path = Path(row["path"]).resolve()
            if path.is_relative_to((KERNEL_DATA / "artifacts").resolve()):
                shutil.rmtree(path.parent, ignore_errors=True)
        with connection() as db:
            # Local project deletion spans the archived schema; remove every
            # project-scoped row and the conversation children, then the project.
            db.execute("PRAGMA foreign_keys=OFF")
            db.execute("DELETE FROM conversation_artifacts WHERE conversation_id IN "
                       "(SELECT id FROM conversations WHERE project_id=?)", (project_id,))
            db.execute("DELETE FROM conversation_calculations WHERE conversation_id IN "
                       "(SELECT id FROM conversations WHERE project_id=?)", (project_id,))
            db.execute("DELETE FROM conversation_history WHERE conversation_id IN "
                       "(SELECT id FROM conversations WHERE project_id=?)", (project_id,))
            db.execute("DELETE FROM conversation_messages WHERE conversation_id IN "
                       "(SELECT id FROM conversations WHERE project_id=?)", (project_id,))
            tables = [row["name"] for row in db.execute(
                "SELECT name FROM sqlite_master WHERE type='table'").fetchall()]
            for table in tables:
                if table == "projects":
                    continue
                columns = [column[1] for column in db.execute(
                    f"PRAGMA table_info({table})").fetchall()]
                if "project_id" in columns:
                    db.execute(f"DELETE FROM {table} WHERE project_id=?", (project_id,))
            db.execute("DELETE FROM projects WHERE id=?", (project_id,))
    return {"ok": True, "deleted": project_id, "name": project["name"]}



@app.post("/v1/projects/{project_id}/requirements")
def save_requirements(project_id: str, p: RequirementsIn):
    project = load_project(project_id)
    existing = {item["code"]: item for item in project.get("requirements") or []}
    for item in p.items:
        existing[item.code] = {**item.model_dump(), "updatedAt": _utc()}
    project["requirements"] = list(existing.values())
    project["review"] = {}
    with _write_lock:
        save_project(project)
    return {"ok": True, "phase": project["phase"], "requirements": project["requirements"]}


@app.post("/v1/projects/{project_id}/plan")
def save_plan(project_id: str, p: PlanIn):
    project = load_project(project_id)
    plan = project.setdefault("plan", {"objects": [], "deliverables": []})
    objects = {item["key"]: item for item in plan.get("objects") or []}
    for item in p.objects:
        objects[item.key] = item.model_dump()
    deliverables = {item["key"]: item for item in plan.get("deliverables") or []}
    for item in p.deliverables:
        deliverables[item.key] = item.model_dump()
    plan["objects"] = list(objects.values())
    plan["deliverables"] = list(deliverables.values())
    project["review"] = {}
    with _write_lock:
        save_project(project)
    return {"ok": True, "phase": project["phase"], "plan": plan}


@app.post("/v1/projects/{project_id}/objects")
def save_object(project_id: str, p: ObjectIn):
    project = load_project(project_id)
    objects = project.setdefault("objects", {})
    storage_key = p.kind + "/" + p.key
    previous = objects.get(storage_key) or {}
    revision = int(previous.get("revision") or 0) + 1
    objects[storage_key] = {
        "kind": p.kind,
        "key": p.key,
        "title": p.title or p.key,
        "quantities": p.quantities,
        "attributes": p.attributes,
        "dependsOn": p.dependsOn,
        "revision": revision,
        "updatedAt": _utc(),
    }
    project["review"] = {}
    with _write_lock:
        save_project(project)
    return {"ok": True, "phase": project["phase"], "object": objects[storage_key]}


# --------------------------------------------------------------------------- #
# execute archived tools
# --------------------------------------------------------------------------- #
def _gate_check(project: dict, tool: str) -> str | None:
    minimum = MIN_PHASE.get(tool)
    if not minimum:
        return None
    current = PHASES.index(project.get("phase") or "intake")
    required = PHASES.index(minimum)
    if current >= required:
        return None
    message = f"当前阶段为 {project.get('phase')}，工具 {tool} 需要至少 {minimum} 阶段"
    if PHASE_GATE == "enforce":
        raise HTTPException(status_code=409, detail=message)
    return message


# --------------------------------------------------------------------------- #
# generated engineering files: spreadsheet and 3D layout
# --------------------------------------------------------------------------- #
DEFAULT_LAYOUT = [
    {"tag": "TK-101", "shape": "cylinder", "x": -14000, "y": 0, "diameter": 8000, "height": 10000},
    {"tag": "E-101", "shape": "box", "x": -6000, "y": 0, "length": 4000, "width": 2000, "height": 2000},
    {"tag": "R-101", "shape": "cylinder", "x": 0, "y": 0, "diameter": 3000, "height": 12000},
    {"tag": "E-102", "shape": "box", "x": -6000, "y": 5000, "length": 4500, "width": 2000, "height": 2000},
    {"tag": "V-101", "shape": "box", "x": 5000, "y": 12000, "length": 6000, "width": 2500, "height": 2500},
    {"tag": "T-101", "shape": "cylinder", "x": 9000, "y": 0, "diameter": 2400, "height": 24000},
    {"tag": "TK-102", "shape": "cylinder", "x": 18000, "y": 0, "diameter": 8000, "height": 10000},
]


def _freecad_executable() -> Path | None:
    candidates: list[Path] = []
    found = shutil.which("freecadcmd") or shutil.which("FreeCADCmd")
    if found:
        candidates.append(Path(found))
    local = os.environ.get("LOCALAPPDATA")
    if local:
        candidates.extend(sorted(Path(local).glob("Programs/FreeCAD*/bin/freecadcmd.exe")))
    candidates.extend(sorted(Path(r"C:\Program Files").glob("FreeCAD*/bin/FreeCADCmd.exe")))
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    return None


def _freecad_script(items: list[dict], fcstd: Path, step: Path) -> str:
    body = [
        "import FreeCAD",
        "import Part",
        "doc = FreeCAD.newDocument('ChemELLM_Layout')",
        "items = " + json.dumps(items, ensure_ascii=False),
        "for item in items:",
        "    tag = str(item.get('tag', 'ITEM'))",
        "    if item.get('shape') == 'cylinder':",
        "        obj = doc.addObject('Part::Cylinder', tag)",
        "        obj.Radius = float(item.get('diameter', 2000)) / 2.0",
        "        obj.Height = float(item.get('height', 4000))",
        "    else:",
        "        obj = doc.addObject('Part::Box', tag)",
        "        obj.Length = float(item.get('length', 2000))",
        "        obj.Width = float(item.get('width', 1500))",
        "        obj.Height = float(item.get('height', 1500))",
        "    obj.Placement.Base = FreeCAD.Vector(float(item.get('x', 0)), float(item.get('y', 0)), float(item.get('z', 0)))",
        "doc.recompute()",
        "doc.saveAs(" + repr(fcstd.as_posix()) + ")",
        "Part.export(doc.Objects, " + repr(step.as_posix()) + ")",
    ]
    return "\n".join(body) + "\n"


def export_xlsx(project: dict) -> dict:
    from openpyxl import Workbook

    ident = uid()
    folder = KERNEL_DATA / "artifacts" / ident
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / "项目数据表.xlsx"
    workbook = Workbook()

    def sheet(title: str, header: list[str], rows: list[list]):
        target = workbook.create_sheet(title) if workbook.sheetnames != ["Sheet"] else workbook.active
        target.title = title
        target.append(header)
        for row in rows:
            target.append(row)
        return target

    sheet("概览", ["字段", "值"], [
        ["项目", project.get("name")],
        ["项目ID", project.get("id")],
        ["阶段", project.get("phase")],
        ["更新时间", project.get("updatedAt")],
        ["要求数", len(project.get("requirements") or [])],
        ["成果计划数", len((project.get("plan") or {}).get("deliverables") or [])],
        ["工程对象数", len(project.get("objects") or {})],
        ["计算记录数", len(project.get("calculations") or [])],
        ["文件数", len(project.get("files") or [])],
        ["自动校核", "ready" if (project.get("review") or {}).get("ready") else "存在缺口"],
    ])
    sheet("要求清单", ["编码", "要求", "必需", "适用性", "原文引句", "来源"],
          [[r.get("code"), r.get("title"), "是" if r.get("mandatory", True) else "否",
            r.get("applicability"), r.get("quote"), r.get("source")]
           for r in (project.get("requirements") or [])])
    sheet("成果计划", ["key", "标题", "要求编码", "格式"],
          [[d.get("key"), d.get("title"), "、".join(d.get("requirementCodes") or []),
            "、".join(d.get("formats") or [])]
           for d in ((project.get("plan") or {}).get("deliverables") or [])])
    object_rows = []
    for obj in (project.get("objects") or {}).values():
        for quantity in obj.get("quantities") or []:
            object_rows.append([obj.get("kind"), obj.get("key"), obj.get("title"),
                                quantity.get("key"), quantity.get("value"), quantity.get("unit"),
                                quantity.get("origin"), quantity.get("conditions"),
                                quantity.get("rationale"), quantity.get("calculationId")])
    sheet("设计依据", ["类型", "对象", "标题", "数值项", "值", "单位", "来源类型", "条件", "理由", "计算ID"],
          object_rows)
    sheet("计算记录", ["工具", "计算ID", "校验", "模型"],
          [[c.get("tool"), c.get("calculationId"),
            "通过" if (c.get("validation") or {}).get("passed") is True else
            "未通过" if (c.get("validation") or {}).get("passed") is False else "无校验",
            c.get("model")]
           for c in (project.get("calculations") or [])])
    sheet("成果文件", ["名称", "类型", "来源", "字节", "SHA256"],
          [[f.get("name"), f.get("kind"), f.get("origin"), f.get("bytes"), f.get("sha256")]
           for f in (project.get("files") or [])])
    workbook.save(path)
    record = artifact_record(project["id"], path, name=path.name, kind="xlsx",
                             metadata={"source": "kernel export_xlsx", "status": "draft"})
    return {"artifacts": [record], "status": "draft", "sheets": workbook.sheetnames}


def _layout_preview_svg(items: list[dict], title: str = "三维设备布置 · 等轴测投影（示意）") -> str:
    """Headless FreeCAD 没有 GUI 视图，这里用几何参数生成等轴测 SVG 预览。"""
    import html
    import math

    cos30 = math.cos(math.radians(30.0))
    sin30 = math.sin(math.radians(30.0))

    def item_corners(item: dict):
        x = float(item.get("x", 0.0))
        y = float(item.get("y", 0.0))
        z = float(item.get("z", 0.0))
        if item.get("shape") == "cylinder":
            radius = float(item.get("diameter", 2000.0)) / 2.0
            height = float(item.get("height", 4000.0))
            points = []
            for index in range(24):
                angle = 2.0 * math.pi * index / 24.0
                px = x + radius * math.cos(angle)
                py = y + radius * math.sin(angle)
                points.append((px, py, z))
                points.append((px, py, z + height))
            return points
        length = float(item.get("length", 2000.0))
        width = float(item.get("width", 1500.0))
        height = float(item.get("height", 1500.0))
        return [
            (x + dx, y + dy, z + dz)
            for dx in (0.0, length) for dy in (0.0, width) for dz in (0.0, height)
        ]

    projected = []
    for item in items:
        points = [
            ((px - py) * cos30, (px + py) * sin30 - pz)
            for px, py, pz in item_corners(item)
        ]
        projected.append((item, points))

    flat = [point for _, points in projected for point in points]
    minx = min(point[0] for point in flat)
    maxx = max(point[0] for point in flat)
    miny = min(point[1] for point in flat)
    maxy = max(point[1] for point in flat)
    width, height, margin = 1200, 760, 70
    spanx = max(maxx - minx, 1e-6)
    spany = max(maxy - miny, 1e-6)
    scale = min((width - 2 * margin) / spanx, (height - 2 * margin) / spany)
    ox = width / 2.0 - (minx + maxx) / 2.0 * scale
    oy = height / 2.0 + (miny + maxy) / 2.0 * scale

    def screen(point):
        return (ox + point[0] * scale, oy - point[1] * scale)

    def hull(points):
        unique = sorted(set((round(point[0], 6), round(point[1], 6)) for point in points))
        if len(unique) <= 2:
            return unique

        def cross(origin, a, b):
            return (a[0] - origin[0]) * (b[1] - origin[1]) - (a[1] - origin[1]) * (b[0] - origin[0])

        lower = []
        for point in unique:
            while len(lower) >= 2 and cross(lower[-2], lower[-1], point) <= 0:
                lower.pop()
            lower.append(point)
        upper = []
        for point in reversed(unique):
            while len(upper) >= 2 and cross(upper[-2], upper[-1], point) <= 0:
                upper.pop()
            upper.append(point)
        return lower[:-1] + upper[:-1]

    palette = ["#cfe3f5", "#d6ecd2", "#f8e3c1", "#e7d8f0", "#f5d6d6", "#d9f0ee", "#e6e6e6"]
    shapes = []
    for index, (item, points) in enumerate(projected):
        hull_points = [screen(point) for point in hull(points)]
        if len(hull_points) < 3:
            continue
        polygon = " ".join("%.1f,%.1f" % (px, py) for px, py in hull_points)
        shapes.append(
            '<polygon points="%s" fill="%s" stroke="#37506b" stroke-width="1.5"/>'
            % (polygon, palette[index % len(palette)])
        )
        cx = sum(point[0] for point in hull_points) / len(hull_points)
        cy = sum(point[1] for point in hull_points) / len(hull_points)
        label = html.escape(str(item.get("tag", "")))
        shapes.append(
            '<text x="%.1f" y="%.1f" font-size="13" text-anchor="middle" fill="#1b2a3a">%s</text>'
            % (cx, cy, label)
        )

    header = [
        '<rect width="%d" height="%d" fill="#ffffff"/>' % (width, height),
        '<text x="24" y="34" font-size="20" fill="#1b2a3a">%s</text>' % html.escape(title),
        '<text x="24" y="58" font-size="12" fill="#5a6b7d">'
        "示意投影（单位 mm，非施工图）；由参数化布置生成，未做碰撞、检修与安全间距校核。</text>",
        '<text x="24" y="76" font-size="12" fill="#5a6b7d">设备数量：%d；FCStd/STEP 请下载后用 FreeCAD 打开。</text>' % len(projected),
    ]
    return (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<svg xmlns="http://www.w3.org/2000/svg" width="%d" height="%d" viewBox="0 0 %d %d">\n'
        % (width, height, width, height)
        + "\n".join(header + shapes)
        + "\n</svg>\n"
    )


def export_freecad(project: dict, args: dict) -> dict:
    executable = _freecad_executable()
    if not executable:
        raise HTTPException(status_code=422, detail="未找到 FreeCADCmd；请安装 FreeCAD 并确认路径")
    items = args.get("items") or DEFAULT_LAYOUT
    if not isinstance(items, list) or not items:
        raise HTTPException(status_code=422, detail="items 必须是非空数组")
    ident = uid()
    folder = KERNEL_DATA / "artifacts" / ident
    folder.mkdir(parents=True, exist_ok=True)
    script = folder / "layout.py"
    fcstd = folder / "layout.FCStd"
    step = folder / "layout.step"
    preview = folder / "layout-preview.svg"
    script.write_text(_freecad_script(items, fcstd, step), encoding="utf-8")
    try:
        completed = subprocess.run([str(executable), str(script)], capture_output=True, timeout=300)
    except (OSError, subprocess.SubprocessError) as exc:
        raise HTTPException(status_code=500, detail="FreeCAD 运行失败：" + str(exc)[:500]) from exc
    if completed.returncode != 0:
        detail = completed.stderr.decode("utf-8", "replace")[-800:]
        raise HTTPException(status_code=500, detail="FreeCAD 返回错误：" + detail)
    preview.write_text(_layout_preview_svg(items), encoding="utf-8")
    records = []
    if fcstd.is_file():
        records.append(artifact_record(
            project["id"], fcstd, name="三维设备布置.FCStd", kind="fcstd", preview=str(preview),
            metadata={"source": "kernel export_freecad", "items": len(items), "preview_mode": "visual"},
        ))
    if step.is_file():
        records.append(artifact_record(
            project["id"], step, name="三维设备布置.step", kind="step", preview=str(preview),
            metadata={"source": "kernel export_freecad", "items": len(items), "preview_mode": "visual"},
        ))
    if preview.is_file():
        records.append(artifact_record(
            project["id"], preview, name="三维布置预览.svg", kind="svg",
            metadata={"source": "kernel export_freecad", "items": len(items)},
        ))
    if not records:
        raise HTTPException(status_code=500, detail="FreeCAD 未生成 FCStd/STEP 文件")
    return {
        "artifacts": records,
        "cad": executable.name,
        "versionHint": "FreeCAD 1.x",
        "items": len(items),
        "warnings": [
            "三维布置为参数化简化模型，未做碰撞、检修空间与安全间距校核。",
            "若在 FreeCAD 中打开后视图空白，请执行菜单“视图 → 全部适应（Fit all）”：模型跨度约 40 m，默认视图可能落在设备内部。",
        ],
    }



def run_dwsim(project: dict, args: dict) -> dict:
    payload = dict(args or {})
    action = str(payload.get("action") or "status").lower()
    if action == "status":
        return {
            "status": dwsim_bridge.status(),
            "warnings": [
                "DWSIM 是外部 GPL 软件，不随插件分发；此处只做本机检测与调用。",
            ],
        }
    if action != "run":
        raise HTTPException(status_code=422, detail="action 必须是 status 或 run")
    raw_path = str(payload.get("flowsheet_path") or "").strip()
    if not raw_path:
        raise HTTPException(status_code=422, detail="run 需要 flowsheet_path")
    path = Path(raw_path)
    if not path.is_absolute():
        path = (ROOT / path).resolve()
    temp = KERNEL_DATA / "dwsim-temp"
    try:
        result = dwsim_bridge.run_flowsheet(path, temp)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)[:1200]) from exc
    ident = uid()
    folder = KERNEL_DATA / "artifacts" / ident
    folder.mkdir(parents=True, exist_ok=True)
    json_path = folder / "DWSIM结果.json"
    csv_path = folder / "DWSIM物流表.csv"
    json_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    csv_path.write_text(dwsim_bridge.streams_csv(result), encoding="utf-8-sig")
    records = [
        artifact_record(project["id"], json_path, name="DWSIM结果.json", kind="json",
                        metadata={"source": "kernel dwsim", "flowsheet": str(path),
                                  "solved": bool(result.get("solved"))}),
        artifact_record(project["id"], csv_path, name="DWSIM物流表.csv", kind="csv",
                        metadata={"source": "kernel dwsim", "flowsheet": str(path)}),
    ]
    warnings = [
        "DWSIM 为外部 GPL 软件，不随插件分发；结果取决于 flowsheet 内的物性包、单元模型与收敛设置。",
        "DWSIM 输出未做项目专业审核，不能替代设计校核或厂家数据。",
        "温度单位 K、压力单位 Pa、质量流量单位 kg/s；组成来自 DWSIM 物流（摩尔分数）。",
    ]
    if not result.get("solved"):
        warnings.insert(0, "DWSIM 未收敛（solved=false），物流结果不可用；请检查 flowsheet 与收敛设置。")
    if not result.get("streams"):
        warnings.append("该 flowsheet 未返回物料流；可能是纯能量/动态/自定义单元流程。")
    return {
        "artifacts": records,
        "summary": {
            "engine": result.get("engine"), "version": result.get("version"),
            "flowsheet": str(path), "solved": bool(result.get("solved")),
            "stream_count": len(result.get("streams") or []),
            "unit_count": len(result.get("unit_operations") or []),
            "error": str(result.get("error") or ""),
        },
        "warnings": warnings,
    }


def export_compliance_matrix(project: dict, args: dict) -> dict:
    payload = dict(args or {})
    ident = uid()
    folder = KERNEL_DATA / "artifacts" / ident
    folder.mkdir(parents=True, exist_ok=True)
    try:
        built = compliance.build_compliance_files(payload, folder)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)[:1200]) from exc
    records = [
        artifact_record(project["id"], built["xlsx"], name="标准符合性矩阵.xlsx", kind="xlsx",
                        metadata={"source": "kernel compliance_matrix"}),
        artifact_record(project["id"], built["docx"], name="标准符合性矩阵.docx", kind="docx",
                        metadata={"source": "kernel compliance_matrix"}),
        artifact_record(project["id"], built["json"], name="标准符合性矩阵.json", kind="json",
                        metadata={"source": "kernel compliance_matrix"}),
    ]
    data = built["data"]
    return {
        "artifacts": records,
        "summary": {"rows": len(data["rows"]), "counts": data["counts"],
                    "by_standard": data["by_standard"]},
        "warnings": [
            "“待核对”条款未取得标准原文，未填写具体限值，需取得原文后人工核对。",
            "“符合”仅表示现有项目数据满足该数值/完成状态，不代表设计文件已通过专业审核。",
            "涉及排放、QRA、重大危险源与施工图深度的条款需环评/安全评价/设计审查确认。",
        ],
    }


def export_hazop_lopa(project: dict, args: dict) -> dict:
    payload = dict(args or {})
    ident = uid()
    folder = KERNEL_DATA / "artifacts" / ident
    folder.mkdir(parents=True, exist_ok=True)
    try:
        built = hazop.build_hazop_files(payload, folder)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)[:1200]) from exc
    records = [
        artifact_record(project["id"], built["xlsx"], name="HAZOP_LOPA工作表.xlsx", kind="xlsx",
                        metadata={"source": "kernel hazop_lopa"}),
        artifact_record(project["id"], built["docx"], name="HAZOP_LOPA工作表.docx", kind="docx",
                        metadata={"source": "kernel hazop_lopa"}),
        artifact_record(project["id"], built["json"], name="HAZOP_LOPA数据.json", kind="json",
                        metadata={"source": "kernel hazop_lopa"}),
    ]
    data = built["data"]
    return {
        "artifacts": records,
        "summary": data["summary"],
        "warnings": list(data["warnings"]) + [
            "HAZOP 为会议驱动方法；本工作表不替代 HAZOP 会议与签署。",
            "LOPA 的初始频率/PFD/目标频率需企业风险标准与项目数据确认；SIL 为初判，未做 SIL 验证。",
        ],
    }


def export_relief_sizing(project: dict, args: dict) -> dict:
    payload = dict(args or {})
    ident = uid()
    folder = KERNEL_DATA / "artifacts" / ident
    folder.mkdir(parents=True, exist_ok=True)
    try:
        built = relief.build_relief_files(payload, folder)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)[:1200]) from exc
    records = [
        artifact_record(project["id"], built["docx"], name="安全阀选型计算书.docx", kind="docx",
                        metadata={"source": "kernel relief_sizing"}),
        artifact_record(project["id"], built["pdf"], name="安全阀选型计算书.pdf", kind="pdf",
                        metadata={"source": "kernel relief_sizing"}),
        artifact_record(project["id"], built["json"], name="安全阀选型计算书.json", kind="json",
                        metadata={"source": "kernel relief_sizing"}),
    ]
    result = built["result"]
    return {
        "artifacts": records,
        "result": {
            "device_type": result["device_type"], "scenario": result["scenario"],
            "regime": result.get("regime"), "required_area_mm2": result.get("required_area_mm2"),
            "selected_orifice": result.get("selected_orifice"),
            "selected_area_mm2": result.get("selected_area_mm2"),
            "capacity_utilization": result.get("capacity_utilization"),
            "residual": result.get("residual"), "passed": result["passed"],
            "checks": result["checks"],
        },
        "warnings": result["warnings"] + [
            "API 520/521/526 原文未取得，公式与限值按公开资料实现，需按正式标准/项目规定核对。",
            "标准孔径为 API 526 常用值，实际以厂家样本与铭牌为准。",
        ],
    }


def export_equipment_datasheet(project: dict, args: dict) -> dict:
    payload = dict(args or {})
    ident = uid()
    folder = KERNEL_DATA / "artifacts" / ident
    folder.mkdir(parents=True, exist_ok=True)
    try:
        built = datasheet.build_datasheet_files(payload, folder)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)[:1200]) from exc
    records = [
        artifact_record(project["id"], built["docx"], name="设备数据表.docx", kind="docx",
                        metadata={"source": "kernel equipment_datasheet"}),
        artifact_record(project["id"], built["pdf"], name="设备数据表.pdf", kind="pdf",
                        metadata={"source": "kernel equipment_datasheet"}),
        artifact_record(project["id"], built["json"], name="设备数据表.json", kind="json",
                        metadata={"source": "kernel equipment_datasheet"}),
    ]
    data = built["data"]
    return {
        "artifacts": records,
        "items": [{"tag": item["tag"], "type": item["equipment_type"],
                   "thickness": item["thickness"], "volume": item["volume"]} for item in data["items"]],
        "warnings": list(data["warnings"]) + [
            "壁厚/重量为简化估算，未计开孔补强、内件与制造公差。",
            "“待定/待厂家确认”项目须由设计、工艺与厂家共同确认后才能询价。",
        ],
    }


def export_pid(project: dict, args: dict) -> dict:
    model = dict(args or {})
    validation = pid_standard.validate_standard(model)
    if not validation["passed"]:
        raise HTTPException(status_code=422, detail="P&ID 校验未通过：" + "；".join(validation["blockers"][:6]))
    ident = uid()
    folder = KERNEL_DATA / "artifacts" / ident
    built = pid_standard.build_pid_files(model, folder)
    spec = {
        "svg": ("P&ID.svg", "svg"),
        "pdf": ("P&ID.pdf", "pdf"),
        "dxf": ("P&ID.dxf", "dxf"),
        "json": ("P&ID.json", "json"),
        "dexpi": ("P&ID.dexpi.xml", "xml"),
    }
    records = []
    for key, (name, kind) in spec.items():
        records.append(artifact_record(
            project["id"], built[key], name=name, kind=kind,
            metadata={"source": "kernel pid", "drawing_type": model.get("drawing_type", "P&ID"),
                      "standard": validation.get("standard")},
        ))
    return {
        "artifacts": records,
        "validation": validation,
        "counts": validation["counts"],
        "warnings": list(validation["warnings"]) + [
            "本图为参照 HG/T 20519 内容深度与 ISA-5.1 符号的规范化示意图，条款级合规与施工图深度仍待核对。",
            "DEXPI 文件为子集，非认证 DEXPI 2.0 文档。",
        ],
    }


def export_pid_data(project: dict, args: dict) -> dict:
    model = dict(args or {})
    ident = uid()
    folder = KERNEL_DATA / "artifacts" / ident
    folder.mkdir(parents=True, exist_ok=True)
    try:
        built = pid_data.build_pid_data_files(model, folder)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)[:1200]) from exc
    records = [
        artifact_record(project["id"], built["xlsx"], name="P&ID数据表.xlsx", kind="xlsx",
                        metadata={"source": "kernel pid_data"}),
        artifact_record(project["id"], built["docx"], name="P&ID数据表.docx", kind="docx",
                        metadata={"source": "kernel pid_data"}),
        artifact_record(project["id"], built["json"], name="P&ID数据.json", kind="json",
                        metadata={"source": "kernel pid_data"}),
    ]
    names = {
        "lines": ("管线表.csv", "管线表"),
        "mto": ("材料表.csv", "材料表"),
        "instruments": ("仪表索引.csv", "仪表索引"),
        "io": ("IO清单.csv", "IO清单"),
        "ce": ("联锁因果表.csv", "联锁因果表"),
    }
    for key, (filename, label) in names.items():
        records.append(artifact_record(
            project["id"], built["csv"][key], name=filename, kind="csv",
            metadata={"source": "kernel pid_data", "table": label},
        ))
    return {
        "artifacts": records,
        "summary": built["data"]["summary"],
        "validation": built["validation"],
        "warnings": list(built["validation"]["warnings"]) + [
            "MTO 中的弯头/三通/法兰为拓扑估算，未按实际走向与管口计算，需配管专业复核。",
            "I/O 信号类型与去向为按位号推断的“待定”建议，需自控专业确认。",
            "数据表由结构化模型生成，不替代设计文件与采购技术协议。",
        ],
    }


def check_pid(project: dict, args: dict) -> dict:
    model = dict(args or {})
    validation = pid_standard.validate_standard(model)
    return {"validation": validation, "model": "pid-rules-standard", "warnings": validation["warnings"]}


def export_cad_part(project: dict, args: dict) -> dict:
    executable = _freecad_executable()
    if not executable:
        raise HTTPException(status_code=422, detail="未找到 FreeCADCmd；请安装 FreeCAD 并确认路径")
    ident = uid()
    folder = KERNEL_DATA / "artifacts" / ident
    folder.mkdir(parents=True, exist_ok=True)
    kind = str((args or {}).get("equipment_type") or "vessel").lower()
    if kind == "gear":
        try:
            built = gear.build_gear_files(dict(args or {}), folder, executable)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail=str(exc)[:800]) from exc
        except RuntimeError as exc:
            raise HTTPException(status_code=500, detail=str(exc)[:800]) from exc
        params = built["params"]
        calc = built["calculation"]
        drawing_svg = built["drawing_svg"]
        records = [
            artifact_record(project["id"], built["fcstd"], name="齿轮模型.FCStd", kind="fcstd",
                            preview=str(drawing_svg),
                            metadata={"source": "kernel cad_part", "equipment_type": "gear",
                                      "tag": params["tag"], "preview_mode": "visual"}),
            artifact_record(project["id"], built["step"], name="齿轮模型.step", kind="step",
                            preview=str(drawing_svg),
                            metadata={"source": "kernel cad_part", "equipment_type": "gear",
                                      "tag": params["tag"], "preview_mode": "visual"}),
            artifact_record(project["id"], drawing_svg, name="齿轮图.svg", kind="svg",
                            metadata={"source": "kernel cad_part", "tag": params["tag"]}),
            artifact_record(project["id"], built["drawing_pdf"], name="齿轮图.pdf", kind="pdf",
                            metadata={"source": "kernel cad_part", "tag": params["tag"]}),
            artifact_record(project["id"], built["dxf"], name="齿轮图.dxf", kind="dxf",
                            metadata={"source": "kernel cad_part", "tag": params["tag"]}),
            artifact_record(project["id"], built["bom"], name="齿轮设计数据.json", kind="json",
                            metadata={"source": "kernel cad_part", "tag": params["tag"]}),
        ]
        return {
            "artifacts": records,
            "params": {key: value for key, value in params.items() if key != "points"},
            "calculation": calc,
            "cad": executable.name,
            "warnings": [
                "标准直齿圆柱齿轮：基本齿廓 α=20°（GB/T 1356，原文未取得，待核对）；无变位、无螺旋角。",
                "齿廓为渐开线（基圆→齿顶），齿根过渡为简化径向/圆弧近似，未计齿根圆角与修形。",
                "精度等级按 GB/T 10095（待核对）；本图不含公差、材料热处理与检验要求，不替代制造图。",
                "重量按齿廓面积×齿宽×7850 kg/m3 估算，不含轮毂/倒角。",
                "若在 FreeCAD 中打开后视图空白，请执行“视图 → 全部适应（Fit all）”。",
            ],
        }
    try:
        built = cad_engineering.build_cad_files(dict(args or {}), folder, executable)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)[:800]) from exc
    except RuntimeError as exc:
        raise HTTPException(status_code=500, detail=str(exc)[:800]) from exc
    params = built["params"]
    calc = built["calculation"]
    drawing_svg = built["drawing_svg"]
    records = [
        artifact_record(
            project["id"], built["fcstd"], name="设备模型.FCStd", kind="fcstd", preview=str(drawing_svg),
            metadata={"source": "kernel cad_part", "equipment_type": params["equipment_type"],
                      "tag": params["tag"], "preview_mode": "visual"},
        ),
        artifact_record(
            project["id"], built["step"], name="设备模型.step", kind="step", preview=str(drawing_svg),
            metadata={"source": "kernel cad_part", "equipment_type": params["equipment_type"],
                      "tag": params["tag"], "preview_mode": "visual"},
        ),
        artifact_record(
            project["id"], drawing_svg, name="设备图.svg", kind="svg",
            metadata={"source": "kernel cad_part", "tag": params["tag"]},
        ),
        artifact_record(
            project["id"], built["drawing_pdf"], name="设备图.pdf", kind="pdf",
            metadata={"source": "kernel cad_part", "tag": params["tag"]},
        ),
        artifact_record(
            project["id"], built["bom"], name="设备设计数据.json", kind="json",
            metadata={"source": "kernel cad_part", "tag": params["tag"]},
        ),
    ]
    return {
        "artifacts": records,
        "params": params,
        "calculation": calc,
        "cad": executable.name,
        "warnings": [
            "实体为名义外形（添加了接管法兰，未抽壳）；壁厚按 GB/T 150.1 简化式计算并在设备图中标注。",
            "设备图不含焊缝详图、开孔补强计算与制造公差，制造前须按正式标准与制造图复核。",
            "设计温度 > 400 ℃ 或外压/真空工况时，本简化式未计蠕变、持久强度与外压稳定性，须另行校核。",
            "若在 FreeCAD 中打开后视图空白，请执行“视图 → 全部适应（Fit all）”。",
        ],
    }


@app.post("/v1/projects/{project_id}/execute")
def execute_tool(project_id: str, p: ExecuteIn):
    project = load_project(project_id)
    if p.tool in {"ask_user", "acceptance"}:
        raise HTTPException(status_code=422, detail="该工具不在此入口执行")
    if p.tool not in SCHEMAS and p.tool not in SPECIAL_TOOLS:
        raise HTTPException(status_code=422, detail="未知工具：" + p.tool)
    warning = _gate_check(project, p.tool)
    ensure_legacy_project(project)
    if p.tool == "export_xlsx":
        result = export_xlsx(project)
    elif p.tool == "export_freecad":
        result = export_freecad(project, p.args or {})
    elif p.tool == "pid":
        result = export_pid(project, p.args or {})
    elif p.tool == "pid_check":
        result = check_pid(project, p.args or {})
    elif p.tool == "pid_data":
        result = export_pid_data(project, p.args or {})
    elif p.tool == "equipment_datasheet":
        result = export_equipment_datasheet(project, p.args or {})
    elif p.tool == "relief_sizing":
        result = export_relief_sizing(project, p.args or {})
    elif p.tool == "hazop_lopa":
        result = export_hazop_lopa(project, p.args or {})
    elif p.tool == "compliance_matrix":
        result = export_compliance_matrix(project, p.args or {})
    elif p.tool == "dwsim":
        result = run_dwsim(project, p.args or {})
    elif p.tool == "cad_part":
        result = export_cad_part(project, p.args or {})
    else:
        try:
            result = run_tool(project_id, p.tool, p.args)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail=str(exc)[:2000]) from exc
        except RuntimeError as exc:  # UncertainToolError is a RuntimeError subclass
            raise HTTPException(status_code=409, detail=str(exc)[:2000]) from exc

    artifacts = _normalize_artifacts(result, [])
    calculation_id = result.get("calculation_id")
    verification = _calc_summary(p.tool, result)
    record = {
        "calculationId": calculation_id,
        "tool": p.tool,
        "createdAt": _utc(),
        "args": p.args,
        **verification,
    }
    project.setdefault("calculations", []).append(record)
    for artifact in artifacts:
        project.setdefault("files", []).append({
            **artifact,
            "origin": "generated",
            "tool": p.tool,
            "createdAt": _utc(),
        })
    if artifacts or p.tool in CALC_TOOLS:
        project["review"] = {}
    with _write_lock:
        save_project(project)

    payload = _truncate_result(result)
    response = {
        "ok": True,
        "phase": project["phase"],
        "calculationId": calculation_id,
        "tool": p.tool,
        "verification": verification,
        "artifacts": artifacts,
        "warnings": [warning] if warning else [],
        **payload,
    }
    return response


@app.get("/v1/projects/{project_id}/calculations")
def list_calculations(project_id: str):
    project = load_project(project_id)
    return {"ok": True, "calculations": project.get("calculations") or []}


@app.get("/v1/projects/{project_id}/calculations/{calculation_id}")
def read_calculation(project_id: str, calculation_id: str):
    load_project(project_id)
    row = one("SELECT * FROM calculations WHERE id=? AND project_id=?", (calculation_id, project_id))
    if not row:
        raise HTTPException(status_code=404, detail="计算记录不存在")
    return {
        "ok": True,
        "calculationId": calculation_id,
        "tool": row["tool"],
        "input": json.loads(row["input"]),
        "result": json.loads(row["result"]),
    }


# --------------------------------------------------------------------------- #
# intake / files
# --------------------------------------------------------------------------- #
@app.post("/v1/projects/{project_id}/intake")
def intake(project_id: str, p: IntakeIn):
    project = load_project(project_id)
    source = Path(p.path)
    if not source.is_absolute():
        source = (ROOT / source).resolve()
    else:
        source = source.resolve()
    allowed = [ROOT.resolve(), DATA_ROOT.resolve()]
    if not any(source.is_relative_to(root) for root in allowed):
        raise HTTPException(status_code=422, detail="只接受工作区内的文件路径")
    if not source.is_file():
        raise HTTPException(status_code=404, detail="文件不存在：" + str(source))
    if source.stat().st_size > MAX_INTAKE:
        raise HTTPException(status_code=413, detail="文件超过 25 MiB")
    content = source.read_bytes()
    name = p.name or source.name
    ensure_legacy_project(project)
    try:
        record = old_ingest(project_id, name, content)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    files = project.setdefault("files", [])
    text_row = one("SELECT length(text) AS chars FROM artifacts WHERE id=?", (record["id"],))
    files.append({
        "artifactId": record["id"],
        "name": record["name"],
        "kind": record["kind"],
        "bytes": (record.get("metadata") or {}).get("bytes"),
        "sha256": (record.get("metadata") or {}).get("sha256"),
        "downloadPath": f"/v1/artifacts/{record['id']}/download",
        "previewPath": f"/v1/artifacts/{record['id']}/preview",
        "origin": "intake",
        "sourcePath": str(source),
        "createdAt": _utc(),
    })
    project["review"] = {}
    with _write_lock:
        save_project(project)
    return {
        "ok": True,
        "phase": project["phase"],
        "artifact": files[-1],
        "pages": len((record.get("metadata") or {}).get("pages") or []),
        "textChars": (text_row or {}).get("chars") or 0,
        "warnings": (record.get("metadata") or {}).get("warnings") or [],
    }


@app.get("/v1/artifacts/{artifact_id}/download")
def artifact_download(artifact_id: str):
    path = _artifact_path(artifact_id)
    return FileResponse(path, filename=path.name)


@app.get("/v1/artifacts/{artifact_id}/preview")
def artifact_preview(artifact_id: str):
    row = one("SELECT * FROM artifacts WHERE id=?", (artifact_id,))
    if not row:
        raise HTTPException(status_code=404, detail="文件不存在")
    preview = row["preview_path"]
    if preview and Path(preview).is_file():
        target = Path(preview).resolve()
        if target.is_relative_to(KERNEL_DATA.resolve()):
            return FileResponse(target, filename=target.name)
    return artifact_download(artifact_id)


# --------------------------------------------------------------------------- #
# review / bundle
# --------------------------------------------------------------------------- #
@app.post("/v1/projects/{project_id}/review")
def review(project_id: str, p: ReviewIn):
    project = load_project(project_id)
    blockers: list[str] = []
    requirements = project.get("requirements") or []
    plan = project.get("plan") or {}
    deliverables = plan.get("deliverables") or []
    objects = [o for o in (project.get("objects") or {}).values() if o.get("quantities")]
    calculations = project.get("calculations") or []
    files = [f for f in (project.get("files") or []) if f.get("origin") == "generated"]
    if not requirements:
        blockers.append("缺少要求清单（chem_requirements）")
    if not deliverables:
        blockers.append("缺少成果计划（chem_plan）")
    else:
        planned_codes = {code for item in deliverables for code in (item.get("requirementCodes") or [])}
        missing = [r["code"] for r in requirements if r.get("mandatory", True) and r["code"] not in planned_codes]
        if missing:
            blockers.append("要求未映射到成果：" + "、".join(missing[:10]))
    if not objects:
        blockers.append("缺少带数值的设计依据对象（chem_object）")
    else:
        incomplete = []
        for obj in objects:
            for quantity in obj.get("quantities") or []:
                missing_fields = [field for field in ("key", "unit", "origin", "conditions", "rationale")
                                  if not str(quantity.get(field) or "").strip()]
                if missing_fields:
                    incomplete.append(f"{obj.get('key')}/{quantity.get('key') or '?'}")
        if incomplete:
            blockers.append("依据数值缺少单位/条件/来源/理由：" + "、".join(incomplete[:8]))
    generated_kinds = {str(item.get("kind") or "").lower() for item in files}
    missing_deliverables = []
    for deliverable in deliverables:
        formats = {str(fmt).lower() for fmt in (deliverable.get("formats") or [])}
        if formats and not (formats & generated_kinds):
            missing_deliverables.append(str(deliverable.get("key")))
        elif not formats and not files:
            missing_deliverables.append(str(deliverable.get("key")))
    if missing_deliverables:
        blockers.append("成果计划缺少对应文件：" + "、".join(missing_deliverables[:10]))
    failed = [c for c in calculations
              if c.get("tool") in CALC_TOOLS and (c.get("validation") or {}).get("passed") is False]
    if failed:
        blockers.append("计算校验未通过：" + "、".join(str(c.get("tool")) for c in failed[:8]))
    if not calculations:
        blockers.append("尚无计算记录（chem_calc）")
    if not files:
        blockers.append("尚无成果文件（chem_generate）")
    unverified = [c for c in calculations
                  if c.get("tool") in CALC_TOOLS and "passed" not in (c.get("validation") or {})]
    result_warnings = []
    if unverified:
        result_warnings.append("以下计算没有显式校验结果，需人工核对：" +
                               "、".join(str(c.get("tool")) for c in unverified[:8]))
    result = {
        "ready": not blockers,
        "blockers": blockers,
        "warnings": result_warnings,
        "checkedAt": _utc(),
        "note": (p.note or "")[:2000],
        "counts": {
            "requirements": len(requirements),
            "deliverables": len(deliverables),
            "objects": len(objects),
            "calculations": len(calculations),
            "generatedFiles": len(files),
        },
        "professionalReview": "not_reviewed",
    }
    project["review"] = result
    with _write_lock:
        save_project(project)
    return {"ok": True, "phase": project["phase"], "review": result, "summary": summary_of(project)}


@app.post("/v1/projects/{project_id}/bundle")
def bundle(project_id: str):
    project = load_project(project_id)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    safe_name = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in project["name"])[:60] or "chemellm"
    filename = f"{safe_name}-{stamp}.zip"
    target = BUNDLES_DIR / filename
    manifest = {
        "kernelVersion": KERNEL_VERSION,
        "project": summary_of(project),
        "generatedAt": _utc(),
        "files": [],
    }
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("project.json", json.dumps(project, ensure_ascii=False, indent=2))
        archive.writestr("requirements.json", json.dumps(project.get("requirements") or [], ensure_ascii=False, indent=2))
        archive.writestr("plan.json", json.dumps(project.get("plan") or {}, ensure_ascii=False, indent=2))
        archive.writestr("calculations.json", json.dumps(project.get("calculations") or [], ensure_ascii=False, indent=2))
        for artifact in project.get("files") or []:
            aid = artifact.get("artifactId")
            if not aid:
                continue
            try:
                path = _artifact_path(aid)
            except HTTPException:
                manifest["files"].append({"artifactId": aid, "name": artifact.get("name"), "missing": True})
                continue
            if path.stat().st_size > MAX_INTAKE:
                manifest["files"].append({"artifactId": aid, "name": artifact.get("name"), "skipped": "too_large"})
                continue
            # Artifact names can repeat (multiple P&ID exports, several CAD parts);
            # prefix with the artifact id so every archive entry stays unique.
            arcname = "files/" + aid[:8] + "-" + path.name
            archive.write(path, arcname)
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            manifest["files"].append({
                "artifactId": aid,
                "name": artifact.get("name"),
                "archivePath": arcname,
                "sha256": digest,
            })
        archive.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
        steps = []
        for record in project.get("calculations") or []:
            tool = record.get("tool")
            args = record.get("args") or {}
            if not tool:
                continue
            steps.append("- `{tid}`：POST `/v1/projects/{{projectId}}/execute`，body = `{body}`".format(
                tid=record.get("calculationId") or "?",
                body=json.dumps({"tool": tool, "args": args}, ensure_ascii=False),
            ))
        recompute = "\n".join([
            "# 复算说明（recompute）",
            "",
            "本包由本地 ChemELLM 内核生成；`project.json` / `requirements.json` / `plan.json` / "
            "`calculations.json` 是工程数据，`files/` 是成果原件，`manifest.json` 记录文件哈希。",
            "",
            "## 复算步骤",
            "1. 启动 Codex MCP，或启动 HTTP 内核：`scripts\\start-kernel.cmd`（默认 http://127.0.0.1:8766）。",
            "2. 新建空项目或复用原项目 ID，按 `calculations.json` 顺序重放下列请求：",
            "",
            *(steps or ["（本项目没有可重放的计算记录）"]),
            "",
            "3. 重新生成成果文件（DOCX/PDF/SVG/DXF/xlsx/FCStd/STEP）并调用 `chem_review` 自检。",
            "4. 校验：对 `files/` 中每个文件计算 SHA256，与 `manifest.json` 比较。",
            "",
            "## 边界",
            "- 自动校核 ≠ 专业审核；本包为工作草稿，审核状态以项目记录为准。",
            "- 复算只保证同一输入下模型可重放；不证明模型假设、物性或工程结论正确。",
            "- 环境差异（Python/FreeCAD/LibreOffice 版本）可能改变文件字节，数值结果以 JSON 字段为准。",
            "",
        ])
        archive.writestr("recompute.md", recompute)
        archive.writestr("README.txt", (
            "ChemELLM local kernel delivery bundle.\n"
            "Generated locally; every file above is referenced by manifest.json.\n"
            "Replay instructions: recompute.md.\n"
            "This bundle is a work draft: automatic checks and professional review are separate.\n"
        ))
    digest = hashlib.sha256(target.read_bytes()).hexdigest()
    entry = {
        "name": filename,
        "bytes": target.stat().st_size,
        "sha256": digest,
        "createdAt": _utc(),
        "downloadPath": f"/v1/bundles/{filename}",
    }
    project["bundle"] = entry
    with _write_lock:
        save_project(project)
    return {"ok": True, "phase": project["phase"], "bundle": entry, "summary": summary_of(project)}


@app.get("/v1/bundles/{filename}")
def bundle_download(filename: str):
    target = (BUNDLES_DIR / filename).resolve()
    if not target.is_file() or not target.is_relative_to(BUNDLES_DIR.resolve()):
        raise HTTPException(status_code=404, detail="成果包不存在")
    return FileResponse(target, filename=target.name)


if __name__ == "__main__":
    import uvicorn

    host = os.getenv("CHEMELLM_KERNEL_HOST", "127.0.0.1")
    port = int(os.getenv("CHEMELLM_KERNEL_PORT", "8766"))
    print(f"ChemELLM local kernel {KERNEL_VERSION} -> http://{host}:{port} (data: {DATA_ROOT})")
    uvicorn.run(app, host=host, port=port, log_level="info")
