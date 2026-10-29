"""P&ID engineering data chain: line list, MTO, instrument index, I/O list, C&E.

The drawing model (nodes/edges) is the single source of truth. This module turns
it into procurement / automation deliverables without inventing tags or sizes:

- line list and material take-off (MTO);
- instrument index and I/O list (signal type inferred from tag function);
- interlock cause & effect matrix and modulating-signal matrix;
- XLSX (multi-sheet), CSV, DOCX and JSON outputs.

MTO quantities that need routing/valve details we do not model (elbows, flanges,
gaskets) are explicitly labelled as estimates in the remark column.
"""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from pid_sheet import EQUIPMENT_KINDS, INSTRUMENT_KINDS, VALVE_KINDS, validate_pid

TAG_RE = re.compile(r"^([A-Z]+)-?(\d+[A-Z]?)$")
DN_RE = re.compile(r"(?:DN|NPS)\s*([0-9]+(?:\.[0-9]+)?)", re.IGNORECASE)

FUNCTION_IO = {
    "TI": "AI", "PI": "AI", "FI": "AI", "LI": "AI", "AI": "AI", "PDI": "AI",
    "TIC": "AI+AO", "PIC": "AI+AO", "FIC": "AI+AO", "LIC": "AI+AO", "AIC": "AI+AO",
    "XI": "DI", "XIC": "DI+DO", "SI": "DI",
    "FV": "AO", "LV": "AO", "PV": "AO", "TV": "AO", "XV": "DO", "HV": "DO",
}


def _function(tag: str) -> str:
    match = TAG_RE.match(tag)
    return match.group(1) if match else tag


def _dn(text: str) -> str:
    match = DN_RE.search(text or "")
    return f"DN{match.group(1)}" if match else "—"


def _ordered_unique(values):
    seen = set()
    result = []
    for value in values:
        if value not in seen:
            seen.add(value)
            result.append(value)
    return result


def build_pid_data(model: dict) -> dict:
    nodes = {str(node.get("id")): node for node in (model.get("nodes") or []) if node.get("id")}
    edges = list(model.get("edges") or [])
    process = [edge for edge in edges if (edge.get("kind") or "process") == "process"]
    signals = [edge for edge in edges if (edge.get("kind") or "process") == "signal"]
    interlocks = [edge for edge in edges if (edge.get("kind") or "process") == "interlock"]

    # pipe degree for tee estimation
    degree = {node_id: 0 for node_id in nodes}
    for edge in process:
        degree[str(edge.get("source"))] = degree.get(str(edge.get("source")), 0) + 1
        degree[str(edge.get("target"))] = degree.get(str(edge.get("target")), 0) + 1

    line_rows = []
    for edge in process:
        source = str(edge.get("source"))
        target = str(edge.get("target"))
        spec = str(edge.get("spec") or "")
        line_rows.append({
            "管线号": str(edge.get("line_no") or "—"),
            "介质": str(edge.get("medium") or edge.get("label") or "—"),
            "起点": source,
            "终点": target,
            "规格": spec or _dn(str(edge.get("line_no") or "")),
            "管道等级": str(edge.get("piping_class") or "—"),
            "长度(m)": str(edge.get("length_m") or ""),
            "保温/伴热": str(edge.get("insulation") or "—"),
            "备注": str(edge.get("remark") or ""),
        })

    # MTO aggregation
    pipe_groups: dict[tuple, dict] = {}
    for edge in process:
        spec = str(edge.get("spec") or _dn(str(edge.get("line_no") or "")))
        key = (spec, str(edge.get("piping_class") or "—"))
        entry = pipe_groups.setdefault(key, {"count": 0, "length": 0.0, "has_length": False})
        entry["count"] += 1
        if edge.get("length_m") not in (None, ""):
            try:
                entry["length"] += float(edge["length_m"])
                entry["has_length"] = True
            except (TypeError, ValueError):
                pass
    mto_rows = []
    for (spec, piping_class), entry in pipe_groups.items():
        mto_rows.append({
            "类别": "管道", "规格": spec, "材质/等级": piping_class,
            "单位": "m" if entry["has_length"] else "条",
            "数量": round(entry["length"], 1) if entry["has_length"] else entry["count"],
            "来源": "P&ID 管线表汇总", "备注": "" if entry["has_length"] else "未填长度，按条数计",
        })
    valves = [node for node in nodes.values() if node.get("kind") in VALVE_KINDS]
    inline_valves = [dict(valve) for edge in process for valve in (edge.get("valves") or [])]
    valve_groups: dict[tuple, int] = {}
    for valve in valves:
        spec = str(valve.get("spec") or "")
        valve_groups[(spec or "—", str(valve.get("type") or "控制阀"))] = valve_groups.get(
            (spec or "—", str(valve.get("type") or "控制阀")), 0) + 1
    for (spec, valve_type), count in valve_groups.items():
        mto_rows.append({"类别": "阀门", "规格": spec, "材质/等级": valve_type, "单位": "台",
                         "数量": count, "来源": "P&ID 控制阀", "备注": "仅统计已建模控制阀"})
    inline_groups: dict[tuple, int] = {}
    for valve in inline_valves:
        spec = str(valve.get("spec") or "—")
        kind = str(valve.get("type") or "手动阀")
        inline_groups[(spec, kind)] = inline_groups.get((spec, kind), 0) + 1
    for (spec, kind), count in inline_groups.items():
        mto_rows.append({"类别": "阀门", "规格": spec, "材质/等级": kind, "单位": "台",
                         "数量": count, "来源": "P&ID 管道阀", "备注": "含手动/止回/安全/爆破片等"})
    instruments = [node for node in nodes.values() if node.get("kind") in INSTRUMENT_KINDS]
    instrument_groups: dict[str, int] = {}
    for instrument in instruments:
        kind = _function(str(instrument.get("id")))
        instrument_groups[kind] = instrument_groups.get(kind, 0) + 1
    for kind, count in instrument_groups.items():
        mto_rows.append({"类别": "仪表", "规格": kind, "材质/等级": "—", "单位": "台",
                         "数量": count, "来源": "P&ID 仪表", "备注": "不含随机备件"})
    elbows = len(process) * 2
    tees = sum(max(0, value - 2) for value in degree.values())
    flanges = len(valves) * 2
    mto_rows.append({"类别": "管件", "规格": "弯头", "材质/等级": "—", "单位": "个", "数量": elbows,
                     "来源": "按每线 2 个估算", "备注": "未按实际走向计算，需配管复核"})
    mto_rows.append({"类别": "管件", "规格": "三通", "材质/等级": "—", "单位": "个", "数量": tees,
                     "来源": "按流程拓扑估算", "备注": "按节点工艺连接数估算"})
    mto_rows.append({"类别": "法兰", "规格": "阀门配套", "材质/等级": "—", "单位": "片", "数量": flanges,
                     "来源": "每台控制阀 2 片估算", "备注": "不含设备管口法兰与垫片"})

    instrument_rows = []
    io_rows = []
    for instrument in instruments:
        tag = str(instrument.get("id"))
        function = _function(tag)
        signal_type = FUNCTION_IO.get(function, "—")
        targets = [str(edge.get("target")) for edge in signals + interlocks
                   if str(edge.get("source")) == tag]
        instrument_rows.append({
            "位号": tag, "功能": function,
            "说明": str(instrument.get("label") or ""),
            "类型": str(instrument.get("type") or "—"),
            "量程/规格": str(instrument.get("range") or "待定"),
            "信号类型": signal_type,
            "连接对象": "、".join(_ordered_unique(targets)) or "—",
            "联锁": str(instrument.get("interlock") or "—"),
        })
        io_rows.append({
            "位号": tag, "信号类型": signal_type,
            "量程/规格": str(instrument.get("range") or "待定"),
            "去向": "SIS/DCS 待定" if (instrument.get("interlock") or signal_type == "DI") else "DCS/PLC 待定",
            "联锁/说明": str(instrument.get("interlock") or "—"),
        })
    for valve in valves:
        tag = str(valve.get("id"))
        function = _function(tag)
        instrument_rows.append({
            "位号": tag, "功能": function, "说明": str(valve.get("label") or ""),
            "类型": str(valve.get("type") or "控制阀"), "量程/规格": str(valve.get("spec") or "—"),
            "信号类型": FUNCTION_IO.get(function, "AO"), "连接对象": "—",
            "联锁": str(valve.get("interlock") or "—"),
        })
        io_rows.append({
            "位号": tag, "信号类型": FUNCTION_IO.get(function, "AO"),
            "量程/规格": str(valve.get("spec") or "—"), "去向": "DCS/PLC 待定",
            "联锁/说明": "调节阀输出",
        })
    for valve in inline_valves:
        instrument_rows.append({
            "位号": str(valve.get("tag") or "—"), "功能": "—",
            "说明": f"管道阀（{valve.get('type') or '手动'}）", "类型": "管道阀",
            "量程/规格": str(valve.get("spec") or "—"), "信号类型": "—",
            "连接对象": "—", "联锁": str(valve.get("interlock") or "—"),
        })

    # cause & effect
    causes = _ordered_unique(str(edge.get("source")) for edge in interlocks)
    effects = _ordered_unique(str(edge.get("target")) for edge in interlocks)
    ce_rows = []
    for cause in causes:
        labels = [str(edge.get("label") or "") for edge in interlocks if str(edge.get("source")) == cause]
        row = {"原因（触发）": cause, "说明": "；".join(filter(None, labels)) or "联锁"}
        for effect in effects:
            hit = any(str(edge.get("source")) == cause and str(edge.get("target")) == effect for edge in interlocks)
            row[effect] = "X" if hit else ""
        ce_rows.append(row)

    # modulating signal matrix
    sig_causes = _ordered_unique(str(edge.get("source")) for edge in signals)
    sig_effects = _ordered_unique(str(edge.get("target")) for edge in signals)
    signal_rows = []
    for cause in sig_causes:
        row = {"来源仪表": cause}
        for effect in sig_effects:
            row[effect] = "●" if any(str(edge.get("source")) == cause and str(edge.get("target")) == effect
                                     for edge in signals) else ""
        signal_rows.append(row)

    summary = {
        "lines": len(line_rows), "mto_items": len(mto_rows),
        "instruments": len(instruments), "valves": len(valves) + len(inline_valves),
        "io_points": len(io_rows), "interlocks": len(interlocks),
        "causes": len(causes), "effects": len(effects),
    }
    return {
        "summary": summary, "lines": line_rows, "mto": mto_rows,
        "instruments": instrument_rows, "io": io_rows,
        "ce": ce_rows, "ce_effects": effects,
        "signals": signal_rows, "signal_effects": sig_effects,
    }


def _write_csv(path: Path, headers: list[str], rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def _write_xlsx(path: Path, data: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    header_fill = PatternFill("solid", fgColor="DCE8F4")
    header_font = Font(bold=True)

    def sheet(title: str, headers: list[str], rows: list[dict]):
        target = workbook.create_sheet(title)
        target.append(headers)
        for cell in target[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in rows:
            target.append([row.get(header, "") for header in headers])
        for index, header in enumerate(headers, start=1):
            width = max(10, min(28, max([len(str(header))] + [len(str(row.get(header, ""))) for row in rows]) + 2))
            target.column_dimensions[get_column_letter(index)].width = width
        target.freeze_panes = "A2"

    sheet("管线表", ["管线号", "介质", "起点", "终点", "规格", "管道等级", "长度(m)", "保温/伴热", "备注"], data["lines"])
    sheet("材料表", ["类别", "规格", "材质/等级", "单位", "数量", "来源", "备注"], data["mto"])
    sheet("仪表索引", ["位号", "功能", "说明", "类型", "量程/规格", "信号类型", "连接对象", "联锁"], data["instruments"])
    sheet("IO清单", ["位号", "信号类型", "量程/规格", "去向", "联锁/说明"], data["io"])
    ce_headers = ["原因（触发）", "说明"] + data["ce_effects"]
    sheet("联锁因果表", ce_headers, data["ce"])
    sig_headers = ["来源仪表"] + data["signal_effects"]
    sheet("调节信号表", sig_headers, data["signals"])
    sheet("说明", ["项目", "内容"], [
        {"项目": "数据来源", "内容": "由 P&ID 结构化模型（nodes/edges）自动生成"},
        {"项目": "估算项", "内容": "弯头/三通/法兰为拓扑估算，需配管专业复核"},
        {"项目": "未确定项", "内容": "量程、信号类型、去向以“待定”标注，需自控专业确认"},
        {"项目": "限制", "内容": "不替代设计文件；不得据此直接采购或施工"},
    ])
    if "Sheet" in workbook.sheetnames:
        del workbook["Sheet"]
    workbook.save(path)


def _write_docx(path: Path, data: dict, model: dict) -> None:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    document.add_heading(str(model.get("title") or "P&ID 数据表"), level=1)
    document.add_paragraph(f"图号：{model.get('drawing_number') or '—'}    "
                           f"版本：{model.get('revision') or 'A'}    单位：mm")
    document.add_paragraph("本文件由 P&ID 结构化模型自动生成；弯头/三通/法兰为拓扑估算，"
                           "量程与信号类型待自控专业确认，不能直接用于采购或施工。")

    def table(title: str, headers: list[str], rows: list[dict]):
        document.add_heading(title, level=2)
        if not rows:
            document.add_paragraph("（无数据）")
            return
        target = document.add_table(rows=1, cols=len(headers))
        target.style = "Table Grid"
        for index, header in enumerate(headers):
            cell = target.rows[0].cells[index]
            cell.text = str(header)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(8)
        for row in rows:
            cells = target.add_row().cells
            for index, header in enumerate(headers):
                cells[index].text = str(row.get(header, ""))
                for paragraph in cells[index].paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(8)

    table("管线表", ["管线号", "介质", "起点", "终点", "规格", "管道等级", "长度(m)", "保温/伴热"], data["lines"])
    table("材料表（MTO）", ["类别", "规格", "材质/等级", "单位", "数量", "来源", "备注"], data["mto"])
    table("仪表索引", ["位号", "功能", "说明", "类型", "量程/规格", "信号类型", "连接对象", "联锁"], data["instruments"])
    table("I/O 清单", ["位号", "信号类型", "量程/规格", "去向", "联锁/说明"], data["io"])
    table("联锁因果表（C&E）", ["原因（触发）", "说明"] + data["ce_effects"], data["ce"])
    table("调节信号表", ["来源仪表"] + data["signal_effects"], data["signals"])
    document.save(path)


def build_pid_data_files(model: dict, out_dir: Path) -> dict:
    validation = validate_pid(model)
    if not validation["passed"]:
        raise ValueError("P&ID 数据表生成失败：" + "；".join(validation["blockers"][:6]))
    out_dir.mkdir(parents=True, exist_ok=True)
    data = build_pid_data(model)
    xlsx = out_dir / "P&ID数据表.xlsx"
    docx = out_dir / "P&ID数据表.docx"
    json_path = out_dir / "P&ID数据.json"
    _write_xlsx(xlsx, data)
    _write_docx(docx, data, model)
    json_path.write_text(json.dumps({"model": model, "tables": data}, ensure_ascii=False, indent=2), encoding="utf-8")
    csv_paths = {
        "lines": out_dir / "管线表.csv",
        "mto": out_dir / "材料表.csv",
        "instruments": out_dir / "仪表索引.csv",
        "io": out_dir / "IO清单.csv",
        "ce": out_dir / "联锁因果表.csv",
    }
    _write_csv(csv_paths["lines"], ["管线号", "介质", "起点", "终点", "规格", "管道等级", "长度(m)", "保温/伴热", "备注"], data["lines"])
    _write_csv(csv_paths["mto"], ["类别", "规格", "材质/等级", "单位", "数量", "来源", "备注"], data["mto"])
    _write_csv(csv_paths["instruments"], ["位号", "功能", "说明", "类型", "量程/规格", "信号类型", "连接对象", "联锁"], data["instruments"])
    _write_csv(csv_paths["io"], ["位号", "信号类型", "量程/规格", "去向", "联锁/说明"], data["io"])
    _write_csv(csv_paths["ce"], ["原因（触发）", "说明"] + data["ce_effects"], data["ce"])
    return {"validation": validation, "data": data, "xlsx": xlsx, "docx": docx, "json": json_path, "csv": csv_paths}
