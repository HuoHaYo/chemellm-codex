"""Equipment datasheet / specification sheet generator.

Turns project data (or explicit arguments) into vendor-facing datasheets:
design conditions, materials, dimensions, nozzles, supply scope, inspection,
standard references and limitations. Missing values stay "待定/待厂家确认" and
are never invented. Outputs DOCX, PDF and JSON.
"""

from __future__ import annotations

import json
import math
import re
from pathlib import Path

TYPES = {"vessel": "容器", "column": "塔器", "tank": "储罐", "exchanger": "换热器",
         "pump": "泵", "compressor": "压缩机"}
TAG_RE = re.compile(r"^[A-Z]{1,4}-?\d+[A-Z]?$")

DEFAULT_SUPPLY = [
    "壳体、封头及接管（含法兰）",
    "裙座/鞍座/支腿等支撑结构",
    "设备铭牌、出厂合格证与质量证明文件",
    "防腐涂装（按项目规定，待确认）",
    "随机备件与专用工具（按合同，待确认）",
]
DEFAULT_INSPECTION = [
    "无损检测比例与方法：按 GB/T 150 及设计文件要求（待定）",
    "水压试验/气密试验：按 GB/T 150.1 与设计文件（待定）",
    "材料复验、热处理与硬度：按材料标准与工艺要求（待定）",
    "涂装与包装运输：按项目技术条件（待定）",
]
LIMITS = [
    "本数据表由项目数据/参数生成，不替代设计图与采购技术协议。",
    "标注“待定/待厂家确认”的项目必须由设计、工艺与厂家共同确认。",
    "壁厚与重量为简化估算（GB/T 150 内压式），未计开孔补强、焊缝、内件与制造公差。",
    "高温蠕变、外压/真空、疲劳、低温冲击等工况未在本表校核，需另行计算。",
]


def _num(value, default=None):
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _round(value, digits=2):
    return round(float(value), digits) if value is not None else None


def _validate(payload: dict) -> list[str]:
    errors: list[str] = []
    items = payload.get("items")
    if not isinstance(items, list) or not items:
        return ["items 必须是非空数组"]
    seen: set[str] = set()
    for index, item in enumerate(items):
        if not isinstance(item, dict):
            errors.append(f"items[{index}] 必须是对象")
            continue
        tag = str(item.get("tag") or "").strip()
        if not tag:
            errors.append(f"items[{index}]：缺少 tag")
        elif tag in seen:
            errors.append(f"items[{index}]：位号重复 {tag}")
        else:
            seen.add(tag)
        kind = str(item.get("equipment_type") or "").strip().lower()
        if kind not in TYPES:
            errors.append(f"{tag or index}：equipment_type 必须是 {'/'.join(TYPES)}")
        pressure = _num(item.get("design_pressure_mpa"))
        if pressure is not None and not (0.0 < pressure <= 35.0):
            errors.append(f"{tag}：设计压力应在 0~35 MPa（GB/T 150.1 范围）")
        temperature = _num(item.get("design_temperature_c"))
        if temperature is not None and not (-269.0 <= temperature <= 900.0):
            errors.append(f"{tag}：设计温度应在 −269~900 ℃（GB/T 150.1 范围）")
        diameter = _num(item.get("diameter_mm"))
        if diameter is not None and not (100.0 <= diameter <= 20000.0):
            errors.append(f"{tag}：直径超出 100~20000 mm")
        length = _num(item.get("tangent_length_mm"))
        if length is not None and not (100.0 <= length <= 120000.0):
            errors.append(f"{tag}：切线长度超出 100~120000 mm")
        quantity = _num(item.get("quantity"), 1.0)
        if quantity is not None and quantity < 1:
            errors.append(f"{tag}：数量必须 ≥1")
        for nozzle_index, nozzle in enumerate(item.get("nozzles") or []):
            dn = _num(nozzle.get("dn"))
            if dn is not None and not (10.0 <= dn <= 2000.0):
                errors.append(f"{tag}：接管[{nozzle_index}] DN 超出 10~2000")
    return errors


def _thickness(item: dict) -> dict:
    pressure = _num(item.get("design_pressure_mpa"))
    diameter = _num(item.get("diameter_mm"))
    stress = _num(item.get("allowable_stress_mpa"))
    phi = _num(item.get("joint_efficiency"), 1.0) or 1.0
    c1 = _num(item.get("mill_tolerance_mm"), 0.3) or 0.3
    c2 = _num(item.get("corrosion_allowance_mm"), 2.0) or 2.0
    if not pressure or not diameter or not stress:
        return {"shell_required_mm": None, "shell_selected_mm": None,
                "head_required_mm": None, "head_selected_mm": None,
                "note": "缺少设计压力/直径/许用应力，壁厚待计算"}
    shell = pressure * diameter / (2.0 * stress * phi - pressure) + c1 + c2
    head = pressure * diameter / (2.0 * stress * phi - 0.5 * pressure) + c1 + c2
    series = [6, 8, 10, 12, 14, 16, 18, 20, 22, 24, 26, 28, 30, 32, 34, 36, 38, 40]
    select = next((plate for plate in series if plate >= shell), series[-1])
    return {"shell_required_mm": _round(shell), "shell_selected_mm": select,
            "head_required_mm": _round(head), "head_selected_mm": select,
            "note": "GB/T 150.1 内压简化式；未计开孔补强"}


def _volume(item: dict) -> dict:
    diameter = _num(item.get("diameter_mm"))
    length = _num(item.get("tangent_length_mm"))
    if not diameter or not length:
        return {"volume_m3": None}
    head = _num(item.get("head_height_mm"), diameter / 4.0) or diameter / 4.0
    shell = math.pi / 4.0 * diameter ** 2 * length / 1e9
    heads = 2.0 * (2.0 / 3.0) * math.pi * (diameter / 2.0) ** 2 * head / 1e9
    return {"volume_m3": _round(shell + heads, 1)}


def _weight(item: dict, thickness: dict) -> dict:
    diameter = _num(item.get("diameter_mm"))
    length = _num(item.get("tangent_length_mm"))
    selected = thickness.get("shell_selected_mm")
    if not diameter or not length or not selected:
        return {"weight_kg": None}
    outer = diameter + 2.0 * selected
    volume = math.pi / 4.0 * (outer ** 2 - diameter ** 2) * length
    weight = volume * 7850e-9
    return {"weight_kg": _round(weight, 0), "note": "仅筒体，不含封头/裙座/内件"}


def _normalize(item: dict, index: int) -> dict:
    kind = str(item.get("equipment_type") or "vessel").lower()
    nozzles = []
    for nozzle_index, nozzle in enumerate(item.get("nozzles") or []):
        nozzles.append({
            "tag": str(nozzle.get("tag") or f"N{nozzle_index + 1}"),
            "dn": _num(nozzle.get("dn")),
            "service": str(nozzle.get("service") or "待定"),
            "flange": str(nozzle.get("flange") or "待定"),
            "position": str(nozzle.get("position") or "待定"),
            "elevation_mm": _num(nozzle.get("elevation_mm")),
        })
    normalized = {
        "tag": str(item.get("tag") or f"EQ-{index + 1}").strip(),
        "name": str(item.get("name") or TYPES.get(kind, "设备")),
        "equipment_type": kind,
        "type_name": TYPES.get(kind, kind),
        "quantity": _num(item.get("quantity"), 1.0) or 1.0,
        "standard": str(item.get("standard") or "GB/T 150.1-2024（如适用）"),
        "design_pressure_mpa": _num(item.get("design_pressure_mpa")),
        "design_temperature_c": _num(item.get("design_temperature_c")),
        "operating_pressure_mpa": _num(item.get("operating_pressure_mpa")),
        "operating_temperature_c": _num(item.get("operating_temperature_c")),
        "medium": str(item.get("medium") or "待定"),
        "corrosion_allowance_mm": _num(item.get("corrosion_allowance_mm")),
        "joint_efficiency": _num(item.get("joint_efficiency")),
        "allowable_stress_mpa": _num(item.get("allowable_stress_mpa")),
        "diameter_mm": _num(item.get("diameter_mm")),
        "tangent_length_mm": _num(item.get("tangent_length_mm")),
        "head_type": str(item.get("head_type") or "2:1 椭圆封头"),
        "head_height_mm": _num(item.get("head_height_mm")),
        "material": str(item.get("material") or "待定"),
        "shell_material": str(item.get("shell_material") or item.get("material") or "待定"),
        "head_material": str(item.get("head_material") or item.get("material") or "待定"),
        "nozzle_material": str(item.get("nozzle_material") or item.get("material") or "待定"),
        "nozzles": nozzles,
        "heat_duty_kw": _num(item.get("heat_duty_kw")),
        "area_m2": _num(item.get("area_m2")),
        "flow_m3_h": _num(item.get("flow_m3_h")),
        "head_m": _num(item.get("head_m")),
        "power_kw": _num(item.get("power_kw")),
        "pressure_ratio": _num(item.get("pressure_ratio")),
        "supply_scope": list(item.get("supply_scope") or DEFAULT_SUPPLY),
        "inspection": list(item.get("inspection") or DEFAULT_INSPECTION),
        "notes": str(item.get("notes") or ""),
    }
    normalized["thickness"] = _thickness(normalized)
    normalized["volume"] = _volume(normalized)
    normalized["weight"] = _weight(normalized, normalized["thickness"])
    return normalized


def build_datasheet(payload: dict) -> dict:
    errors = _validate(payload)
    if errors:
        raise ValueError("设备数据表参数错误：" + "；".join(errors[:8]))
    items = [_normalize(item, index) for index, item in enumerate(payload.get("items") or [])]
    warnings = []
    for item in items:
        if item["design_pressure_mpa"] is None:
            warnings.append(f"{item['tag']}：设计压力待定")
        if item["design_temperature_c"] is None:
            warnings.append(f"{item['tag']}：设计温度待定")
        if item["equipment_type"] in {"vessel", "column", "tank"} and item["diameter_mm"] is None:
            warnings.append(f"{item['tag']}：直径待定")
        if item["equipment_type"] == "exchanger" and item["area_m2"] is None:
            warnings.append(f"{item['tag']}：换热面积待定（可由 heat_exchanger 计算后回填）")
        if item["equipment_type"] == "pump" and item["flow_m3_h"] is None:
            warnings.append(f"{item['tag']}：泵流量/扬程待厂家确认")
    return {
        "project": str(payload.get("project") or ""),
        "items": items,
        "warnings": warnings,
        "limits": list(payload.get("limits") or LIMITS),
        "generated_by": "kernel equipment_datasheet",
    }


def _design_table(item: dict) -> list[list[str]]:
    def value(number, unit, digits=2):
        return "待定" if number is None else f"{number:.{digits}f} {unit}".strip()

    return [
        ["设计压力", value(item["design_pressure_mpa"], "MPa(G)")],
        ["设计温度", value(item["design_temperature_c"], "℃", 0)],
        ["操作压力", value(item["operating_pressure_mpa"], "MPa(G)")],
        ["操作温度", value(item["operating_temperature_c"], "℃", 0)],
        ["介质", item["medium"]],
        ["腐蚀裕量", value(item["corrosion_allowance_mm"], "mm")],
        ["焊接接头系数", value(item["joint_efficiency"], "")],
        ["设计标准", item["standard"]],
    ]


def _material_table(item: dict) -> list[list[str]]:
    rows = [
        ["壳体材料", item["shell_material"]],
        ["封头材料", item["head_material"]],
        ["接管材料", item["nozzle_material"]],
        ["内径 / 切线长", f"{item['diameter_mm'] or '待定'} / {item['tangent_length_mm'] or '待定'} mm"],
        ["封头形式", item["head_type"]],
        ["容积（估算）", f"{item['volume']['volume_m3']} m3" if item["volume"]["volume_m3"] else "待定"],
        ["筒体壁厚（计算/选取）", (
            f"{item['thickness']['shell_required_mm']} / {item['thickness']['shell_selected_mm']} mm"
            if item["thickness"]["shell_selected_mm"] else "待计算（需许用应力）")],
        ["筒体重量（估算）", f"{item['weight']['weight_kg']:.0f} kg" if item["weight"]["weight_kg"] else "待定"],
    ]
    if item["equipment_type"] == "exchanger":
        rows.extend([
            ["换热面积", f"{item['area_m2']} m²" if item["area_m2"] else "待定"],
            ["热负荷", f"{item['heat_duty_kw']} kW" if item["heat_duty_kw"] else "待定"],
        ])
    if item["equipment_type"] == "pump":
        rows.extend([
            ["流量", f"{item['flow_m3_h']} m3/h" if item["flow_m3_h"] else "待厂家确认"],
            ["扬程", f"{item['head_m']} m" if item["head_m"] else "待厂家确认"],
            ["轴功率", f"{item['power_kw']} kW" if item["power_kw"] else "待厂家确认"],
        ])
    if item["equipment_type"] == "compressor":
        rows.extend([
            ["流量", f"{item['flow_m3_h']} m3/h" if item["flow_m3_h"] else "待厂家确认"],
            ["压比", f"{item['pressure_ratio']}" if item["pressure_ratio"] else "待厂家确认"],
            ["轴功率", f"{item['power_kw']} kW" if item["power_kw"] else "待厂家确认"],
        ])
    return rows


def _nozzle_table(item: dict) -> list[list[str]]:
    if not item["nozzles"]:
        return [["—", "—", "—", "—", "—"]]
    rows = []
    for nozzle in item["nozzles"]:
        rows.append([
            nozzle["tag"],
            f"DN{nozzle['dn']:.0f}" if nozzle["dn"] else "待定",
            nozzle["service"],
            nozzle["flange"],
            nozzle["position"],
        ])
    return rows


def render_docx(path: Path, data: dict) -> None:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    document.add_heading("设备数据表 / 规格书", level=1)
    if data["project"]:
        document.add_paragraph("项目：" + data["project"])
    document.add_paragraph("说明：本数据表由项目数据生成，供询价与技术协议使用；"
                           "标注“待定/待厂家确认”的项目必须由设计、工艺与厂家确认。")

    for item in data["items"]:
        document.add_heading(f"{item['tag']}  {item['name']}（{item['type_name']}）", level=2)
        document.add_paragraph(f"数量：{item['quantity']:.0f} 台    设计标准：{item['standard']}")

        def table(title, headers, rows):
            document.add_heading(title, level=3)
            target = document.add_table(rows=1, cols=len(headers))
            target.style = "Table Grid"
            for index, header in enumerate(headers):
                cell = target.rows[0].cells[index]
                cell.text = header
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.size = Pt(8)
            for row in rows:
                cells = target.add_row().cells
                for index, value in enumerate(row):
                    cells[index].text = str(value)
                    for paragraph in cells[index].paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(8)

        table("设计条件", ["项目", "数值/说明"], _design_table(item))
        table("材料与主要尺寸", ["项目", "数值/说明"], _material_table(item))
        table("接管表", ["位号", "DN", "用途", "法兰", "方位"], _nozzle_table(item))
        table("供货范围", ["序号", "内容"], [[str(i + 1), text] for i, text in enumerate(item["supply_scope"])])
        table("检验与试验", ["序号", "要求"], [[str(i + 1), text] for i, text in enumerate(item["inspection"])])
        if item["notes"]:
            document.add_paragraph("备注：" + item["notes"])
        document.add_paragraph("")
    document.add_heading("限制与说明", level=2)
    for index, text in enumerate(data["limits"]):
        document.add_paragraph(f"{index + 1}. {text}")
    document.save(path)


def render_pdf(path: Path, data: dict) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font_name = "Helvetica"
    for candidate in (Path(r"C:\Windows\Fonts\simhei.ttf"), Path(r"C:\Windows\Fonts\Deng.ttf"),
                      Path(r"C:\Windows\Fonts\msyh.ttc")):
        if not candidate.is_file():
            continue
        try:
            if "ChemCJK" not in pdfmetrics.getRegisteredFontNames():
                font = TTFont("ChemCJK", str(candidate), subfontIndex=0) if candidate.suffix.lower() == ".ttc" else TTFont("ChemCJK", str(candidate))
                pdfmetrics.registerFont(font)
            font_name = "ChemCJK"
            break
        except Exception:
            continue

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("title", parent=styles["Heading1"], fontName=font_name, fontSize=16)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontName=font_name, fontSize=12, spaceBefore=8)
    h3 = ParagraphStyle("h3", parent=styles["Heading3"], fontName=font_name, fontSize=10, spaceBefore=6)
    body = ParagraphStyle("body", parent=styles["BodyText"], fontName=font_name, fontSize=8.5, leading=12)
    small = ParagraphStyle("small", parent=body, fontSize=7.5, leading=10)

    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=15 * mm, bottomMargin=15 * mm,
                            title="设备数据表")
    story = [Paragraph("设备数据表 / 规格书", title_style)]
    if data["project"]:
        story.append(Paragraph("项目：" + data["project"], body))
    story.append(Paragraph("本数据表由项目数据生成，供询价与技术协议使用；"
                           "标注“待定/待厂家确认”的项目必须由设计、工艺与厂家确认。", small))
    story.append(Spacer(1, 4))

    def make_table(title, headers, rows):
        story.append(Paragraph(title, h3))
        table_data = [[Paragraph(f"<b>{header}</b>", small) for header in headers]]
        for row in rows:
            table_data.append([Paragraph(str(cell), small) for cell in row])
        table = Table(table_data, repeatRows=1, colWidths=[70 * mm, 105 * mm] if len(headers) == 2 else None)
        table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE8F4")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(table)
        story.append(Spacer(1, 3))

    for item in data["items"]:
        story.append(Paragraph(f"{item['tag']}  {item['name']}（{item['type_name']}）", h2))
        story.append(Paragraph(f"数量：{item['quantity']:.0f} 台    设计标准：{item['standard']}", body))
        make_table("设计条件", ["项目", "数值/说明"], _design_table(item))
        make_table("材料与主要尺寸", ["项目", "数值/说明"], _material_table(item))
        make_table("接管表", ["位号", "DN", "用途", "法兰", "方位"], _nozzle_table(item))
        make_table("供货范围", ["序号", "内容"], [[str(i + 1), text] for i, text in enumerate(item["supply_scope"])])
        make_table("检验与试验", ["序号", "要求"], [[str(i + 1), text] for i, text in enumerate(item["inspection"])])
        if item["notes"]:
            story.append(Paragraph("备注：" + item["notes"], small))

    story.append(Paragraph("限制与说明", h2))
    for index, text in enumerate(data["limits"]):
        story.append(Paragraph(f"{index + 1}. {text}", small))
    doc.build(story)


def build_datasheet_files(payload: dict, out_dir: Path) -> dict:
    data = build_datasheet(payload)
    out_dir.mkdir(parents=True, exist_ok=True)
    docx = out_dir / "设备数据表.docx"
    pdf = out_dir / "设备数据表.pdf"
    json_path = out_dir / "设备数据表.json"
    render_docx(docx, data)
    render_pdf(pdf, data)
    json_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"data": data, "docx": docx, "pdf": pdf, "json": json_path}
