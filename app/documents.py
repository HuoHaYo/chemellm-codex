import html
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from app.config import ARTIFACTS, MAX_UPLOAD
from app.db import uid, now, execute, one
from app.competition import seed_requirements

ALLOWED={".pdf",".docx",".pptx",".xlsx",".txt",".md",".csv",".json",".xml",".html",".png",".jpg",".jpeg",".svg",".dxf"}

def safe_name(name):
    return re.sub(r'[\\/:*?"<>|\x00-\x1f]',"_",Path(name).name)[:180] or "document"

def artifact_record(project_id,path,name=None,kind=None,preview=None,text="",metadata=None,artifact_id=None):
    ident=artifact_id or uid()
    details={**(metadata or {}),"bytes":path.stat().st_size,"sha256":hashlib.sha256(path.read_bytes()).hexdigest()}
    execute("INSERT INTO artifacts(id,project_id,name,kind,path,preview_path,text,metadata,created_at) VALUES(?,?,?,?,?,?,?,?,?)",(ident,project_id,name or path.name,kind or path.suffix.lstrip("."),str(path),str(preview) if preview else None,text,json.dumps(details,ensure_ascii=False),now()))
    return public_artifact(one("SELECT * FROM artifacts WHERE id=?",(ident,)))

def public_artifact(row):
    result={k:v for k,v in row.items() if k not in {"path","preview_path","text"}}
    result["metadata"]=json.loads(result["metadata"])
    result["download_url"]=f'/api/artifacts/{row["id"]}/download'
    result["preview_url"]=f'/api/artifacts/{row["id"]}/preview'
    result["has_visual_preview"]=bool(row["preview_path"]) or row["kind"] in {"pdf","svg","png","jpg","jpeg"}
    return result

def validate_office(path):
    with zipfile.ZipFile(path) as archive:
        files=archive.infolist()
        if len(files)>5000 or sum(i.file_size for i in files)>100*1024*1024:
            raise ValueError("Office 文件解压尺寸超限")
        if any("vbaproject" in i.filename.lower() for i in files):
            raise ValueError("不接受含宏的 Office 文件")

def office_pdf(path):
    executable=shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        return None
    with tempfile.TemporaryDirectory(prefix="chemellm-lo-") as temp:
        subprocess.run([executable,"-env:UserInstallation="+Path(temp).as_uri(),"--headless","--convert-to","pdf","--outdir",str(path.parent),str(path)],capture_output=True,timeout=60,check=True)
    target=path.with_suffix(".pdf")
    return target if target.exists() else None

def ingest(project_id,name,content,*,html_extraction=None,metadata_extra=None):
    name=safe_name(name)
    suffix=Path(name).suffix.lower()
    if suffix not in ALLOWED:
        raise ValueError("支持 PDF、DOCX、PPTX、XLSX、文本、JATS XML、HTML、JSON、CSV、SVG/DXF 和 PNG/JPEG")
    if len(content)>MAX_UPLOAD:
        raise ValueError("文件不得超过 25 MiB")
    ident=uid()
    folder=ARTIFACTS/ident
    folder.mkdir()
    path=folder/name
    path.write_bytes(content)
    pages=[]
    text=""
    preview=None
    meta={"bytes":len(content),"sha256":hashlib.sha256(content).hexdigest(),"warnings":[]}
    try:
        if suffix==".pdf":
            if not content.startswith(b"%PDF-"):
                raise ValueError("文件不是有效 PDF")
            reader=PdfReader(path)
            if len(reader.pages)>600:
                raise ValueError("PDF 页数超过600页限制")
            pages=[{"page":i+1,"text":page.extract_text() or ""} for i,page in enumerate(reader.pages)]
            text="\n".join(p["text"] for p in pages)
            if not text.strip():
                meta["warnings"].append("扫描 PDF 没有可提取的原生文本；可显式请求AI图像转录，不能自动解析其要求。")
        elif suffix in {".docx",".pptx"}:
            validate_office(path)
            if suffix==".docx":
                doc=Document(path)
                text="\n".join([p.text for p in doc.paragraphs]+[" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
            else:
                pres=Presentation(path)
                pages=[{"page":i+1,"text":"\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)} for i,slide in enumerate(pres.slides)]
                text="\n".join(f'第 {p["page"]} 页\n{p["text"]}' for p in pages)
            try:
                preview=office_pdf(path)
            except (OSError,subprocess.SubprocessError):
                meta["warnings"].append("Office 保真转换失败，提供文本预览与原文件下载。")
            if not preview:
                meta["warnings"].append("当前为文本预览；安装 LibreOffice 可启用 PDF 保真预览。")
        elif suffix in {".svg",".dxf"}:
            from app.drawing_io import svg_summary,dxf_summary
            if suffix==".svg":
                summary=svg_summary(content)
            else:
                target=folder/"drawing-preview.svg"
                summary=dxf_summary(path,target)
                preview=target if target.exists() else None
            text="\n".join(t["text"] for t in summary["text_elements"])
            meta["drawing_summary"]={k:v for k,v in summary.items() if k!="text_elements"}
            meta["warnings"].append("矢量文字/图元读取与预览，不自动认定管线、仪表连接；内置拓扑待核对。")
        elif suffix==".xml":
            from app.fulltext import extract_jats
            parsed=extract_jats(content)
            text="\n\n".join(b["text"] for b in parsed["blocks"])
            meta["jats"]={k:parsed[k] for k in ("title","DOI","license_text","authors","journal","year")}
            meta["warnings"].extend(parsed["warnings"])
        elif suffix==".html":
            if html_extraction is None:
                from app.public_sources import extract_html
                html_extraction=extract_html(content)
            text="\n\n".join(block["text"] for block in html_extraction["blocks"])
            meta["html"]={key:html_extraction[key] for key in ("title","links","license_text","encoding","extractor")}
            meta["warnings"].extend(html_extraction["warnings"])
        elif suffix in {".csv",".xlsx"}:
            from app.tabular import read_table
            table=read_table(path)
            meta["table"]={"columns":table["columns"],"rows":len(table["rows"]),"sheet":table["sheet"],"sheets":table["sheets"],"blank_rows":table["blank_rows"]}
            text="\t".join(table["columns"])+"\n"+"\n".join("\t".join(str(v) if v is not None else "" for v in r) for r in table["rows"])
            meta["warnings"].append("表格预览不保留格式；数据分析中不会执行公式或自动补齐缺失值。")
        elif suffix in {".png",".jpg",".jpeg"}:
            from PIL import Image
            with Image.open(path) as im:
                im.verify()
            meta["warnings"].append("图片原件可供AI视觉转录；不自动认定图纸管线、仪表连接或尺寸语义。")
        else:
            text=content.decode("utf-8-sig")
        meta["pages"]=pages
        meta.update(metadata_extra or {})
        meta["competition_match"]=all(token in text for token in ["2026","苯乙烯","设计任务书","gbcbig.shx"])
        record=artifact_record(project_id,path,name,preview=preview,text=text[:2000000],metadata=meta,artifact_id=ident)
        if meta["competition_match"]:
            seed_requirements(project_id)
        return record
    except Exception:
        shutil.rmtree(folder)
        raise

def text_pdf(path,title,body):
    if "STSong-Light" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles=getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName="STSong-Light"
        style.wordWrap="CJK"
    styles["BodyText"].leading=17
    story=[Paragraph(html.escape(title),styles["Title"]),Spacer(1,15)]
    for line in body.splitlines():
        story.append(Paragraph(html.escape(line) or "&#160;", styles["BodyText"]))
    SimpleDocTemplate(str(path),title=title).build(story)

def export_report(project_id,title,body,formats):
    ident=uid()
    folder=ARTIFACTS/ident
    folder.mkdir()
    title=safe_name(title)[:90]
    body="状态：工作草稿，尚未完成工程审核。\n\n"+body
    records=[]
    for fmt in dict.fromkeys(formats):
        # Each format has its own folder: Office conversion must never overwrite a separately authored PDF.
        target=folder/fmt
        target.mkdir()
        path=target/(title+"."+fmt)
        preview=None
        if fmt=="pdf":
            text_pdf(path,title,body)
        elif fmt=="docx":
            doc=Document()
            doc.add_heading(title,0)
            for line in body.splitlines():
                doc.add_paragraph(line)
            doc.save(path)
            try:
                preview=office_pdf(path)
            except (OSError,subprocess.SubprocessError):
                preview=None
        elif fmt=="pptx":
            pres=Presentation()
            pres.slide_width=Inches(13.333)
            pres.slide_height=Inches(7.5)
            chunks=[body[i:i+650] for i in range(0,len(body),650)] or [body]
            for i,chunk in enumerate(chunks):
                slide=pres.slides.add_slide(pres.slide_layouts[1])
                slide.shapes.title.text=title+f" · {i+1}"
                frame=slide.placeholders[1].text_frame
                frame.text=chunk
                for para in frame.paragraphs:
                    para.font.size=Pt(18)
            pres.save(path)
            try:
                preview=office_pdf(path)
            except (OSError,subprocess.SubprocessError):
                preview=None
        elif fmt=="md":
            path.write_text("# "+title+"\n\n"+body,encoding="utf-8")
        else:
            raise ValueError("不支持的导出格式")
        records.append(artifact_record(project_id,path,preview=preview,text=body,metadata={"status":"draft","preview_mode":"visual" if preview or fmt=="pdf" else "text"}))
    return {"artifacts":records,"status":"draft"}
