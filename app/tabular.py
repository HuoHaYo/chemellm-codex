"""Bounded CSV/XLSX reading; formulas are never evaluated."""
import csv
import io
import zipfile
from pathlib import Path

MAX_ROWS = 5000
MAX_COLUMNS = 100


def read_table(path: Path, sheet=None):
    sheets = []
    workbook = None
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        with zipfile.ZipFile(path) as archive:
            if len(archive.infolist()) > 5000 or sum(f.file_size for f in archive.infolist()) > 100*1024*1024:
                raise ValueError("Excel 解压尺寸超限")
            if any("vbaproject" in f.filename.lower() for f in archive.infolist()):
                raise ValueError("不接受含宏的 Excel 文件")
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        sheets = workbook.sheetnames
        selected = sheet if sheet is not None else sheets[0]
        if selected not in sheets:
            workbook.close()
            raise ValueError("指定工作表不存在")
        worksheet = workbook[selected]
        if worksheet.max_column is not None and worksheet.max_column > MAX_COLUMNS:
            workbook.close()
            raise ValueError("Excel 列数超过100；请导出需要分析的表格区域")
        iterator = worksheet.iter_rows(values_only=True)
        sheet = selected
    elif path.suffix.lower() == ".csv":
        if sheet:
            raise ValueError("CSV 不支持工作表参数")
        content = path.read_text(encoding="utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(content[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        iterator = csv.reader(io.StringIO(content), dialect=dialect, strict=True)
    else:
        raise ValueError("请选择 CSV 或 XLSX 文件；CSV 须为 UTF-8 编码")
    try:
        header = next(iterator, None)
        if not header or len(header) > MAX_COLUMNS:
            raise ValueError("表格首行须为列名，最多100列")
        headers = [str(v).strip() if v is not None else "" for v in header]
        if any(not name or len(name) > 120 for name in headers) or len(set(headers)) != len(headers):
            raise ValueError("列名必须非空、唯一且不超过120字")
        data, row_numbers, blank_rows = [], [], []
        for number, values in enumerate(iterator, 2):
            if number > MAX_ROWS+1:
                raise ValueError("表格最多5000行数据；请先选择分析区域")
            if not values or all(v is None or str(v).strip() == "" for v in values):
                blank_rows.append(number)
                continue
            if len(values) != len(headers):
                raise ValueError(f"第{number}行列数与表头不一致")
            if any(len(str(v)) > 1000 for v in values):
                raise ValueError(f"第{number}行单元格过长")
            data.append(list(values))
            row_numbers.append(number)
        if not data:
            raise ValueError("表格没有数据行")
        return {"columns": headers, "rows": data, "row_numbers": row_numbers,
                "blank_rows": blank_rows, "sheet": sheet, "sheets": sheets}
    except (csv.Error, UnicodeError, zipfile.BadZipFile) as exc:
        raise ValueError("表格无法解析，请检查编码和文件格式") from exc
    finally:
        if workbook is not None:
            workbook.close()


def numeric_column(table, name):
    import math
    if name not in table["columns"]:
        raise ValueError("数据列不存在："+name)
    index = table["columns"].index(name)
    values = []
    for row, number in zip(table["rows"], table["row_numbers"]):
        value = row[index]
        try:
            if isinstance(value, bool) or value is None or str(value).strip() == "":
                raise ValueError()
            value = float(value)
            if not math.isfinite(value):
                raise ValueError()
        except (TypeError, ValueError, OverflowError):
            raise ValueError(f"第{number}行的 {name} 列须为有限数值；缺失值、公式或文本不能自动填补") from None
        values.append(value)
    return values
