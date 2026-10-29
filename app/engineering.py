"""Shared engineering records, pinned deliverables and evidence-based gap reporting.

No engineering answers or automatic professional approvals are stored here.
Records are immutable by revision; dependencies pin exact revisions and file hashes.
"""
import hashlib
import json
import math
import re
import zipfile
from contextlib import nullcontext
from pathlib import Path
from typing import Literal

from pydantic import Field, field_validator, model_validator
from app import config
from app.db import connection, now, rows, one, uid
from app.schemas import Inputs
from app.design_data import units
from app.workflow import at, KEY


def encoded(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False)


def digest(value):
    return hashlib.sha256(value).hexdigest()


def local_file(record):
    path = Path(record['path']).resolve()
    if not path.is_relative_to(config.ARTIFACTS.resolve()) or not path.is_file():
        raise ValueError('项目文件不存在或路径越界')
    return path


class Quantity(Inputs):
    key: str = Field(pattern=KEY)
    value: float | None = Field(default=None, description='计算来源可省略，平台按结果指针读取并换算，避免抄录与舍入误差；其他来源必填')
    unit: str = Field(min_length=1, max_length=80)
    origin: Literal['assumption', 'design_choice', 'document', 'calculation']
    conditions: str = Field(min_length=1, max_length=2000)
    rationale: str = Field(min_length=1, max_length=3000)
    artifact_id: str | None = None
    page: int | None = Field(default=None, ge=1)
    snapshot_id: str | None = Field(default=None, max_length=64)
    block_id: str | None = Field(default=None, max_length=64)
    quote: str = Field(default='', max_length=10000)
    calculation_id: str | None = None
    result_path: str = Field(default='', max_length=400)
    result_unit: str = Field(default='', max_length=80)

    @model_validator(mode='after')
    def evidence(self):
        if self.origin == 'document' and not ((self.artifact_id and self.page) or
                (self.snapshot_id and self.block_id and len(self.quote.strip()) >= 5)):
            raise ValueError('文档数值必须关联原件页码或带精确引句的不可变快照段落')
        if self.origin == 'calculation' and (not self.calculation_id or not self.result_path or not self.result_unit):
            raise ValueError('计算数值须指定计算记录、结果JSON指针及结果单位')
        if self.origin != 'calculation' and self.value is None:
            raise ValueError('非计算来源须填写数值')
        if self.page and not self.artifact_id:
            raise ValueError('页码必须关联原件')
        if bool(self.snapshot_id) != bool(self.block_id) or self.snapshot_id and len(self.quote.strip()) < 5:
            raise ValueError('全文快照引用须同时给出段落编号及至少5字精确引句')
        if self.origin != 'document' and (self.snapshot_id or self.block_id or self.quote):
            raise ValueError('仅文档来源可以关联文档快照引句')
        return self


class ObjectWrite(Inputs):
    key: str = Field(pattern=KEY)
    kind: Literal['basis', 'component', 'stream', 'unit_operation', 'equipment', 'utility',
                  'cost', 'site', 'control', 'pipe', 'hazard', 'method']
    title: str = Field(min_length=1, max_length=200)
    quantities: list[Quantity] = Field(default_factory=list, max_length=100)
    attributes: dict = Field(default_factory=dict, description='非数值条件、方法、规格、来源说明；工程数值使用quantities')
    depends_on: list[str] = Field(default_factory=list, max_length=100)
    expected_revision: int = Field(default=0, ge=0, description='新建为0，更新须给出当前版本')

    @model_validator(mode='after')
    def unique(self):
        if len({q.key for q in self.quantities}) != len(self.quantities):
            raise ValueError('同一对象内数值名称不能重复')
        if len(set(self.depends_on)) != len(self.depends_on) or self.key in self.depends_on:
            raise ValueError('对象依赖重复或依赖自身')
        if len(encoded(self.attributes)) > 30000:
            raise ValueError('对象属性过长')
        return self


class FileLocation(Inputs):
    artifact_id: str
    location: str = Field(min_length=1, max_length=1000, description='实际章节、页码、图号或源文件模型入口')


class DeliverableWrite(Inputs):
    key: str = Field(pattern=KEY)
    title: str = Field(min_length=1, max_length=200)
    requirement_codes: list[str] = Field(min_length=1, max_length=100)
    object_keys: list[str] = Field(default_factory=list, max_length=200)
    depends_on: list[str] = Field(default_factory=list, max_length=100, description='上游成果key')
    formats: list[str] = Field(min_length=1, max_length=20)
    files: list[FileLocation] = Field(default_factory=list, max_length=100)
    acceptance_criteria: str = Field(min_length=1, max_length=6000)
    native_software: str = Field(default='', max_length=300, description='专业软件及目标版本；填写不表示已运行验证')
    expected_revision: int = Field(default=0, ge=0)

    @model_validator(mode='after')
    def unique(self):
        for values in [self.requirement_codes, self.object_keys, self.depends_on, self.formats,
                       [f.artifact_id for f in self.files]]:
            if len(values) != len(set(values)):
                raise ValueError('成果的要求、依赖、格式及文件不能重复')
        if self.key in self.depends_on:
            raise ValueError('成果不能依赖自身')
        if any(not re.fullmatch(r'[A-Za-z0-9_-]{1,30}', f) for f in self.formats):
            raise ValueError('格式使用不含点的扩展名')
        return self


class EngineeringQuery(Inputs):
    include_values: bool = True
    requirement_codes: list[str] = Field(default_factory=list, max_length=100)


class EngineeringPlan(Inputs):
    objects: list[ObjectWrite] = Field(default_factory=list, max_length=100)
    deliverables: list[DeliverableWrite] = Field(default_factory=list, max_length=100)

    @model_validator(mode='after')
    def valid_batch(self):
        if not self.objects and not self.deliverables:
            raise ValueError('批量计划不能为空')
        for records in [self.objects, self.deliverables]:
            if len({r.key for r in records}) != len(records):
                raise ValueError('同一批次不能重复更新同一对象或成果')
        return self


class CalculationRead(Inputs):
    calculation_id: str = Field(min_length=1, max_length=64)
    result_path: str = Field(default='', max_length=400, description='JSON Pointer，空字符串读取根节点；例如 /streams/0')
    offset: int = Field(default=0, ge=0)
    limit: int = Field(default=20, ge=1, le=100)


class RequirementSource(Inputs):
    code: str = Field(pattern=r'^[A-Z][A-Z0-9_-]{0,29}$')
    title: str = Field(min_length=1, max_length=300)
    mandatory: bool = True
    applicability: str = Field(min_length=1, max_length=3000, description='保留原文适用条件；额外或条件项不能删除')
    source_kind: Literal['document', 'user_request'] = 'document'
    snapshot_id: str | None = None
    block_id: str | None = None
    quote: str = Field(min_length=5, max_length=10000, description='原文逐字引句，不含自行补写的要求')
    formats: list[str] = Field(default_factory=list, max_length=20)

    @field_validator('formats')
    @classmethod
    def valid_formats(cls, values):
        if len(set(values)) != len(values) or any(not re.fullmatch(r'[A-Za-z0-9_-]{1,30}', value) for value in values):
            raise ValueError('要求格式必须唯一并使用不含点的扩展名')
        return values


class RequirementBatch(Inputs):
    requirements: list[RequirementSource] = Field(min_length=1, max_length=100)


def save_requirements(project_id, p):
    """Add or source-bind requirements without allowing original requirements to change."""
    from app.fulltext import read_document, DocumentReadInput, read_snapshot
    prepared = []
    for requirement in p.requirements:
        if requirement.source_kind == 'document':
            if not requirement.snapshot_id or not requirement.block_id:
                raise ValueError('文档要求须指定原件全文快照及段落编号')
            source = read_document(project_id, DocumentReadInput(snapshot_id=requirement.snapshot_id, block_ids=[requirement.block_id]))
            block = source['blocks'][0]
            if requirement.quote not in block['text']:
                raise ValueError('要求引句不在指定原文段落中：' + requirement.code)
            snapshot, _ = read_snapshot(project_id, requirement.snapshot_id)
            prepared.append((requirement, {**requirement.model_dump(), 'artifact_id': source['artifact_id'],
                                          'source_sha256': source['source_sha256'], 'locator': block['locator'],
                                          'snapshot_sha256': snapshot['content_sha256'],
                                          'ocr': block.get('ocr', False), 'review': 'not_reviewed'}))
        else:
            run = one("SELECT r.id,r.request,e.snapshot FROM runs r JOIN engineering_run_inputs e ON e.run_id=r.id "
                      "WHERE r.project_id=? AND r.status IN ('queued','running') ORDER BY r.created_at DESC LIMIT 1", (project_id,))
            if not run:
                raise ValueError('用户请求引用必须来自当前活动且已冻结的工程任务')
            request = json.loads(run['request'])
            message = request.get('message', '')
            frozen = json.loads(run['snapshot'])
            if message != frozen.get('request', {}).get('message', '') or requirement.quote not in message:
                raise ValueError('要求引句不在当前任务冻结的用户请求中：' + requirement.code)
            prepared.append((requirement, {**requirement.model_dump(), 'source_run_id': run['id'],
                                          'source_sha256': digest(message.encode()),
                                          'locator': {'page': None, 'section': 'user request', 'paragraph': None},
                                          'ocr': False, 'review': 'not_reviewed'}))
    with connection() as db:
        db.execute('BEGIN IMMEDIATE')
        codes = {r['code']: dict(r) for r in db.execute('SELECT code,title,mandatory,page FROM requirements WHERE project_id=?', (project_id,))}
        bound = {r['code'] for r in db.execute('SELECT code FROM engineering_requirement_sources WHERE project_id=?', (project_id,))}
        incoming = [r.code for r, _ in prepared]
        if len(set(incoming)) != len(incoming):
            raise ValueError('同一批要求编码重复')
        added, linked = [], []
        for requirement, source in prepared:
            existing = codes.get(requirement.code)
            if existing:
                if requirement.code in bound:
                    raise ValueError('要求编码已经关联原文；不允许模型覆盖或删除现有原题要求')
                if existing['title'] != requirement.title or bool(existing['mandatory']) != requirement.mandatory:
                    raise ValueError('已有要求的标题或必需属性与原记录不一致，不能覆盖：' + requirement.code)
                if existing['page'] is not None and source['locator'].get('page') != existing['page']:
                    raise ValueError('要求原文定位页与已核对题目页码不一致：' + requirement.code)
                linked.append(requirement.code)
            else:
                db.execute('INSERT INTO requirements(id,project_id,code,title,page,mandatory) VALUES(?,?,?,?,?,?)',
                           (uid(), project_id, requirement.code, requirement.title, source['locator'].get('page'), int(requirement.mandatory)))
                added.append(requirement.code)
            db.execute('INSERT INTO engineering_requirement_sources VALUES(?,?,?,?)', (project_id, requirement.code, encoded(source), now()))
    return {'added': added, 'source_bound': linked, 'status': 'source_linked_not_reviewed', 'phase': engineering_phase(project_id)}


def calculation_read(project_id, p):
    record = one('SELECT * FROM calculations WHERE project_id=? AND id=?', (project_id, p.calculation_id))
    if not record:
        raise ValueError('计算记录不属于当前项目或不存在')
    result = json.loads(record['result'])
    selected = at(result, p.result_path) if p.result_path else result
    total = None
    if isinstance(selected, list):
        total = len(selected)
        selected = selected[p.offset:p.offset+p.limit]
    elif isinstance(selected, dict):
        total = len(selected)
        selected = dict(list(selected.items())[p.offset:p.offset+p.limit])
    if len(encoded(selected)) > 20000:
        return {'source_calculation_id': p.calculation_id, 'tool': record['tool'], 'result_path': p.result_path,
                'result_sha256': digest(record['result'].encode()), 'total_items': total,
                'keys': list(selected) if isinstance(selected, dict) else None,
                'notice': '所选结果仍过长；用更具体的JSON指针或更小limit读取，完整原值没有被截断替代。'}
    return {'source_calculation_id': p.calculation_id, 'tool': record['tool'], 'result_path': p.result_path,
            'result_sha256': digest(record['result'].encode()), 'total_items': total, 'offset': p.offset, 'value': selected}


def _save(db, table, project_id, p, data):
    old = db.execute(f'SELECT revision FROM {table} WHERE project_id=? AND key=?', (project_id, p.key)).fetchone()
    revision = old['revision'] if old else 0
    if revision != p.expected_revision:
        raise ValueError(f'版本冲突：{p.key} 当前版本为{revision}')
    timestamp = now()
    text = encoded(data)
    db.execute(f'INSERT INTO {table} VALUES(?,?,?,?,?,?) ON CONFLICT(project_id,key) DO UPDATE SET revision=excluded.revision,data=excluded.data,updated_at=excluded.updated_at',
               (project_id, p.key, revision + 1, text, timestamp, timestamp))
    history = 'engineering_object_revisions' if table == 'engineering_objects' else 'engineering_deliverable_revisions'
    db.execute(f'INSERT INTO {history} VALUES(?,?,?,?,?)', (project_id, p.key, revision + 1, text, timestamp))
    return {'key': p.key, 'revision': revision + 1, 'data': data, 'professional_review': 'not_reviewed'}


def _acyclic(db, table, project_id, key, parents):
    graph = {r['key']: json.loads(r['data']).get('depends_on', []) for r in db.execute(f'SELECT key,data FROM {table} WHERE project_id=?', (project_id,))}
    graph[key] = parents
    def visit(node, path):
        if node in path:
            raise ValueError('工程依赖存在循环')
        for parent in graph.get(node, []):
            if parent not in graph:
                raise ValueError('上游对象或成果不存在：' + parent)
            visit(parent, path | {node})
    visit(key, set())


def _pin(db, project_id, table, keys):
    result = []
    for key in keys:
        row = db.execute(f'SELECT revision FROM {table} WHERE project_id=? AND key=?', (project_id, key)).fetchone()
        if not row:
            raise ValueError('依赖不属于当前项目或不存在：' + key)
        result.append({'key': key, 'revision': row['revision']})
    return result


def save_object(project_id, p, _db=None):
    with (connection() if _db is None else nullcontext(_db)) as db:
        if _db is None:
            db.execute('BEGIN IMMEDIATE')
        _acyclic(db, 'engineering_objects', project_id, p.key, p.depends_on)
        data = p.model_dump(exclude={'expected_revision'})
        data['object_revisions'] = _pin(db, project_id, 'engineering_objects', p.depends_on)
        for q, entry in zip(p.quantities, data['quantities']):
            if q.calculation_id:
                calc = db.execute('SELECT * FROM calculations WHERE project_id=? AND id=?', (project_id, q.calculation_id)).fetchone()
                if not calc:
                    raise ValueError('来源计算不属于当前项目')
                entry['calculation_sha256'] = digest(calc['result'].encode())
                if q.origin == 'calculation':
                    try:
                        result = at(json.loads(calc['result']), q.result_path)
                        if isinstance(result, bool) or not isinstance(result, (int, float)):
                            raise ValueError('指定结果不是数值')
                        converted = units.Quantity(result, q.result_unit).to(q.unit).magnitude
                        if q.value is not None and not math.isclose(q.value, converted, rel_tol=1e-9, abs_tol=1e-10):
                            raise ValueError('数值与指定计算结果不一致')
                        entry['value'] = float(converted)
                    except Exception as exc:
                        raise ValueError('计算结果引用错误：' + str(exc)) from exc
            try:
                value = units.Quantity(entry['value'], q.unit).to_base_units()
                if not math.isfinite(value.magnitude):
                    raise ValueError('换算结果非有限值')
            except Exception as exc:
                raise ValueError('数值单位无效：' + q.key) from exc
            entry.update(base_value=float(value.magnitude), base_unit=str(value.units))
            if q.artifact_id:
                source = db.execute('SELECT * FROM artifacts WHERE project_id=? AND id=?', (project_id, q.artifact_id)).fetchone()
                if not source:
                    raise ValueError('来源文件不属于当前项目')
                entry['source_sha256'] = digest(local_file(source).read_bytes())
                pages = json.loads(source['metadata']).get('pages', [])
                if q.page and not any(page['page'] == q.page for page in pages):
                    raise ValueError('来源页码不可验证')
            if q.snapshot_id:
                from app.fulltext import read_document, DocumentReadInput, read_snapshot
                source = read_document(project_id, DocumentReadInput(snapshot_id=q.snapshot_id, block_ids=[q.block_id]))
                block = source['blocks'][0]
                if q.quote not in block['text']:
                    raise ValueError('数值来源引句不在指定快照段落中：' + q.key)
                snapshot, _ = read_snapshot(project_id, q.snapshot_id)
                if q.artifact_id and q.artifact_id != source['artifact_id']:
                    raise ValueError('快照原件与指定原件不一致')
                if q.page and block['locator'].get('page') != q.page:
                    raise ValueError('快照页码与指定页码不一致')
                entry.update(artifact_id=source['artifact_id'], source_sha256=source['source_sha256'],
                             source_snapshot_id=q.snapshot_id, source_snapshot_sha256=snapshot['content_sha256'],
                             source_block_id=q.block_id, source_quote=q.quote, source_locator=block['locator'],
                             source_ocr=block.get('ocr', False))
        return _save(db, 'engineering_objects', project_id, p, data)


def save_deliverable(project_id, p, _db=None):
    with (connection() if _db is None else nullcontext(_db)) as db:
        if _db is None:
            db.execute('BEGIN IMMEDIATE')
        _acyclic(db, 'engineering_deliverables', project_id, p.key, p.depends_on)
        codes = {r['code'] for r in db.execute('SELECT code FROM requirements WHERE project_id=?', (project_id,))}
        if set(p.requirement_codes) - codes:
            raise ValueError('成果引用不存在的要求编码')
        data = p.model_dump(exclude={'expected_revision'})
        data['object_revisions'] = _pin(db, project_id, 'engineering_objects', p.object_keys)
        data['deliverable_revisions'] = _pin(db, project_id, 'engineering_deliverables', p.depends_on)
        for entry in data['files']:
            record = db.execute('SELECT * FROM artifacts WHERE project_id=? AND id=?', (project_id, entry['artifact_id'])).fetchone()
            if not record:
                raise ValueError('成果文件不属于当前项目')
            entry.update(sha256=digest(local_file(record).read_bytes()), name=record['name'], kind=record['kind'])
        return _save(db, 'engineering_deliverables', project_id, p, data)


def save_plan(project_id, p):
    """One transaction, including dependency pins and all revision histories."""
    def ordered(records):
        pending = {r.key: r for r in records}
        result = []
        while pending:
            ready = [r for r in pending.values() if not set(r.depends_on) & pending.keys()]
            if not ready:
                raise ValueError('批量计划包含循环依赖')
            for record in ready:
                result.append(record)
                del pending[record.key]
        return result
    with connection() as db:
        db.execute('BEGIN IMMEDIATE')
        objects = [save_object(project_id, r, db) for r in ordered(p.objects)]
        deliverables = [save_deliverable(project_id, r, db) for r in ordered(p.deliverables)]
    return {'objects': [{'key': r['key'], 'revision': r['revision']} for r in objects],
            'deliverables': [{'key': r['key'], 'revision': r['revision']} for r in deliverables],
            'phase': engineering_phase(project_id), 'professional_review': 'not_reviewed'}


def engineering_phase(project_id, requirement_codes=None):
    requirements = {r['code'] for r in rows('SELECT code FROM requirements WHERE project_id=?', (project_id,))}
    if requirement_codes:
        if set(requirement_codes)-requirements:
            raise ValueError('阶段检查范围包含不存在的要求')
        requirements=set(requirement_codes)
    source_linked = {r['code'] for r in rows('SELECT code FROM engineering_requirement_sources WHERE project_id=?', (project_id,))}
    planned = {code for r in _records(project_id, 'engineering_deliverables').values() for code in r['data']['requirement_codes']}
    objects = _records(project_id, 'engineering_objects')
    basis = [r['key'] for r in objects.values() if r['data']['kind'] == 'basis' and r['data']['quantities']
             and not _object_issues(project_id, objects, r['key'])]
    missing = sorted(requirements - planned)
    source_missing = sorted(requirements - source_linked)
    # A free-form task may have no formal requirement records.  In that case the
    # user's request is the brief and the model can establish its deliverable plan;
    # document-backed requirements, when present, must all be source-bound first.
    stage = 'requirements' if not requirements or source_missing else 'plan' if missing else 'basis' if not basis else 'execute'
    return {'stage': stage,
            'unplanned_requirements': missing, 'source_unlinked_requirements': source_missing, 'basis_keys': basis,
            'note': '执行顺序约束只检查原文要求、映射和共同依据存在，不代表设计正确或专业审核通过。'}


# The executor uses this policy in addition to hiding tools from the model.  Keeping the
# policy here makes a saved engineering state, rather than a prior model response, decide
# whether a stale batch may run.  Direct API calls remain available for deliberate human QA.
_ENGINEERING_RESEARCH = frozenset({
    'project_context', 'acceptance', 'engineering_status', 'engineering_bundle', 'ask_user',
    'document_extract', 'document_read', 'project_search', 'source_fetch', 'literature_search',
    'open_literature_search', 'pubchem', 'fulltext_import', 'claim_save', 'evidence_add',
    'evidence_check', 'evidence_export', 'dataset_inspect', 'drawing_read',
})
_ENGINEERING_PLAN = _ENGINEERING_RESEARCH | frozenset({'engineering_requirements', 'engineering_plan', 'engineering_deliverable'})
_ENGINEERING_BASIS = _ENGINEERING_PLAN | frozenset({'engineering_object', 'calculation_read', 'equipment_query'})


def engineering_tool_policy(project_id, requirement_codes=None):
    """Return current phase and the model/executor tool allowlist; ``None`` means execute phase."""
    phase = engineering_phase(project_id, requirement_codes)
    if phase['stage'] == 'requirements':
        return phase, _ENGINEERING_RESEARCH | frozenset({'engineering_requirements'})
    if phase['stage'] == 'plan':
        return phase, _ENGINEERING_PLAN
    if phase['stage'] == 'basis':
        # Calculations are needed to establish a numerical design basis.  Keep
        # generation and packaging tools behind the basis gate, while exposing
        # every registered calculation so this remains useful beyond one case.
        from app.calculations import REGISTRY
        return phase, _ENGINEERING_BASIS | frozenset(REGISTRY)
    return phase, None


def require_engineering_tool(project_id, name, requirement_codes=None):
    phase, allowed = engineering_tool_policy(project_id, requirement_codes)
    if allowed is not None and name not in allowed:
        missing = phase['source_unlinked_requirements'] if phase['stage'] == 'requirements' else phase['unplanned_requirements']
        detail = '、'.join(missing[:12]) or '共同设计依据'
        raise ValueError('工程阶段为' + phase['stage'] + '，当前不能执行 ' + name + '；先完成：' + detail)
    return phase


def _records(project_id, table):
    return {r['key']: {**r, 'data': json.loads(r['data'])} for r in rows(f'SELECT * FROM {table} WHERE project_id=? ORDER BY key', (project_id,))}


def _object_issues(project_id, objects, key, visiting=None):
    visiting = visiting or set()
    if key in visiting:
        return ['循环依赖']
    record = objects.get(key)
    if not record:
        return ['对象不存在：' + key]
    issues = []
    for pin in record['data']['object_revisions']:
        parent = objects.get(pin['key'])
        if not parent or parent['revision'] != pin['revision']:
            issues.append('上游版本变化：' + pin['key'])
        elif _object_issues(project_id, objects, pin['key'], visiting | {key}):
            issues.append('上游数据失效：' + pin['key'])
    for q in record['data']['quantities']:
        if q.get('artifact_id'):
            source = one('SELECT * FROM artifacts WHERE project_id=? AND id=?', (project_id, q['artifact_id']))
            try:
                if not source or digest(local_file(source).read_bytes()) != q['source_sha256']:
                    issues.append('来源文件已改变：' + q['key'])
            except (ValueError, OSError):
                issues.append('来源文件不可用：' + q['key'])
        if q.get('calculation_id'):
            source = one('SELECT result FROM calculations WHERE project_id=? AND id=?', (project_id, q['calculation_id']))
            if not source or digest(source['result'].encode()) != q['calculation_sha256']:
                issues.append('来源计算已改变：' + q['key'])
        if q.get('source_snapshot_id'):
            from app.fulltext import read_snapshot, source_artifact
            try:
                snapshot, content = read_snapshot(project_id, q['source_snapshot_id'])
                _, _, artifact_hash = source_artifact(project_id, snapshot['artifact_id'])
                block = next((b for b in content['blocks'] if b['id'] == q.get('source_block_id')), None)
                if (snapshot['content_sha256'] != q.get('source_snapshot_sha256') or
                    artifact_hash != q.get('source_sha256') or not block or q.get('source_quote') not in block['text']):
                    issues.append('来源快照、段落或引句已改变：' + q['key'])
            except (ValueError, OSError):
                issues.append('来源快照不可用：' + q['key'])
    return issues


def _inspect_file(record, pin, competition=False):
    issues = []
    text = ''
    try:
        path = local_file(record)
        if digest(path.read_bytes()) != pin['sha256']:
            issues.append('文件哈希变化')
        if record['kind'] == 'docx':
            from docx import Document
            doc = Document(path)
            text = '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
        elif record['kind'] == 'pdf':
            from pypdf import PdfReader
            reader = PdfReader(path)
            text = '\n'.join(page.extract_text() or '' for page in reader.pages)
            if not reader.pages:
                issues.append('PDF没有页面')
        elif record['kind'] == 'dxf':
            import ezdxf
            drawing = ezdxf.readfile(path)
            if competition and drawing.dxfversion != 'AC1018':
                issues.append('未验证AutoCAD 2004兼容版本')
            if competition:
                used = {e.dxf.style for e in drawing.modelspace() if e.dxftype() in {'TEXT', 'MTEXT'}}
                if not used or any(drawing.styles.get(style).dxf.get('bigfont', '').lower() != 'gbcbig.shx' for style in used):
                    issues.append('图纸文字未全部使用gbcbig.shx')
            issues.append('缺少目标CAD软件实际打开、字体及图纸视觉核验')
        elif record['kind'] in {'md', 'txt', 'csv', 'json'}:
            text = path.read_text(encoding='utf-8')
        elif record['kind'] not in {'svg', 'png', 'jpg', 'jpeg', 'xlsx', 'pptx'}:
            issues.append('缺少专业原生文件打开及运行核验')
        if record['kind'] in {'docx', 'pdf', 'md', 'txt'} and not text.strip():
            issues.append('文档没有可提取的实质文本')
        if re.search(r'\bTODO\b|\bTBD\b|待补(?:充|写|齐)|此处填写|占位(?:符|内容)', text, re.I):
            issues.append('发现占位或待补内容')
        if json.loads(record['metadata']).get('workflow_freshness') == 'stale':
            issues.append('文件的研究任务上游已失效')
    except Exception as exc:
        issues.append('文件无法解析：' + type(exc).__name__)
    return issues, text


def engineering_status(project_id, p=None, run_id=None):
    p = p or EngineeringQuery()
    objects = _records(project_id, 'engineering_objects')
    deliverables = _records(project_id, 'engineering_deliverables')
    requirements = rows('SELECT * FROM requirements WHERE project_id=? ORDER BY code', (project_id,))
    files = {r['id']: r for r in rows('SELECT * FROM artifacts WHERE project_id=?', (project_id,))}
    # Tests and imported projects can seed the verified competition index without
    # carrying the PDF artifact in the same database.  Require the distinctive
    # report codes as a second, conservative recognition path.
    requirement_codes = {r['code'] for r in requirements}
    competition = any(json.loads(f['metadata']).get('competition_match') for f in files.values()) or (
        {'D01', 'D02', 'D03'} <= requirement_codes and len(requirements) >= 30)
    requirement_sources = {r['code']: json.loads(r['data']) for r in rows('SELECT * FROM engineering_requirement_sources WHERE project_id=?', (project_id,))}
    if run_id is None:
        active = one("SELECT r.id FROM runs r JOIN engineering_run_inputs e ON e.run_id=r.id WHERE r.project_id=? AND r.status IN ('queued','running') ORDER BY r.created_at DESC LIMIT 1", (project_id,))
        if active:
            run_id = active['id']
    blockers = []
    object_state = []
    for key, record in objects.items():
        issues = _object_issues(project_id, objects, key)
        object_state.append({**record, 'issues': issues, 'fresh': not issues})
        if issues:
            blockers.append({'kind': 'engineering_data', 'key': key, 'issues': issues})
    cache = {}
    def check(key, visiting=None):
        if key in cache:
            return cache[key]
        visiting = visiting or set()
        if key in visiting or key not in deliverables:
            return ['成果循环依赖或上游缺失']
        data = deliverables[key]['data']
        issues = []
        for pin in data['object_revisions']:
            if pin['key'] not in objects or objects[pin['key']]['revision'] != pin['revision'] or _object_issues(project_id, objects, pin['key']):
                issues.append('工程数据版本失效：' + pin['key'])
        for pin in data['deliverable_revisions']:
            if pin['key'] not in deliverables or deliverables[pin['key']]['revision'] != pin['revision'] or check(pin['key'], visiting | {key}):
                issues.append('上游成果缺失或失效：' + pin['key'])
        kinds = {f['kind'].lower() for f in data['files']}
        if not data['files']:
            issues.append('尚未生成文件')
        for fmt in data['formats']:
            if fmt.lower() not in kinds:
                issues.append('缺少计划文件格式：' + fmt)
        codes = set(data['requirement_codes'])
        for code, limit in [('D01', 40000), ('D02', 150000), ('D03', 50000)]:
            if competition and code in codes and not {'docx', 'pdf'} <= kinds:
                issues.append(code + '要求DOCX及PDF原件')
        if competition and codes & {'D04', 'P04', 'P05', 'L02', 'L03', 'L05'} and (not kinds & {'dxf', 'dwg'} or 'pdf' not in kinds):
            issues.append('图纸须同时提供AutoCAD 2004兼容原件与PDF')
        if data['native_software'] or competition and codes & {'P03', 'D05', 'L04'}:
            issues.append('缺少声明软件/版本的独立打开及运行验证')
        for pin in data['files']:
            record = files.get(pin['artifact_id'])
            if not record:
                issues.append('文件记录缺失：' + pin['artifact_id'])
                continue
            file_issues, text = _inspect_file(record, pin, competition=competition)
            issues.extend(record['name'] + '：' + i for i in file_issues)
            for code, limit in [('D01', 40000), ('D02', 150000), ('D03', 50000)]:
                if competition and code in codes and record['kind'] in {'docx', 'pdf'} and len(re.sub(r'\s', '', text)) > limit:
                    issues.append(code + '文本字符数超过原题上限；需按期刊口径复核')
        if competition and codes & {'D01', 'D02', 'D03'}:
            issues.append('期刊格式、跨文件数值与专业内容尚无独立校核证据')
        for code in codes:
            for fmt in requirement_sources.get(code, {}).get('formats', []):
                if fmt.lower() not in kinds:
                    issues.append(code + '原文映射要求的格式缺失：' + fmt)
        cache[key] = issues
        return issues
    deliveries = []
    for key, record in deliverables.items():
        issues = check(key)
        deliveries.append({**record, 'issues': issues, 'files_checked': not issues, 'professional_review': 'not_reviewed'})
        if issues:
            blockers.append({'kind': 'deliverable', 'key': key, 'issues': issues})
    coverage = []
    for requirement in requirements:
        linked = [d for d in deliveries if requirement['code'] in d['data']['requirement_codes']]
        coverage.append({'code': requirement['code'], 'title': requirement['title'], 'mandatory': bool(requirement['mandatory']),
                         'page': requirement['page'], 'deliverables': [d['key'] for d in linked],
                         'files_checked': bool(linked) and all(d['files_checked'] for d in linked),
                         'professional_review': requirement['status']})
        if requirement['code'] in requirement_sources:
            source = requirement_sources[requirement['code']]
            coverage[-1]['source'] = source
            try:
                if source.get('source_kind') == 'user_request':
                    run = one('SELECT request FROM runs WHERE id=? AND project_id=?', (source['source_run_id'], project_id))
                    frozen_source = one('SELECT snapshot FROM engineering_run_inputs WHERE run_id=? AND project_id=?', (source['source_run_id'], project_id))
                    if not run:
                        raise ValueError('request removed')
                    message = json.loads(run['request']).get('message', '')
                    if not frozen_source or message != json.loads(frozen_source['snapshot']).get('request', {}).get('message'):
                        raise ValueError('frozen request changed')
                    if digest(message.encode()) != source['source_sha256'] or source['quote'] not in message:
                        raise ValueError('request changed')
                else:
                    record = files.get(source['artifact_id'])
                    if not record or digest(local_file(record).read_bytes()) != source['source_sha256']:
                        raise ValueError('changed')
                    from app.fulltext import read_snapshot
                    snapshot, content = read_snapshot(project_id, source['snapshot_id'])
                    block = next((b for b in content['blocks'] if b['id'] == source['block_id']), None)
                    if snapshot['content_sha256'] != source.get('snapshot_sha256') or not block or source['quote'] not in block['text']:
                        raise ValueError('snapshot changed')
            except (ValueError, OSError):
                blockers.append({'kind': 'requirement_source', 'key': requirement['code'], 'issues': ['要求原文版本已变化或缺失']})
        if not linked:
            blockers.append({'kind': 'requirement_unplanned', 'key': requirement['code'], 'issues': ['要求尚未映射到任务/成果']})
        if requirement['status'] != 'verified':
            blockers.append({'kind': 'professional_review', 'key': requirement['code'], 'issues': ['没有当前要求的专业审核证据']})
        if requirement['code'] not in requirement_sources:
            blockers.append({'kind': 'requirement_source', 'key': requirement['code'], 'issues': ['要求尚未关联可核验的原文快照段落']})
    if not requirements:
        blockers.append({'kind': 'requirements', 'key': '', 'issues': ['没有从任务原文建立要求清单']})
    frozen = None
    if run_id:
        row = one('SELECT snapshot FROM engineering_run_inputs WHERE run_id=? AND project_id=?', (run_id, project_id))
        if row:
            frozen = json.loads(row['snapshot'])
            frozen_path = Path(frozen['frozen_archive']).resolve()
            if not frozen_path.is_relative_to((config.DATA / 'engineering-runs').resolve()) or not frozen_path.is_file() or digest(frozen_path.read_bytes()) != frozen['frozen_archive_sha256']:
                blockers.append({'kind': 'frozen_archive', 'key': run_id, 'issues': ['冻结原件和平台源码包缺失或哈希变化']})
            if platform_manifest() != frozen['platform_files']:
                blockers.append({'kind': 'frozen_platform', 'key': run_id, 'issues': ['平台代码与启动时冻结版本不同']})
            for pin in frozen['inputs']:
                try:
                    if pin['id'] not in files or digest(local_file(files[pin['id']]).read_bytes()) != pin['sha256']:
                        raise ValueError('changed')
                except ValueError:
                    blockers.append({'kind': 'frozen_input', 'key': pin['id'], 'issues': ['输入原件缺失或哈希变化']})
            generated = set()
            def find_ids(value):
                if isinstance(value, dict):
                    if 'download_url' in value and value.get('id') in files:
                        generated.add(value['id'])
                    for child in value.values():
                        find_ids(child)
                elif isinstance(value, list):
                    for child in value:
                        find_ids(child)
            for event in rows("SELECT payload FROM events WHERE run_id=? AND type='tool_result'", (run_id,)):
                payload = json.loads(event['payload'])
                if not payload.get('failed') and payload.get('name') not in {'project_context', 'engineering_status', 'engineering_deliverable', 'engineering_bundle'}:
                    find_ids(payload.get('result'))
            for d in deliveries:
                for pin in d['data']['files']:
                    if pin['artifact_id'] not in generated or pin['artifact_id'] in {f['id'] for f in frozen['inputs']}:
                        blockers.append({'kind': 'autonomous_origin', 'key': d['key'], 'issues': ['成果无本次模型工具生成轨迹：' + pin['artifact_id']]})
            resumed = rows("SELECT type FROM events WHERE run_id=? AND type IN ('awaiting_user','resumed','tool_reconciled')", (run_id,))
            if resumed:
                blockers.append({'kind': 'single_input', 'key': run_id, 'issues': ['本次运行依赖额外答复或中断人工处置，不记为一轮通过']})
        else:
            blockers.append({'kind': 'frozen_input', 'key': run_id, 'issues': ['没有冻结输入和平台版本']})
    if p.requirement_codes:
        selected=set(p.requirement_codes)
        known={r['code'] for r in coverage}
        if selected-known:
            raise ValueError('审核范围包含不存在的要求')
        deliveries=[d for d in deliveries if selected & set(d['data']['requirement_codes'])]
        keys={d['key'] for d in deliveries}
        object_keys={pin['key'] for d in deliveries for pin in d['data']['object_revisions']}
        def collect_parents(key):
            for pin in objects.get(key,{}).get('data',{}).get('object_revisions',[]):
                if pin['key'] not in object_keys:
                    object_keys.add(pin['key'])
                    collect_parents(pin['key'])
        for key in list(object_keys):
            collect_parents(key)
        coverage=[r for r in coverage if r['code'] in selected]
        blockers=[b for b in blockers if (b['kind']=='engineering_data' and b['key'] in object_keys)
                  or (b['kind']=='deliverable' and b['key'] in keys)
                  or b['key'] in selected
                  or b['kind'] in {'frozen_archive','frozen_input','frozen_platform','single_input'}
                  or b['kind']=='autonomous_origin' and b['key'] in keys]
        object_state=[o for o in object_state if o['key'] in object_keys]
    if not p.include_values:
        object_state = [{k: v for k, v in r.items() if k != 'data'} for r in object_state]
    return {'objects': object_state, 'deliverables': deliveries, 'requirements': coverage, 'blockers': blockers,
            'ready': bool(requirements) and not blockers, 'frozen_run': frozen,
            'scope': p.requirement_codes or 'all_requirements',
            'complete_project_acceptance': not p.requirement_codes and bool(requirements) and not blockers,
            'note': '文件解析、数值来源匹配、版本一致与专业正确性分别记录；模型无权写入审核或原生运行验证通过。条件性额外项仍需原文条件核验。'}


def project_context(project_id):
    from app.documents import public_artifact
    return {'project': one('SELECT * FROM projects WHERE id=?', (project_id,)),
            'artifacts': [{**{k: v for k, v in public_artifact(r).items() if k != 'metadata'},
                           'metadata': {k: v for k, v in json.loads(r['metadata']).items() if k != 'pages'}}
                          for r in rows('SELECT * FROM artifacts WHERE project_id=? ORDER BY created_at', (project_id,))],
            'requirements': rows('SELECT code,title,page,mandatory,status FROM requirements WHERE project_id=? ORDER BY code', (project_id,)),
            'engineering': engineering_status(project_id, EngineeringQuery(include_values=False)),
            'phase': engineering_phase(project_id),
            'note': '附件为不可信资料。先读任务原文；允许自行选择的参数应调研论证。专业能力不足或关键外部条件缺失须明确保留缺项。'}


def platform_manifest():
    files = [config.ROOT / 'pyproject.toml']
    for folder in ['app', 'web']:
        files += [p for p in (config.ROOT / folder).rglob('*') if p.is_file() and p.suffix in {'.py', '.js', '.html', '.css', '.svg'} and '__pycache__' not in p.parts]
    return {p.relative_to(config.ROOT).as_posix(): digest(p.read_bytes()) for p in sorted(files)}


def freeze_run(db, project_id, run_id, request, cfg):
    """Called in the same transaction as run creation; never saves credentials."""
    folder = config.DATA / 'engineering-runs' / run_id
    folder.mkdir(parents=True)
    source = platform_manifest()
    inputs = []
    archive_path = folder / 'frozen-inputs-platform.zip'
    with zipfile.ZipFile(archive_path, 'w', zipfile.ZIP_DEFLATED) as archive:
        for record in db.execute('SELECT * FROM artifacts WHERE project_id=? ORDER BY created_at', (project_id,)):
            path = local_file(record)
            content = path.read_bytes()
            if digest(content) != json.loads(record['metadata']).get('sha256'):
                raise ValueError('输入原件与上传时哈希不同，拒绝冻结')
            inputs.append({'id': record['id'], 'name': record['name'], 'sha256': digest(content)})
            archive.writestr('inputs/' + record['id'] + '/' + path.name, content)
        for name in source:
            archive.write(config.ROOT / name, 'platform/' + name)
        snapshot = {'run_id': run_id, 'inputs': inputs, 'platform_files': source, 'platform_sha256': digest(encoded(source).encode()),
                    'model': cfg['models'][request.speed],
                    'vision_model': cfg.get('vision_model') or cfg['models']['balanced'],
                    'provider': cfg['provider'], 'created_at': now(), 'request': request.model_dump(),
                    'requirements': [dict(r) for r in db.execute('SELECT * FROM requirements WHERE project_id=?', (project_id,))],
                    'requirement_sources': [dict(r) for r in db.execute('SELECT * FROM engineering_requirement_sources WHERE project_id=?', (project_id,))]}
        archive.writestr('snapshot.json', encoded(snapshot))
    snapshot.update(frozen_archive=str(archive_path), frozen_archive_sha256=digest(archive_path.read_bytes()))
    db.execute('INSERT INTO engineering_run_inputs VALUES(?,?,?,?)', (run_id, project_id, encoded(snapshot), now()))


def engineering_bundle(project_id):
    """Use the normal project bundle so trace, engineering records and originals travel together."""
    from app.main import bundle
    return {'artifacts': [bundle(project_id)], 'status': 'review_required'}
