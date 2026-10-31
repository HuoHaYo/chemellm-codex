"""Shared originals, immutable native extracts and session-authored reference notes.

No model API is called here. The calling Codex session reads text or rendered
pages and saves its own summaries/transcriptions, always linked to source blocks.
"""
from __future__ import annotations

from collections import Counter
from contextlib import contextmanager
import hashlib
import json
import math
import os
from pathlib import Path
import re
import shutil
import tempfile
import threading
import time
from typing import Literal

from pydantic import Field, model_validator

from app import config
from app.db import now, one, execute
from app.documents import ALLOWED, safe_name, validate_office
from app.pricing import _atomic
from app.schemas import Inputs

_lock = threading.RLock()


class ReferenceConflict(ValueError):
    pass


class ReferenceSource(Inputs):
    title: str = Field(min_length=1, max_length=500)
    source: str = Field(min_length=1, max_length=4000, description='用户原话、提供者、文件出处或网页来源；不详时如实注明待补')
    standard_number: str = Field(default='', max_length=200)
    edition: str = Field(default='', max_length=200)
    provider: str = Field(default='', max_length=500)
    url: str = Field(default='', max_length=4000)
    notes: str = Field(default='', max_length=12000, description='可编辑的来源和适用背景备注；截图/节选注明不完整，不自行认定标准现行')

    @model_validator(mode='after')
    def nonblank(self):
        if not self.title.strip() or not self.source.strip():
            raise ValueError('资料名称和来源不能是空白')
        return self


class ReferenceImport(Inputs):
    path: str | None = Field(default=None, max_length=4000)
    text: str | None = Field(default=None, min_length=1)
    url: str | None = Field(default=None, max_length=4000)
    name: str | None = Field(default=None, max_length=200)
    source: ReferenceSource

    @model_validator(mode='after')
    def one_input(self):
        if sum(value is not None for value in (self.path, self.text, self.url)) != 1:
            raise ValueError('path、text、url 必须且只能提供一个')
        return self


class ReferenceEntry(Inputs):
    key: str = Field(pattern=r'^[A-Za-z0-9][A-Za-z0-9_.-]{0,47}$')
    kind: Literal['summary', 'transcription', 'table', 'formula', 'figure'] = 'summary'
    text: str = Field(min_length=1, max_length=16000)
    source_block_ids: list[str] = Field(min_length=1, max_length=100)
    cells: list[list[str]] = Field(default_factory=list, max_length=500)
    conditions: str = Field(default='', max_length=4000)
    notes: str = Field(default='', max_length=4000)


class ReferenceContent(Inputs):
    overview: str = Field(default='', max_length=16000)
    entries: list[ReferenceEntry] = Field(default_factory=list, max_length=10000)

    @model_validator(mode='after')
    def unique(self):
        if len({entry.key for entry in self.entries}) != len(self.entries):
            raise ValueError('整理条目 key 重复')
        return self


class ReferenceSave(Inputs):
    expected_revision: int = Field(ge=1)
    reason: str = Field(min_length=1, max_length=3000)
    source: ReferenceSource | None = None
    overview: str | None = Field(default=None, max_length=16000)
    entries: list[ReferenceEntry] = Field(default_factory=list, max_length=500)
    remove_keys: list[str] = Field(default_factory=list, max_length=500)
    brief: str | None = Field(default=None, max_length=200000, description='可选 Markdown 精简稿；省略则由当前结构化条目生成')


class ReferenceUse(Inputs):
    reference_id: str = Field(pattern=r'^[a-f0-9]{64}$')
    revision: int = Field(ge=1)
    block_ids: list[str] = Field(min_length=1, max_length=100, description='原文 b... 段落，或已保存的逐字转录 t-<条目key>；精简改写不能作为原文证据')
    purpose: str = Field(min_length=1, max_length=3000)


def _dump(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False)


def _hash(value: bytes):
    return hashlib.sha256(value).hexdigest()


def _atomic_bytes(path, content):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, delete=False) as output:
            temporary = Path(output.name)
            output.write(content)
            output.flush()
            os.fsync(output.fileno())
        os.replace(temporary, path)
    finally:
        if temporary and temporary.exists():
            temporary.unlink()


def _tokens(text):
    english = re.findall(r'[a-z][a-z0-9_.-]*|\d+(?:\.\d+)?', text.lower())
    chinese = re.findall(r'[\u3400-\u9fff]+', text)
    return english + [word[i:i+2] for word in chinese for i in range(max(1, len(word)-1))]


def native_extract(path: Path):
    """Extract all supported native content without model calls or silent clipping."""
    from app.fulltext import _pdf, ExtractInput, extract_jats
    from app.public_sources import extract_html
    blocks, warnings, total_pages = [], [], None
    kind = path.suffix.lower()

    def add(text, locator, block_kind='paragraph', **extra):
        if text and str(text).strip():
            text = str(text).strip()
            for offset in range(0, len(text), 8000):
                blocks.append({'kind': block_kind, 'text': text[offset:offset+8000],
                               'locator': {**locator, 'text_offset': offset}, 'ocr': False,
                               **{key: value for key, value in extra.items() if key != 'cells' or offset == 0}})

    if kind == '.pdf':
        from pypdf import PdfReader
        total_pages = len(PdfReader(path).pages)
        if total_pages > 600:
            raise ValueError('PDF 页数超过600页限制，原件仍保留')
        for first in range(1, total_pages+1, 50):
            extracted = _pdf(path, ExtractInput(artifact_id='shared', first_page=first,
                                              last_page=min(first+49, total_pages)))
            warnings.extend(extracted['warnings'])
            for block in extracted['blocks']:
                add(block['text'], block['locator'], block['kind'],
                    **{k: v for k, v in block.items() if k not in {'id', 'text', 'kind', 'locator', 'ocr'}})
        # A page anchor exists even when native text is present: mixed pages and
        # vector diagrams must remain available to the current session visually.
        for page in range(1, total_pages+1):
            blocks.append({'kind': 'page_image', 'text': '', 'locator': {'page': page}, 'ocr': False})
    elif kind in {'.png', '.jpg', '.jpeg'}:
        from PIL import Image
        with Image.open(path) as image:
            image.verify()
        blocks.append({'kind': 'image', 'text': '', 'locator': {'page': 1}, 'ocr': False})
        total_pages = 1
    elif kind == '.docx':
        from docx import Document
        validate_office(path)
        document = Document(path)
        for number, paragraph in enumerate(document.paragraphs, 1):
            add(paragraph.text, {'paragraph': number, 'page': None})
        for number, table in enumerate(document.tables, 1):
            cells = [[cell.text for cell in row.cells] for row in table.rows]
            add('\n'.join('\t'.join(row) for row in cells), {'table': number, 'page': None}, 'table', cells=cells)
        warnings.append('DOCX 使用段落/表格定位，没有可靠出版版页码；内嵌图片可通过版面预览检查')
    elif kind == '.pptx':
        from pptx import Presentation
        validate_office(path)
        slides = Presentation(path).slides
        total_pages = len(slides)
        for page, slide in enumerate(slides, 1):
            for shape in slide.shapes:
                if shape.has_text_frame:
                    add(shape.text, {'page': page}, 'slide_text')
            blocks.append({'kind': 'page_image', 'text': '', 'locator': {'page': page}, 'ocr': False})
    elif kind in {'.xlsx', '.csv'}:
        if kind == '.xlsx':
            validate_office(path)
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            try:
                for sheet in workbook:
                    if (sheet.max_row or 0) > 5000 or (sheet.max_column or 0) > 100:
                        raise ValueError('原生表格提取最多5000行、100列；原件完整保留')
                    for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        if number > 5000 or len(row) > 100:
                            raise ValueError('原生表格提取最多5000行、100列；原件完整保留')
                        add('\t'.join('' if cell is None else str(cell) for cell in row),
                            {'sheet': sheet.title, 'row': number, 'page': None}, 'table_row')
            finally:
                workbook.close()
        else:
            import csv
            with path.open(encoding='utf-8-sig', newline='') as source:
                for number, row in enumerate(csv.reader(source), 1):
                    if number > 5000 or len(row) > 100:
                        raise ValueError('原生表格提取最多5000行、100列；原件完整保留')
                    add('\t'.join(row), {'row': number, 'page': None}, 'table_row', cells=[row])
    elif kind == '.xml':
        content = extract_jats(path.read_bytes())
        warnings.extend(content['warnings'])
        blocks.extend(content['blocks'])
    elif kind == '.html':
        content = extract_html(path.read_bytes())
        warnings.extend(content['warnings'])
        blocks.extend(content['blocks'])
    elif kind in {'.svg', '.dxf'}:
        from app.drawing_io import svg_summary, dxf_summary
        # dxf_summary writes only to the caller's staging directory.
        result = svg_summary(path.read_bytes()) if kind == '.svg' else dxf_summary(path, path.parent/'drawing-preview.svg')
        for number, item in enumerate(result['text_elements'], 1):
            add(item['text'], {'element': number, 'page': None}, 'drawing_text')
        warnings.append('矢量文字不证明连接关系或尺寸语义，图形内容须核对原件')
    else:
        text = path.read_text(encoding='utf-8-sig')
        position = 0
        for paragraph in re.split(r'(\n\s*\n)', text):
            add(paragraph, {'char_offset': position, 'page': None})
            position += len(paragraph)
    # Normalize foreign extractors and long HTML/JATS blocks to bounded chunks.
    normalized = []
    for block in blocks:
        text = block['text']
        for offset in range(0, max(1, len(text)), 8000):
            normalized.append({**block, 'id': 'b'+str(len(normalized)+1).zfill(6),
                               'text': text[offset:offset+8000],
                               'locator': {**block['locator'], **({'chunk_offset': offset} if offset else {})}})
            if offset:
                normalized[-1].pop('cells', None)
    return {'blocks': normalized, 'total_pages': total_pages, 'warnings': list(dict.fromkeys(warnings))}


class ReferenceLibrary:
    def __init__(self, data_root: Path, workspace: Path):
        self.root = (data_root / 'references').resolve()
        self.allowed = [workspace.resolve(), data_root.resolve()]
        self._local = threading.local()

    @contextmanager
    def _guard(self):
        """Serialize shared-library writers across MCP processes and threads."""
        with _lock:
            if getattr(self._local, 'locked', False):
                yield
                return
            self.root.mkdir(parents=True, exist_ok=True)
            with (self.root/'.lock').open('a+b') as handle:
                if handle.tell() == 0:
                    handle.write(b'0')
                    handle.flush()
                deadline = time.monotonic()+15
                while True:
                    try:
                        handle.seek(0)
                        if os.name == 'nt':
                            import msvcrt
                            msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
                        else:
                            import fcntl
                            fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
                        break
                    except OSError:
                        if time.monotonic() >= deadline:
                            raise ReferenceConflict('资料库正被其他进程使用，请稍后重新读取')
                        time.sleep(.05)
                self._local.locked = True
                try:
                    yield
                finally:
                    self._local.locked = False
                    handle.seek(0)
                    if os.name == 'nt':
                        import msvcrt
                        msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
                    else:
                        import fcntl
                        fcntl.flock(handle.fileno(), fcntl.LOCK_UN)

    def folder(self, reference_id):
        if not re.fullmatch(r'[a-f0-9]{64}', reference_id):
            raise ValueError('无效资料标识符')
        folder = (self.root / reference_id).resolve()
        if folder.parent != self.root:
            raise ValueError('资料路径越界')
        return folder

    def _meta(self, reference_id):
        folder = self.folder(reference_id)
        if not (folder/'manifest.json').is_file():
            raise ValueError('资料不存在')
        meta = json.loads((folder/'manifest.json').read_text(encoding='utf-8'))
        original = (folder / meta['original']).resolve()
        if not original.is_relative_to(folder) or not original.is_file() or _hash(original.read_bytes()) != reference_id:
            raise ValueError('资料原件缺失或哈希变化；请保留文件并恢复原件')
        native = (folder/'native.json').read_bytes()
        if _hash(native) != meta['native_sha256']:
            raise ValueError('原文提取结果哈希变化；不能用编辑稿覆盖原文层')
        return meta, json.loads(native)

    def _revision(self, folder, revision):
        if not isinstance(revision, int) or revision < 1:
            raise ValueError('资料版本须为正整数')
        record = json.loads((folder/'revisions'/f'{revision:06d}.json').read_text(encoding='utf-8'))
        digest = record.pop('sha256')
        if _hash(_dump(record).encode()) != digest or record['revision'] != revision:
            raise ValueError('资料历史版本校验失败')
        return record

    def _editable(self, folder, document):
        return {'source.json': (_dump(document['source'])+'\n').encode(),
                'organized.json': (_dump(document['content'])+'\n').encode(),
                'brief.md': document['brief'].encode('utf-8')}

    def _hashes(self, folder):
        return {name: _hash((folder/name).read_bytes()) if (folder/name).exists() else None
                for name in ('source.json', 'organized.json', 'brief.md')}

    def _recover(self, folder):
        pending_path = folder/'.pending.json'
        if not pending_path.exists():
            return
        pending = json.loads(pending_path.read_text(encoding='utf-8'))
        record = pending['record']
        target = folder/'revisions'/f"{record['revision']:06d}.json"
        if target.exists():
            existing = json.loads(target.read_text(encoding='utf-8'))
            if existing != record:
                raise ReferenceConflict('资料提交日志与历史不一致')
        else:
            _atomic(target, record)
        head_path = folder/'head.json'
        previous = json.loads(head_path.read_text(encoding='utf-8'))['revision'] if head_path.exists() else 0
        if previous not in {record['revision']-1, record['revision']}:
            raise ReferenceConflict('资料提交日志与当前版本不一致')
        _atomic(head_path, {'revision': record['revision']})
        actual = self._hashes(folder)
        for name, content in self._editable(folder, record['document']).items():
            if actual[name] in {None, pending['before'][name], _hash(content)}:
                _atomic_bytes(folder/name, content)
            # Preserve edits made during an interrupted commit; _sync records
            # them as a later revision instead of replacing them on recovery.
        pending_path.unlink()

    def _validate(self, document, native):
        document['source'] = ReferenceSource.model_validate(document['source']).model_dump()
        document['content'] = ReferenceContent.model_validate(document['content']).model_dump()
        if len(_dump(document['content']).encode()) > config.MAX_UPLOAD:
            raise ValueError('结构化整理稿超过25 MiB，请精简或拆分资料')
        if not isinstance(document['brief'], str) or len(document['brief']) > 200000:
            raise ValueError('Markdown 精简稿超过200000字符或格式无效')
        ids = {block['id'] for block in native['blocks']}
        for entry in document['content']['entries']:
            if set(entry['source_block_ids']) - ids:
                raise ValueError('整理条目引用了不存在的原文段落：'+entry['key'])
            if entry['kind'] == 'transcription' and any(
                    block['kind'] not in {'image', 'page_image'} for block in native['blocks'] if block['id'] in entry['source_block_ids']):
                raise ValueError('图像逐字转录必须绑定图像/页面锚点；原生文字应直接引用原文段落')
            if entry['kind'] == 'transcription' and len(entry['source_block_ids']) != 1:
                raise ValueError('逐字转录每条绑定一个图像/页面，以免引用位置混淆；多页请分条保存')
        return document

    def _commit(self, folder, revision, document, actor, reason, before):
        if self._hashes(folder) != before:
            raise ReferenceConflict('资料文件已被编辑，请重新读取并合并更新')
        record = {'revision': revision+1, 'recorded_at': now(), 'actor': actor,
                  'reason': reason, 'document': document}
        record['sha256'] = _hash(_dump(record).encode())
        _atomic(folder/'.pending.json', {'record': record, 'before': before})
        self._recover(folder)

    def _sync(self, reference_id):
        folder = self.folder(reference_id)
        self._recover(folder)
        meta, native = self._meta(reference_id)
        for _ in range(4):
            revision = json.loads((folder/'head.json').read_text(encoding='utf-8'))['revision']
            document = self._revision(folder, revision)['document']
            before = self._hashes(folder)
            candidate = dict(document)
            for name, key in (('source.json', 'source'), ('organized.json', 'content')):
                if (folder/name).exists():
                    candidate[key] = json.loads((folder/name).read_text(encoding='utf-8-sig'))
            if (folder/'brief.md').exists():
                candidate['brief'] = (folder/'brief.md').read_text(encoding='utf-8-sig')
            # Normalize optional fields before rendering user-edited JSON.
            # Handwritten files need not repeat every empty source/entry field.
            self._validate(candidate, native)
            if candidate['brief'] != document['brief']:
                candidate['brief_custom'] = True
            elif not document.get('brief_custom', False) and (
                    candidate['source'] != document['source'] or candidate['content'] != document['content']):
                candidate['brief'] = self._brief(candidate['source'], candidate['content'])
            self._validate(candidate, native)
            if candidate == document:
                if self._hashes(folder) != before:
                    continue
                # Missing editable copies can be reconstructed from history.
                expected = dict(before)
                for name, content in self._editable(folder, document).items():
                    if before[name] is None and not (folder/name).exists():
                        _atomic_bytes(folder/name, content)
                        expected[name] = _hash(content)
                if self._hashes(folder) != expected:
                    continue
                return meta, native, revision, document, expected
            self._commit(folder, revision, candidate, 'local_edit', '用户直接编辑来源备注、结构化整理稿或 Markdown 精简稿', before)
        raise ReferenceConflict('资料文件正在连续编辑，请保存后重新读取')

    def _brief(self, source, content):
        lines = ['# '+source['title'], '', '来源：'+source['source'],
                 '备注：'+source['notes'], '', content['overview'], '']
        for key, label in (('standard_number', '标准号'), ('edition', '版本'), ('provider', '提供者'), ('url', '来源链接')):
            if source[key]:
                lines.extend([label+'：'+source[key], ''])
        for entry in content['entries']:
            lines.extend(['## '+entry['key'], '', entry['text'], ''])
            if entry['cells']:
                width = max(len(row) for row in entry['cells'])
                if width:
                    for number, row in enumerate(entry['cells']):
                        cells = [str(value).replace('|', '\\|').replace('\n', '<br>') for value in row]
                        lines.append('| '+' | '.join(cells+['']*(width-len(cells)))+' |')
                        if number == 0:
                            lines.append('| '+' | '.join(['---']*width)+' |')
                    lines.append('')
            lines.extend(['原文定位：'+', '.join(entry['source_block_ids']),
                          '适用条件：'+entry['conditions'], '备注：'+entry['notes'], ''])
            if entry['kind'] == 'transcription':
                lines.extend(['图片逐字转录，待对照原图核对。', ''])
        return '\n'.join(lines)

    def _view(self, reference_id, meta, native, revision, document):
        covered = {anchor for entry in document['content']['entries'] for anchor in entry['source_block_ids']}
        pending = [block['id'] for block in native['blocks'] if block['id'] not in covered]
        folder = self.folder(reference_id)
        return {'reference_id': reference_id, 'revision': revision, 'source': document['source'],
                'created_at': meta['created_at'], 'source_sha256': reference_id,
                'extraction_status': meta['extraction_status'], 'warnings': meta['warnings'],
                'acquisitions': document['acquisitions'][-20:], 'acquisition_count': len(document['acquisitions']),
                'progress': {'total_blocks': len(native['blocks']), 'organized_blocks': len(covered),
                             'pending_count': len(pending), 'pending_block_ids': pending[:100],
                             'total_pages': native['total_pages']},
                'files': {'original': str((folder/meta['original']).resolve()),
                          **{key: str(folder/name) for key, name in [('source', 'source.json'),
                             ('native', 'native.json'), ('organized', 'organized.json'), ('brief', 'brief.md')]}}}

    def import_reference(self, request: ReferenceImport):
        from app.public_sources import fetch_public
        with self._guard():
            acquisition = {'source': request.source.model_dump(), 'retrieved_at': now()}
            if request.path is not None:
                path = Path(request.path).expanduser().resolve()
                if not path.is_file() or not any(path.is_relative_to(root) for root in self.allowed):
                    raise ValueError('导入文件必须位于本仓库或当前数据目录内')
                if path.stat().st_size > config.MAX_UPLOAD:
                    raise ValueError('文件不得超过25 MiB')
                name, content = request.name or path.name, path.read_bytes()
                acquisition['path'] = str(path)
            elif request.text is not None:
                name, content = request.name or '用户提供的原始文本.txt', request.text.encode('utf-8')
                acquisition['input'] = 'user_text'
            else:
                fetched = fetch_public(request.url)
                content = fetched['content']
                content_type = fetched['headers'].get('content-type', '').lower()
                suffix = '.pdf' if content.startswith(b'%PDF-') else '.html' if 'html' in content_type else '.txt' if content_type.startswith('text/plain') else None
                if suffix is None:
                    raise ValueError('公开网址仅接受原始 PDF、HTML 或 UTF-8 文本')
                name = '网页原件'+suffix
                acquisition.update({k: v for k, v in fetched.items() if k != 'content'})
                acquisition['requested_url'] = request.url
                if not request.source.url:
                    request.source.url = fetched['final_url']
                    acquisition['source'] = request.source.model_dump()
            name = safe_name(name)
            if Path(name).suffix.lower() not in ALLOWED:
                raise ValueError('不支持的原件类型')
            if len(content) > config.MAX_UPLOAD:
                raise ValueError('原件不得超过25 MiB')
            reference_id = _hash(content)
            folder = self.folder(reference_id)
            if folder.exists():
                meta, native, revision, document, before = self._sync(reference_id)
                # Identical repeated imports do not create revision noise; a new
                # provider or URL is retained as an additional acquisition.
                comparable = {k: v for k, v in acquisition.items() if k != 'retrieved_at'}
                if not any({k: v for k, v in old.items() if k != 'retrieved_at'} == comparable for old in document['acquisitions']):
                    document['acquisitions'].append(acquisition)
                    self._commit(folder, revision, document, 'import', '相同原件补充来源记录', before)
                result = self.read(reference_id, view='metadata')
                result['duplicate'] = True
                return result
            self.root.mkdir(parents=True, exist_ok=True)
            with tempfile.TemporaryDirectory(prefix='.import-', dir=self.root) as temporary:
                stage = Path(temporary)/reference_id
                original = stage/'original'/name
                _atomic_bytes(original, content)
                try:
                    native = native_extract(original)
                    extraction_status = 'extracted' if any(block['text'] for block in native['blocks']) else 'awaiting_session'
                except Exception as exc:
                    native = {'blocks': [], 'total_pages': None,
                              'warnings': ['原生提取失败，原件已保留：'+type(exc).__name__+': '+str(exc)[:500]]}
                    extraction_status = 'unavailable'
                _atomic(stage/'native.json', native)
                meta = {'reference_id': reference_id, 'original': 'original/'+name,
                        'original_name': name, 'created_at': now(), 'bytes': len(content),
                        'native_sha256': _hash((stage/'native.json').read_bytes()),
                        'extraction_status': extraction_status, 'warnings': native['warnings']}
                _atomic(stage/'manifest.json', meta)
                source = request.source.model_dump()
                organized = ReferenceContent().model_dump()
                document = {'source': source, 'content': organized,
                            'brief': self._brief(source, organized), 'acquisitions': [acquisition]}
                self._commit(stage, 0, document, 'import', '保存原件、来源与原文提取结果', self._hashes(stage))
                os.replace(stage, folder)
            result = self.read(reference_id, view='metadata')
            result['duplicate'] = False
            return result

    def save(self, reference_id, request: ReferenceSave):
        with self._guard():
            meta, native, revision, document, before = self._sync(reference_id)
            if revision != request.expected_revision:
                raise ReferenceConflict(f'资料版本冲突：当前为 {revision}，请重新读取并合并更新')
            if len({entry.key for entry in request.entries}) != len(request.entries):
                raise ValueError('同批次整理条目 key 重复')
            if request.source is not None:
                document['source'] = request.source.model_dump()
            content = document['content']
            if request.overview is not None:
                content['overview'] = request.overview
            entries = {entry['key']: entry for entry in content['entries'] if entry['key'] not in request.remove_keys}
            entries.update({entry.key: entry.model_dump() for entry in request.entries})
            content['entries'] = list(entries.values())
            if request.brief is not None:
                document['brief'] = request.brief
                document['brief_custom'] = True
            elif not document.get('brief_custom', False):
                document['brief'] = self._brief(document['source'], content)
            self._validate(document, native)
            old = self._revision(self.folder(reference_id), revision)['document']
            if old != document:
                self._commit(self.folder(reference_id), revision, document, 'codex_session', request.reason, before)
            return self.read(reference_id, view='metadata')

    def _state(self, reference_id, revision=None):
        if revision is None:
            return self._sync(reference_id)[:4]
        folder = self.folder(reference_id)
        self._recover(folder)
        meta, native = self._meta(reference_id)
        return meta, native, revision, self._revision(folder, revision)['document']

    def _render(self, reference_id, meta, page):
        from PIL import Image
        folder = self.folder(reference_id)
        original = folder/meta['original']
        suffix = original.suffix.lower()
        target = folder/'previews'/f'page-{page:04d}.png'
        if target.exists():
            return target
        target.parent.mkdir(parents=True, exist_ok=True)
        if suffix in {'.png', '.jpg', '.jpeg'}:
            if page != 1:
                raise ValueError('图片仅有第1页')
            with Image.open(original) as source:
                image = source.convert('RGB')
        else:
            if suffix in {'.docx', '.pptx', '.xlsx'}:
                from app import documents
                cached_source = target.parent/original.name
                if not cached_source.exists():
                    shutil.copyfile(original, cached_source)
                pdf = documents.office_pdf(cached_source)
                if not pdf:
                    raise ValueError('当前 Office 版面渲染不可用，原件和原生文本仍保留')
            elif suffix == '.pdf':
                pdf = original
            else:
                raise ValueError('图像视图支持 PDF、PNG/JPEG 和可用的 Office 版面预览')
            import pypdfium2 as pdfium
            document = pdfium.PdfDocument(pdf)
            try:
                if not 1 <= page <= len(document):
                    raise ValueError('图像页码超出范围')
                rendered = document[page-1]
                try:
                    scale = min(2, 1800/max(rendered.get_size()))
                    bitmap = rendered.render(scale=scale)
                    try:
                        image = bitmap.to_pil().convert('RGB')
                    finally:
                        bitmap.close()
                finally:
                    rendered.close()
            finally:
                document.close()
        try:
            image.thumbnail((1800, 1800))
            import io
            buffer = io.BytesIO()
            image.save(buffer, format='PNG')
            _atomic_bytes(target, buffer.getvalue())
        finally:
            image.close()
        return target

    def read(self, reference_id, *, revision=None, view='source', page=None,
             block_ids=None, offset=0, limit=10):
        with self._guard():
            if view not in {'metadata', 'source', 'organized', 'image'} or offset < 0 or not 1 <= limit <= 50:
                raise ValueError('视图无效，或 offset/limit 超出范围')
            if page is not None and (not isinstance(page, int) or page < 1):
                raise ValueError('页码须为正整数')
            meta, native, revision, document = self._state(reference_id, revision)
            result = self._view(reference_id, meta, native, revision, document)
            if view == 'metadata':
                return result
            if view == 'image':
                page = page or 1
                path = self._render(reference_id, meta, page)
                result['image'] = {'localPath': str(path), 'mimeType': 'image/png',
                                   'page': page, 'sha256': _hash(path.read_bytes())}
                result['source_block_ids'] = [block['id'] for block in native['blocks']
                                             if block['locator'].get('page') == page and block['kind'] in {'image', 'page_image'}]
                result['note'] = '请当前 Codex 会话识读；无法辨认处标[不清]，逐字转录和精简说明分别保存'
                return result
            if view == 'source':
                items = native['blocks']
            else:
                items = [{**entry, 'id': ('t-' if entry['kind'] == 'transcription' else 's-')+entry['key'],
                          'needs_review': entry['kind'] == 'transcription'} for entry in document['content']['entries']]
                result['overview'] = document['content']['overview']
                # The complete human-editable Markdown is a local file; bounded
                # pages prevent a whole manual from flooding a tool response.
                result['brief_excerpt'] = document['brief'][:8000]
                result['brief_truncated'] = len(document['brief']) > 8000
            if page is not None:
                source_ids = {block['id'] for block in native['blocks'] if block['locator'].get('page') == page}
                items = [item for item in items if item['id'] in source_ids or source_ids.intersection(item.get('source_block_ids', []))]
            if block_ids:
                wanted = set(block_ids)
                selected = [item for item in items if item['id'] in wanted]
                if {item['id'] for item in selected} != wanted:
                    raise ValueError('段落或条目编号与当前资料/页码不匹配')
                items = selected
            result.update({'total_items': len(items), 'offset': offset, 'items': []})
            size = 0
            for item in items[offset:offset+limit]:
                exposed = dict(item)
                if len(_dump(exposed.get('cells', []))) > 8000:
                    exposed.pop('cells', None)
                    exposed['cells_note'] = '完整行列保存在 native/organized JSON 文件中；本次返回完整的分块文字'
                item_size = len(_dump(exposed))
                if result['items'] and size+item_size > 24000:
                    break
                result['items'].append(exposed)
                size += item_size
            result['next_offset'] = offset+len(result['items']) if offset+len(result['items']) < len(items) else None
            return result

    def search(self, query='', *, offset=0, limit=10):
        with self._guard():
            if offset < 0 or not 1 <= limit <= 50 or len(query) > 1000:
                raise ValueError('检索参数超出范围')
            terms, candidates, errors = set(_tokens(query)), [], []
            for folder in sorted(self.root.glob('[a-f0-9]'*64)):
                try:
                    meta, native, revision, document = self._sync(folder.name)[:4]
                    source = document['source']
                    header = ' '.join(str(value) for value in source.values())
                    index_text = header+'\n'+document['content']['overview']+'\n'+document['brief']
                    for item in ([{'id': 'metadata', 'text': index_text, 'locator': {}}] + native['blocks'] + [
                            {**entry, 'id': ('t-' if entry['kind'] == 'transcription' else 's-')+entry['key']}
                            for entry in document['content']['entries']]):
                        words = Counter(_tokens(item['text']+' '+header))
                        if terms and not terms.intersection(words):
                            continue
                        candidates.append((folder.name, revision, source, item, words))
                except (ValueError, OSError, KeyError) as exc:
                    errors.append({'reference_id': folder.name, 'error': str(exc)[:500]})
            frequencies = {term: sum(term in words for *_, words in candidates) for term in terms}
            average = sum(sum(words.values()) for *_, words in candidates)/max(1, len(candidates))
            found = []
            for reference_id, revision, source, item, words in candidates:
                score = sum(math.log(1+(len(candidates)-frequencies[term]+.5)/(frequencies[term]+.5)) *
                            words[term]*2.2/(words[term]+1.2*(.25+.75*sum(words.values())/max(average, 1)))
                            for term in terms if words[term])
                if not terms and item['id'] != 'metadata':
                    continue
                found.append({'reference_id': reference_id, 'revision': revision, 'source': source,
                              'block_id': item['id'], 'text': item['text'][:1600], 'score': score,
                              'locator': item.get('locator'), 'source_block_ids': item.get('source_block_ids', []),
                              'needs_review': item.get('kind') == 'transcription',
                              'kind': item.get('kind', 'metadata')})
            found.sort(key=lambda item: (-item['score'], item['reference_id'], item['block_id']))
            return {'results': found[offset:offset+limit], 'total': len(found), 'offset': offset,
                    'errors': errors, 'note': '相关性仅用于查找；引用前回看原文/图像，结合来源备注和当前任务判断'}

    def use(self, project_id, request: ReferenceUse):
        """Copy a pinned edition into the existing project artifact/evidence store."""
        from app.documents import artifact_record, public_artifact
        with self._guard():
            meta, native, revision, document = self._state(request.reference_id, request.revision)
            by_id = {block['id']: block for block in native['blocks']}
            for entry in document['content']['entries']:
                if entry['kind'] != 'transcription':
                    continue
                anchor = by_id[entry['source_block_ids'][0]]
                by_id['t-'+entry['key']] = {
                    'id': 't-'+entry['key'], 'kind': 'ai_image_transcription', 'text': entry['text'],
                    'locator': anchor['locator'], 'ocr': True, 'transcription_review_required': True,
                    'transcription_model': 'current_codex_session',
                    'transcription_prompt_version': 'reference-session-v1',
                    'transcription_sha256': _hash(entry['text'].encode()),
                    'source_block_ids': entry['source_block_ids']}
            if len(set(request.block_ids)) != len(request.block_ids):
                raise ValueError('引用段落不能重复')
            selected = []
            for block_id in sorted(request.block_ids):
                block = by_id.get(block_id)
                if not block or not block['text']:
                    raise ValueError('引用须选原文文字段落或 t- 前缀的逐字转录，精简稿/图像锚点不能冒充原文：'+block_id)
                selected.append(block)
            pin = {'reference_id': request.reference_id, 'revision': revision,
                   'block_ids': sorted(request.block_ids), 'purpose': request.purpose}
            use_id = _hash(_dump({'project_id': project_id, **pin}).encode())[:32]
            folder = config.ARTIFACTS/use_id
            folder.mkdir(parents=True, exist_ok=True)
            original = self.folder(request.reference_id)/meta['original']
            snapshot_content = {'blocks': selected, 'warnings': native['warnings'], 'shared_reference': pin}
            content = _dump(snapshot_content)
            snapshot_id = use_id
            source_metadata = {'shared_reference': pin, 'fulltext_source': {'source_url': document['source']['url']},
                               'source_notes': document['source'], 'retrieved_at': meta['created_at']}
            payloads = {'original/'+meta['original_name']: original.read_bytes(),
                        'source.json': (_dump(document['source'])+'\n').encode(),
                        'organized.json': (_dump(document['content'])+'\n').encode(),
                        'brief.md': document['brief'].encode(),
                        'native.json': content.encode(),
                        'citation.json': (_dump({**pin, 'snapshot_id': snapshot_id,
                                                'source_sha256': request.reference_id,
                                                'retrieved_at': meta['created_at'],
                                                'acquisitions': document['acquisitions']})+'\n').encode()}
            if any(len(data) > config.MAX_UPLOAD for data in payloads.values()):
                raise ValueError('引用副本超过25 MiB，请缩小选用片段或精简整理稿')
            artifacts, original_id = [], None
            for name, data in payloads.items():
                artifact_id = _hash((use_id+'/'+name).encode())[:32]
                path = folder/name
                if path.exists() and path.read_bytes() != data:
                    raise ValueError('项目引用副本内容已变化，不能静默覆盖：'+name)
                if not path.exists():
                    _atomic_bytes(path, data)
                row = one('SELECT * FROM artifacts WHERE id=? AND project_id=?', (artifact_id, project_id))
                if row:
                    artifact = public_artifact(row)
                else:
                    artifact = artifact_record(project_id, path, name=path.name, metadata=source_metadata,
                                               text='\n'.join(block['text'] for block in selected) if name.startswith('original/') else '',
                                               artifact_id=artifact_id)
                artifacts.append(artifact)
                if name.startswith('original/'):
                    original_id = artifact_id
            existing = one('SELECT * FROM document_snapshots WHERE id=? AND project_id=?', (snapshot_id, project_id))
            if not existing:
                execute('INSERT INTO document_snapshots VALUES(?,?,?,?,?,?,?,?)',
                        (snapshot_id, project_id, original_id, request.reference_id, _hash(content.encode()),
                         _dump({'shared_reference': pin}), content, now()))
            elif existing['content'] != content or existing['artifact_id'] != original_id:
                raise ValueError('项目引用快照与资料固定版本不一致')
            return {'use_id': use_id, **pin, 'snapshot_id': snapshot_id, 'artifact_id': original_id,
                    'source_sha256': request.reference_id, 'artifacts': artifacts,
                    'citations': [{'snapshot_id': snapshot_id, 'block_id': block['id'],
                                   'locator': block['locator'], 'ocr': block.get('ocr', False)} for block in selected]}
