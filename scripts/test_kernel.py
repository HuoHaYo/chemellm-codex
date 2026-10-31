"""Kernel regression checks in an isolated data directory; no running service needed.

Run: python -X utf8 scripts/test_kernel.py
"""
from __future__ import annotations

import hashlib
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
_temporary = tempfile.TemporaryDirectory(prefix="chemellm-kernel-tests-")
os.environ.update({
    "CHEMELLM_DATA_ROOT": _temporary.name,
    "CHEMELLM_KERNEL_DATA": str(Path(_temporary.name) / "kernel"),
    "CHEMELLM_DATA_DIR": str(Path(_temporary.name) / "kernel"),
    "CHEMELLM_SNAPSHOT_DIR": str(Path(_temporary.name) / ".chem"),
    "CHEMELLM_KERNEL_TOKEN": "regression-token",
    "CHEMELLM_PHASE_GATE": "enforce",
})

from fastapi.testclient import TestClient  # noqa: E402
from kernel.server import app  # noqa: E402
from app.pricing import PriceBook, PriceWrite  # noqa: E402


def price_item(key='equipment-E-101', **updates):
    return {
        'key': key, 'category': 'equipment', 'title': '测试换热器', 'equipment_tag': 'E-101',
        'model_spec': '试验面积 20 m2', 'material': '304',
        'design_pressure': {'value': 1.0, 'unit': 'MPa'},
        'design_temperature': {'value': 160, 'unit': 'degC'},
        'quantity': 2, 'quantity_unit': '台', 'unit_price': 100000, 'currency': 'CNY',
        'cost_basis': 'one_time', 'conditions': '测试专用数据；表压；同一供货范围',
        'tax': {'status': 'included', 'rate': 0.13}, 'freight': {'status': 'included'},
        'installation_scope': '不含现场安装', 'source': '测试用户原话：单台 10 万元；不是市场报价',
        'quote_date': '2026-01-01', 'valid_until': '2099-12-31',
        'origin': 'quotation', 'rationale': '回归测试资料', **updates,
    }


class KernelWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app, headers={"x-kernel-token": "regression-token"})
        created = self.client.post("/v1/projects", json={"name": "重构回归"})
        self.assertEqual(created.status_code, 200, created.text)
        self.base = "/v1/projects/" + created.json()["project"]["projectId"]

    def tearDown(self):
        self.client.delete(self.base)
        self.client.close()

    def post(self, suffix, payload):
        response = self.client.post(self.base + suffix, json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def test_token_and_executable_catalog(self):
        with TestClient(app) as anonymous:
            self.assertEqual(anonymous.get("/v1/health").status_code, 401)
        catalog = self.client.get("/v1/tools").json()
        self.assertIn("pid", catalog["executable"])
        self.assertNotIn("ask_user", catalog["executable"])
        self.assertNotIn("acceptance", catalog["executable"])

    def test_full_workflow_and_review_invalidation(self):
        for tool in ("properties", "export_xlsx", "export_freecad"):
            response = self.client.post(self.base + "/execute", json={"tool": tool, "args": {}})
            self.assertEqual(response.status_code, 409, response.text)
        self.post("/requirements", {"items": [{"code": "R01", "title": "物性计算与报告"}]})
        self.post("/plan", {"deliverables": [{
            "key": "report", "title": "计算报告", "requirementCodes": ["R01"], "formats": ["md"],
        }]})
        basis = {"kind": "basis", "key": "feed", "quantities": [{
            "key": "flow", "value": 100, "unit": "kmol/h", "origin": "assumption",
            "conditions": "298 K, 101325 Pa", "rationale": "回归验证用假设",
        }]}
        self.post("/objects", basis)
        calculation = self.post("/execute", {"tool": "properties", "args": {"compound": "water"}})
        self.assertTrue(calculation["calculationId"])
        report = self.post("/execute", {"tool": "export_report", "args": {
            "title": "回归计算报告", "body": "本文件仅用于验证确定性生成和成果打包。", "formats": ["md"],
        }})
        artifact = report["artifacts"][0]
        downloaded = self.client.get("/v1/artifacts/" + artifact["artifactId"] + "/download")
        self.assertEqual(downloaded.status_code, 200)
        review = self.post("/review", {"note": "refactor regression"})
        self.assertTrue(review["review"]["ready"], review)
        bundle = self.post("/bundle", {})
        response = self.client.get("/v1/bundles/" + bundle["bundle"]["name"])
        self.assertEqual(response.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            manifest = json.loads(archive.read("manifest.json"))
            for item in manifest["files"]:
                self.assertEqual(hashlib.sha256(archive.read(item["archivePath"])).hexdigest(), item["sha256"])
        basis["quantities"][0]["value"] = 120
        self.post("/objects", basis)
        project = self.client.get(self.base).json()["project"]
        self.assertEqual(project["review"], {})

    def prices(self):
        response = self.client.get(self.base + '/prices')
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()['prices']

    def save_prices(self, items, revision, **updates):
        return self.post('/prices', {'items': items, 'expected_revision': revision,
                                    'reason': '测试更新', **updates})['prices']

    def test_price_updates_manual_edits_and_history(self):
        self.assertEqual(self.prices()['revision'], 0)
        first = self.save_prices([price_item()], 0)
        second = self.save_prices([price_item(unit_price=120000, material='316L')], 1)
        self.assertEqual(second['revision'], 2)
        self.assertEqual(second['items'][0]['material'], '316L')
        unchanged = self.save_prices(second['items'], 2)
        self.assertEqual(unchanged['revision'], 2)
        path = Path(first['files']['current'])
        document = json.loads(path.read_text(encoding='utf-8'))
        document['items'][0]['unit_price'] = 115000
        path.write_text(json.dumps(document, ensure_ascii=False), encoding='utf-8-sig')
        project = self.client.get(self.base).json()['project']
        self.assertEqual(project['prices']['revision'], 3)
        self.assertEqual(project['prices']['items'][0]['unit_price'], 115000)
        events = self.client.get(self.base + '/prices/history', params={'key': 'equipment-E-101'}).json()['events']
        self.assertEqual(len(events), 3)
        self.assertEqual(events[0]['changes'][0]['before'], None)
        self.assertEqual(events[1]['changes'][0]['before']['unit_price'], 100000)
        self.assertEqual(events[1]['changes'][0]['after']['unit_price'], 120000)
        self.assertEqual(events[2]['actor'], 'local_edit')
        self.assertEqual(events[2]['changes'][0]['after']['unit_price'], 115000)
        # A new object reconstructs from disk, without relying on process memory.
        self.assertEqual(PriceBook(path.parent).read()['revision'], 3)
        other = self.client.post('/v1/projects', json={'name': '其他项目'}).json()['project']['projectId']
        try:
            self.assertEqual(self.client.get('/v1/projects/' + other + '/prices').json()['prices']['items'], [])
        finally:
            self.client.delete('/v1/projects/' + other)

    def test_missing_estimates_protect_real_prices_and_other_costs(self):
        self.save_prices([price_item()], 0)
        assumed = price_item('civil-building', category='civil', equipment_tag='', model_spec='',
                             material='', design_pressure=None, design_temperature=None,
                             origin='assumption', unit_price=500000, quantity=1, quantity_unit='项',
                             source='工程判断的暂估，非实际报价', rationale='测试暂估及不确定性，待现场方案确认',
                             estimate_low=400000, estimate_high=700000, assumed_fields=['quantity'])
        result = self.save_prices([price_item(origin='assumption', unit_price=1), assumed], 1, mode='fill_missing')
        self.assertEqual(result['revision'], 2)
        self.assertEqual(result['items'][0]['unit_price'], 100000)
        self.assertEqual([item['key'] for item in result['pending_review']], ['civil-building'])
        rejected = self.client.post(self.base + '/prices', json={
            'items': [price_item(origin='assumption', unit_price=1)], 'expected_revision': 2, 'reason': '非法覆盖',
        })
        self.assertEqual(rejected.status_code, 422)
        self.assertEqual(self.prices()['revision'], 2)
        confirmed = {**assumed, 'origin': 'quotation', 'unit_price': 600000,
                     'source': '测试用用户价格', 'assumed_fields': [], 'estimate_low': None, 'estimate_high': None}
        updated = self.save_prices([confirmed], 2)
        self.assertEqual(updated['revision'], 3)
        self.assertEqual(updated['pending_review'], [])

    def test_invalid_prices_and_stale_revision_preserve_user_file(self):
        self.save_prices([price_item()], 0)
        for updates in ({'unit_price': -1}, {'quantity': 0}, {'equipment_tag': ''},
                        {'design_pressure': {'value': 1, 'unit': 'degC'}},
                        {'valid_until': '2025-01-01'}, {'currency': '元'},
                        {'assumed_fields': ['unit_price']}, {'estimate_low': 200000, 'estimate_high': 300000}):
            response = self.client.post(self.base + '/prices', json={
                'items': [price_item(**updates)], 'expected_revision': 1, 'reason': '无效输入',
            })
            self.assertEqual(response.status_code, 422, response.text)
        paths = self.prices()['files']
        current = Path(paths['current'])
        history = Path(paths['history'])
        prior_history = history.read_bytes()
        current.write_text('{ broken user edit', encoding='utf-8')
        response = self.client.get(self.base + '/prices')
        self.assertEqual(response.status_code, 422, response.text)
        self.assertEqual(current.read_text(encoding='utf-8'), '{ broken user edit')
        self.assertEqual(history.read_bytes(), prior_history)
        current.unlink()
        restored = self.prices()
        self.assertEqual(restored['revision'], 1)
        doc = json.loads(current.read_text(encoding='utf-8'))
        doc['items'][0]['unit_price'] = 110000
        current.write_text(json.dumps(doc), encoding='utf-8')
        stale = self.client.post(self.base + '/prices', json={
            'items': [price_item(unit_price=150000)], 'expected_revision': 1, 'reason': '旧版本',
        })
        self.assertEqual(stale.status_code, 409, stale.text)
        self.assertEqual(self.prices()['items'][0]['unit_price'], 110000)

    def test_local_deletion_and_paginated_history(self):
        state = self.save_prices([price_item(), price_item('equipment-E-102', equipment_tag='E-102')], 0)
        path = Path(state['files']['current'])
        doc = json.loads(path.read_text(encoding='utf-8'))
        doc['items'] = doc['items'][1:]
        path.write_text(json.dumps(doc), encoding='utf-8')
        self.assertEqual(len(self.prices()['items']), 1)
        history = self.client.get(self.base + '/prices/history', params={'key': 'equipment-E-101', 'offset': 1, 'limit': 1}).json()
        self.assertEqual(history['total'], 2)
        self.assertEqual(history['events'][0]['changes'][0]['after'], None)
        self.assertIsNotNone(history['events'][0]['changes'][0]['before'])
        self.assertEqual(self.client.get(self.base + '/prices/history?limit=0').status_code, 422)

    def test_interrupted_price_publication_preserves_later_user_edit(self):
        import app.pricing as pricing
        folder = Path(self.prices()['files']['current']).parent
        book = PriceBook(folder)
        original_atomic = pricing._atomic

        def fail_current(path, value):
            if path == book.current:
                raise OSError('simulated interruption before publication')
            return original_atomic(path, value)

        with patch.object(pricing, '_atomic', side_effect=fail_current):
            with self.assertRaises(OSError):
                book.save(PriceWrite(items=[price_item()], expected_revision=0, reason='模拟中断'))
        self.assertTrue(book.pending.is_file())
        # Editor still has revision zero after history revision one was committed.
        book.current.write_text(json.dumps({'schema_version': 1, 'revision': 0,
                                           'items': [price_item(unit_price=90000)]}), encoding='utf-8')
        recovered = PriceBook(folder).read()
        self.assertEqual(recovered['revision'], 2)
        self.assertEqual(recovered['items'][0]['unit_price'], 90000)
        self.assertFalse(book.pending.exists())
        events = book.events()['events']
        self.assertEqual(events[1]['changes'][0]['before']['unit_price'], 100000)
        self.assertEqual(events[1]['actor'], 'local_edit')

    def test_price_save_racing_with_local_edit_retains_both_changes(self):
        import app.pricing as pricing
        state = self.save_prices([price_item()], 0)
        book = PriceBook(Path(state['files']['current']).parent)
        original_atomic = pricing._atomic
        edited = False

        def edit_before_journal(path, value):
            nonlocal edited
            if path == book.pending and not edited:
                edited = True
                document = json.loads(book.current.read_text(encoding='utf-8'))
                document['items'][0]['unit_price'] = 95000
                book.current.write_text(json.dumps(document), encoding='utf-8')
            return original_atomic(path, value)

        with patch.object(pricing, '_atomic', side_effect=edit_before_journal):
            saved = book.save(PriceWrite(items=[price_item(unit_price=120000)], expected_revision=1, reason='并发修改测试'))
        self.assertEqual(saved['items'][0]['unit_price'], 95000)
        events = book.events()['events']
        self.assertEqual(events[1]['changes'][0]['after']['unit_price'], 120000)
        self.assertEqual(events[2]['changes'][0]['after']['unit_price'], 95000)
        self.assertEqual(events[2]['actor'], 'local_edit')

    def test_price_review_invalidation_economics_and_bundle(self):
        self.save_prices([price_item(origin='assumption', quote_date=None, valid_until=None)], 0)
        self.post('/requirements', {'items': [{'code': 'R01', 'title': '价格测试报告'}]})
        self.post('/plan', {'deliverables': [{'key': 'report', 'title': '报告', 'requirementCodes': ['R01'], 'formats': ['md']}]})
        self.post('/objects', {'kind': 'basis', 'key': 'cost', 'quantities': [{
            'key': 'capital', 'value': 200000, 'unit': 'CNY', 'origin': 'assumption',
            'conditions': '测试', 'rationale': '回归测试暂估',
        }]})
        self.post('/execute', {'tool': 'economics', 'args': {
            'capital_CNY': 200000, 'annual_revenue_CNY': 50000, 'annual_cost_CNY': 10000,
        }})
        self.post('/execute', {'tool': 'export_report', 'args': {'title': '价格测试', 'body': '回归用报告', 'formats': ['md']}})
        reviewed = self.post('/review', {})['review']
        self.assertTrue(reviewed['ready'], reviewed)
        self.assertTrue(any('价格为假设' in warning for warning in reviewed['warnings']))
        bundle = self.post('/bundle', {})['bundle']
        archive_bytes = self.client.get('/v1/bundles/' + bundle['name']).content
        with zipfile.ZipFile(io.BytesIO(archive_bytes)) as archive:
            self.assertIn('prices/prices.json', archive.namelist())
            self.assertIn('prices/history.json', archive.namelist())
            for entry in json.loads(archive.read('manifest.json'))['files']:
                self.assertEqual(hashlib.sha256(archive.read(entry['archivePath'])).hexdigest(), entry['sha256'])
        self.save_prices([price_item(unit_price=130000)], 1)
        project = self.client.get(self.base).json()['project']
        self.assertEqual(project['review'], {})
        self.assertIsNone(project['bundle'])
        reviewed = self.post('/review', {})['review']
        self.assertTrue(any('价格台账已更新' in item for item in reviewed['blockers']))
        self.post('/execute', {'tool': 'economics', 'args': {
            'capital_CNY': 260000, 'annual_revenue_CNY': 50000, 'annual_cost_CNY': 10000,
        }})
        self.assertTrue(self.post('/review', {})['review']['ready'])


class ReferenceLibraryTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app, headers={'x-kernel-token': 'regression-token'})
        self.project_id = self.client.post('/v1/projects', json={'name': '资料库回归'}).json()['project']['projectId']
        self.source = {'title': '合成测试资料 '+self.project_id, 'source': '回归测试生成，非工程标准', 'notes': '测试适用范围'}

    def tearDown(self):
        self.client.delete('/v1/projects/'+self.project_id)
        self.client.close()

    def post(self, path, payload):
        result = self.client.post(path, json=payload)
        self.assertEqual(result.status_code, 200, result.text)
        return result.json()

    def imported(self, **inputs):
        return self.post('/v1/references/import', {'source': self.source, **inputs})['reference']

    def read(self, ident, **params):
        result = self.client.get('/v1/references/'+ident, params=params)
        self.assertEqual(result.status_code, 200, result.text)
        return result.json()['reference']

    def save(self, ident, revision, **fields):
        return self.post('/v1/references/'+ident, {'expected_revision': revision, 'reason': '整理回归', **fields})['reference']

    def test_original_dedup_edition_long_text_and_search(self):
        text = '测试资料：设计压力按用户给出的适用条件核对。\n\n'+('长段落保留单位和例外条件；'*2000)
        result = self.imported(text=text)
        ident = result['reference_id']
        self.assertEqual(Path(result['files']['original']).read_text(encoding='utf-8'), text)
        self.assertEqual(ident, hashlib.sha256(text.encode()).hexdigest())
        again = self.imported(text=text)
        self.assertTrue(again['duplicate'])
        self.assertEqual(again['revision'], 1)
        other = self.imported(text=text+'不同版本')
        self.assertNotEqual(ident, other['reference_id'])
        current, offset, restored = self.read(ident, limit=1), 0, []
        while True:
            page = self.read(ident, offset=offset, limit=1)
            restored.extend(block['text'] for block in page['items'])
            if page['next_offset'] is None:
                break
            offset = page['next_offset']
        self.assertEqual(''.join(restored), text.replace('\n\n', ''))
        self.assertGreater(len(restored), 2)
        key = current['items'][0]['id']
        saved = self.save(ident, 1, overview='设计压力资料索引', entries=[{
            'key': 'pressure', 'text': '压力及适用条件须核对。', 'source_block_ids': [key]}])
        self.assertGreater(saved['progress']['pending_count'], 0)
        search = self.client.get('/v1/references', params={'query': '设计压力'}).json()
        self.assertTrue(any(item['reference_id'] == ident for item in search['results']))
        listed = self.client.get('/v1/references', params={'query': '', 'limit': 50}).json()
        self.assertTrue(any(item['reference_id'] == ident for item in listed['results']))

    def test_manual_edits_history_conflict_and_invalid_edit_preservation(self):
        imported = self.imported(text='原文记录：限值为 160 degC，仅用于测试。')
        ident = imported['reference_id']
        block = self.read(ident)['items'][0]
        self.save(ident, 1, entries=[{'key': 'limit', 'text': '采用限值前核对条件。', 'source_block_ids': [block['id']]}])
        source = Path(imported['files']['source'])
        metadata = json.loads(source.read_text(encoding='utf-8'))
        metadata['notes'] = '用户补充：这是题目附录节选，资料不完整。'
        source.write_text(json.dumps(metadata, ensure_ascii=False), encoding='utf-8-sig')
        brief = Path(imported['files']['brief'])
        brief.write_text('用户精简稿：只用于理解，引用需回看原文。', encoding='utf-8')
        edited = self.read(ident, view='organized')
        self.assertEqual(edited['revision'], 3)
        self.assertEqual(edited['brief_excerpt'], brief.read_text(encoding='utf-8'))
        original = self.read(ident, revision=2, view='organized')
        self.assertEqual(original['source']['notes'], '测试适用范围')
        self.assertNotEqual(original['brief_excerpt'], edited['brief_excerpt'])
        preserved = self.save(ident, 3, overview='补充其他章节索引')
        self.assertEqual(self.read(ident, view='organized')['brief_excerpt'], edited['brief_excerpt'])
        self.assertEqual(preserved['revision'], 4)
        stale = self.client.post('/v1/references/'+ident, json={'expected_revision': 2, 'reason': '过期版本', 'overview': '旧修改'})
        self.assertEqual(stale.status_code, 409)
        source.write_text('{ user invalid edit', encoding='utf-8')
        failed = self.client.get('/v1/references/'+ident)
        self.assertEqual(failed.status_code, 422)
        self.assertEqual(source.read_text(encoding='utf-8'), '{ user invalid edit')
        self.assertEqual(self.read(ident, revision=1)['revision'], 1)

    def test_image_transcription_and_native_summary_separation(self):
        from PIL import Image, ImageDraw
        from kernel.server import DATA_ROOT
        path = DATA_ROOT/(self.project_id+'.png')
        with Image.new('RGB', (500, 160), 'white') as image:
            ImageDraw.Draw(image).text((25, 50), 'Limit: 160 degC; test only', fill='black')
            image.save(path)
        with patch('app.vision.transcribe_image', side_effect=AssertionError('不应调用模型API')):
            imported = self.imported(path=str(path))
            ident = imported['reference_id']
            self.assertEqual(imported['extraction_status'], 'awaiting_session')
            anchor = self.read(ident)['items'][0]['id']
            rendered = self.read(ident, view='image', page=1)
            self.assertTrue(Path(rendered['image']['localPath']).is_file())
            self.assertEqual(rendered['source_block_ids'], [anchor])
            self.save(ident, 1, entries=[
                {'key': 'image-text', 'kind': 'transcription', 'text': 'Limit: 160 degC; test only', 'source_block_ids': [anchor]},
                {'key': 'image-summary', 'kind': 'summary', 'text': '测试图像说明限值，尚需核对。', 'source_block_ids': [anchor]},
            ])
        use_path = '/v1/projects/'+self.project_id+'/references'
        bad = self.client.post(use_path, json={'reference_id': ident, 'revision': 2,
                               'block_ids': ['s-image-summary'], 'purpose': '禁止摘要冒充原文'})
        self.assertEqual(bad.status_code, 422)
        used = self.post(use_path, {'reference_id': ident, 'revision': 2, 'block_ids': ['t-image-text'], 'purpose': '引用测试图像'})
        from app.fulltext import read_document, DocumentReadInput
        snapshot = read_document(self.project_id, DocumentReadInput(snapshot_id=used['snapshot_id']))
        self.assertTrue(snapshot['blocks'][0]['ocr'])
        self.assertTrue(snapshot['blocks'][0]['transcription_review_required'])
        self.assertEqual(self.client.get('/v1/references/'+ident+'?view=image&page=2').status_code, 422)
        path.unlink()

    def test_mixed_and_scanned_pdf_full_index_and_render(self):
        from PIL import Image
        from reportlab.pdfgen.canvas import Canvas
        from reportlab.lib.utils import ImageReader
        from kernel.server import DATA_ROOT
        path = DATA_ROOT/(self.project_id+'.pdf')
        with Image.new('RGB', (200, 100), 'white') as image:
            drawing = Canvas(str(path))
            drawing.drawString(50, 750, 'Mixed page native text: limit 160 degC')
            drawing.drawImage(ImageReader(image), 50, 500, width=200, height=100)
            drawing.showPage()
            drawing.drawImage(ImageReader(image), 50, 500, width=200, height=100)
            drawing.showPage()
            drawing.save()
        with patch('app.vision.transcribe_image', side_effect=AssertionError('不应调用模型API')):
            imported = self.imported(path=str(path))
            ident = imported['reference_id']
            self.assertEqual(imported['progress']['total_pages'], 2)
            blocks = self.read(ident, limit=50)['items']
            self.assertTrue(any(block['text'].startswith('Mixed page') for block in blocks))
            self.assertEqual(sum(block['kind'] == 'page_image' for block in blocks), 2)
            for page in (1, 2):
                rendered = self.read(ident, view='image', page=page)
                self.assertEqual(rendered['image']['page'], page)
                self.assertTrue(rendered['source_block_ids'])
        path.unlink()

    def test_long_pdf_indexes_beyond_default_extraction_window(self):
        from reportlab.pdfgen.canvas import Canvas
        from kernel.server import DATA_ROOT
        path = DATA_ROOT/(self.project_id+'.pdf')
        canvas = Canvas(str(path))
        for page in range(1, 53):
            canvas.drawString(50, 750, f'Page {page} - reference extraction test')
            canvas.showPage()
        canvas.save()
        imported = self.imported(path=str(path))
        last = self.read(imported['reference_id'], page=52)
        self.assertTrue(any('Page 52' in item['text'] for item in last['items']))
        self.assertEqual(imported['progress']['total_pages'], 52)
        path.unlink()

    def test_shared_reuse_pinned_citation_evidence_and_bundle_hashes(self):
        imported = self.imported(text='测试原文条款：设计温度不得超过 160 degC，使用前核对适用条件。')
        ident = imported['reference_id']
        block = self.read(ident)['items'][0]
        self.save(ident, 1, entries=[{'key': 'temperature', 'text': '设计温度限值摘要。', 'source_block_ids': [block['id']]}])
        request = {'reference_id': ident, 'revision': 2, 'block_ids': [block['id']], 'purpose': '测试设计依据'}
        used = self.post('/v1/projects/'+self.project_id+'/references', request)
        again = self.post('/v1/projects/'+self.project_id+'/references', request)
        self.assertEqual(used['snapshot_id'], again['snapshot_id'])
        project = self.client.get('/v1/projects/'+self.project_id).json()['project']
        self.assertEqual(len(project['references']), 1)
        self.assertEqual(len(project['files']), 6)
        from app.evidence import save_claim, add_evidence, ClaimInput, EvidenceInput
        claim = save_claim(self.project_id, ClaimInput(text='回归测试引用设计温度限值原文'))
        linked = add_evidence(self.project_id, EvidenceInput(claim_id=claim['id'], expected_revision=claim['revision'],
                              snapshot_id=used['snapshot_id'], block_id=block['id'], quote='设计温度不得超过 160 degC'))
        self.assertTrue(linked['evidence'][0]['integrity']['quote_found'])
        self.assertEqual(linked['evidence'][0]['data']['shared_reference']['revision'], 2)
        self.save(ident, 2, source={**self.source, 'notes': '共享资料的新备注'})
        pinned = self.client.get('/v1/projects/'+self.project_id).json()['project']['references'][0]
        self.assertEqual(pinned['revision'], 2)
        bundle = self.post('/v1/projects/'+self.project_id+'/bundle', {})['bundle']
        archive_bytes = self.client.get('/v1/bundles/'+bundle['name']).content
        with zipfile.ZipFile(io.BytesIO(archive_bytes)) as archive:
            self.assertEqual(json.loads(archive.read('references.json'))[0]['revision'], 2)
            entries = json.loads(archive.read('manifest.json'))['files']
            citation = next(entry for entry in entries if entry['name'] == 'citation.json')
            provenance = json.loads(archive.read(citation['archivePath']))
            self.assertEqual(provenance['source_sha256'], ident)
            self.assertEqual(provenance['retrieved_at'], imported['created_at'])
            self.assertEqual(provenance['acquisitions'][0]['source'], imported['source'])
            self.assertEqual(provenance['acquisitions'][0]['retrieved_at'], imported['acquisitions'][0]['retrieved_at'])
            for entry in entries:
                self.assertEqual(hashlib.sha256(archive.read(entry['archivePath'])).hexdigest(), entry['sha256'])
        self.client.delete('/v1/projects/'+self.project_id)
        other = self.client.post('/v1/projects', json={'name': '另一个设计项目'}).json()['project']['projectId']
        try:
            other_use = self.post('/v1/projects/'+other+'/references', request)
            self.assertNotEqual(used['snapshot_id'], other_use['snapshot_id'])
            self.assertTrue(Path(imported['files']['original']).is_file())
        finally:
            self.client.delete('/v1/projects/'+other)

    def test_immutable_source_layer_and_canonical_citation_order(self):
        imported = self.imported(text='第一段原文：温度限值为 160 degC。\n\n第二段原文：压力限值为 1.0 MPa。')
        ident = imported['reference_id']
        blocks = self.read(ident)['items']
        ids = [block['id'] for block in blocks]
        use_path = '/v1/projects/'+self.project_id+'/references'
        request = {'reference_id': ident, 'revision': 1, 'block_ids': ids, 'purpose': '两段条款引用'}
        first = self.post(use_path, request)
        second = self.post(use_path, {**request, 'block_ids': list(reversed(ids))})
        self.assertEqual(first['snapshot_id'], second['snapshot_id'])
        original = Path(imported['files']['original'])
        original.write_bytes(b'changed source')
        self.assertEqual(self.client.get('/v1/references/'+ident).status_code, 422)
        # The project's independent snapshot still verifies against its own copy.
        from app.fulltext import read_document, DocumentReadInput
        snapshot = read_document(self.project_id, DocumentReadInput(snapshot_id=first['snapshot_id']))
        self.assertIn('160 degC', snapshot['blocks'][0]['text'])
        failed_search = self.client.get('/v1/references', params={'query': self.source['title']}).json()
        self.assertTrue(any(error['reference_id'] == ident for error in failed_search['errors']))

    def test_docx_table_formula_and_structured_partial_updates(self):
        from docx import Document
        from kernel.server import DATA_ROOT
        path = DATA_ROOT/(self.project_id+'.docx')
        document = Document()
        document.add_paragraph('Test formula: Q = U A deltaT; verify fouling conditions.')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = 'P', 'T'
        table.cell(1, 0).text, table.cell(1, 1).text = '1.0 MPa', '160 degC'
        document.save(path)
        imported = self.imported(path=str(path))
        ident = imported['reference_id']
        blocks = self.read(ident)['items']
        raw_table = next(block for block in blocks if block['kind'] == 'table')
        self.assertEqual(raw_table['cells'][1], ['1.0 MPa', '160 degC'])
        first = self.save(ident, 1, entries=[{'key': 'formula', 'kind': 'formula', 'text': 'Q = U A deltaT',
                                           'source_block_ids': [blocks[0]['id']], 'conditions': '核对污垢条件'}])
        self.save(ident, first['revision'], entries=[{'key': 'table', 'kind': 'table', 'text': 'P=1.0 MPa; T=160 degC',
                  'source_block_ids': [raw_table['id']], 'cells': raw_table['cells']}])
        organized = self.read(ident, view='organized')
        self.assertEqual(len(organized['items']), 2)
        self.assertEqual(organized['items'][1]['cells'][1][0], '1.0 MPa')
        self.assertIn('| 1.0 MPa | 160 degC |', organized['brief_excerpt'])
        self.assertEqual(self.read(ident, revision=2, view='organized')['total_items'], 1)
        path.unlink()

    def test_shared_library_process_lock(self):
        import subprocess
        import time
        from kernel.server import REFERENCE_LIBRARY, DATA_ROOT
        ready, completed = DATA_ROOT/(self.project_id+'-ready'), DATA_ROOT/(self.project_id+'-completed')
        code = (
            "from pathlib import Path; import sys; from app.references import ReferenceLibrary; "
            "library=ReferenceLibrary(Path(sys.argv[1]),Path(sys.argv[2])); "
            "Path(sys.argv[3]).write_text('ready'); "
            "\nwith library._guard(): Path(sys.argv[4]).write_text('completed')\n")
        child = None
        try:
            with REFERENCE_LIBRARY._guard():
                child = subprocess.Popen([sys.executable, '-X', 'utf8', '-c', code, str(DATA_ROOT), str(ROOT),
                                          str(ready), str(completed)], cwd=ROOT,
                                         stdout=subprocess.PIPE, stderr=subprocess.PIPE)
                deadline = time.monotonic()+10
                while not ready.exists() and time.monotonic() < deadline:
                    if child.poll() is not None:
                        break
                    time.sleep(.02)
                self.assertTrue(ready.exists(), '子进程未准备好')
                self.assertFalse(completed.exists(), '共享库进程锁没有阻止并发写入')
            output, error = child.communicate(timeout=15)
            self.assertEqual(child.returncode, 0, error.decode('utf-8', errors='replace'))
            self.assertTrue(completed.exists())
        finally:
            if child and child.poll() is None:
                child.terminate()
                child.communicate(timeout=10)

    def test_archive_failure_url_source_and_boundary(self):
        from kernel.server import DATA_ROOT
        broken = DATA_ROOT/(self.project_id+'.pdf')
        broken.write_bytes(b'%PDF-broken-original')
        imported = self.imported(path=str(broken))
        self.assertEqual(imported['extraction_status'], 'unavailable')
        self.assertEqual(Path(imported['files']['original']).read_bytes(), broken.read_bytes())
        broken.unlink()
        fake = {'content': ('<html><body><p>来源网页 '+self.project_id+'</p></body></html>').encode(),
                'final_url': 'https://example.com/reference', 'headers': {'content-type': 'text/html'},
                'redirects': [{'url': 'https://example.com/reference', 'status': 200}]}
        with patch('app.public_sources.fetch_public', return_value=fake):
            fetched = self.imported(url='https://example.com/reference')
        self.assertEqual(fetched['source']['url'], fake['final_url'])
        self.assertTrue(self.read(fetched['reference_id'])['items'])
        with tempfile.TemporaryDirectory(prefix='chemellm-outside-') as outside:
            path = Path(outside)/'outside.txt'
            path.write_text('outside', encoding='utf-8')
            response = self.client.post('/v1/references/import', json={'source': self.source, 'path': str(path)})
            self.assertEqual(response.status_code, 422)
        for fields in ({}, {'path': 'one', 'text': 'two'}):
            response = self.client.post('/v1/references/import', json={'source': self.source, **fields})
            self.assertEqual(response.status_code, 422)
        self.assertEqual(self.client.get('/v1/references/not-an-id').status_code, 422)

    def test_interrupted_publication_retains_later_manual_edit(self):
        import app.references as references
        from kernel.server import REFERENCE_LIBRARY
        imported = self.imported(text='中断恢复测试原文：保留修改记录和原始材料。')
        ident = imported['reference_id']
        source = Path(imported['files']['source'])
        original_atomic_bytes = references._atomic_bytes

        def interrupt(path, content):
            if path == source:
                raise OSError('模拟发布中断')
            return original_atomic_bytes(path, content)

        with patch.object(references, '_atomic_bytes', side_effect=interrupt):
            failed = self.client.post('/v1/references/'+ident, json={
                'expected_revision': 1, 'reason': '模拟中断', 'overview': 'AI 已保存的概览'})
            self.assertEqual(failed.status_code, 422)
        metadata = json.loads(source.read_text(encoding='utf-8'))
        metadata['notes'] = '中断后用户补充的备注'
        source.write_text(json.dumps(metadata, ensure_ascii=False), encoding='utf-8')
        recovered = self.read(ident, view='organized')
        self.assertEqual(recovered['revision'], 3)
        self.assertEqual(recovered['source']['notes'], '中断后用户补充的备注')
        self.assertEqual(recovered['overview'], 'AI 已保存的概览')
        record = REFERENCE_LIBRARY._revision(REFERENCE_LIBRARY.folder(ident), 3)
        self.assertEqual(record['actor'], 'local_edit')
        self.assertFalse((source.parent/'.pending.json').exists())


if __name__ == "__main__":
    try:
        unittest.main(verbosity=2)
    finally:
        _temporary.cleanup()
